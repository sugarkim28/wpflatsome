"""Kiểm thử với cổng giả lập: python -m unittest test_taihoadon.py"""

import io
import json
import os
import random
import re
import shutil
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import taihoadon as t

# Hình ký tự giả (toạ độ tuyệt đối từ gốc 0,0); captcha dịch chuyển từng ký tự tới vị trí ngẫu nhiên.
GLYPHS = {
    "A": "M0 10 Q2 5 5 0 Q7 5 10 10 Z M3 7 Q5 6 7 7 Z",
    "B": "M0 0 Q0 6 0 12 Q8 12 8 6 Q8 0 0 0 Z",
    "C": "M9 1 Q0 -1 0 6 Q0 13 9 11 Q8 11 9 1 Z",
    "D": "M0 0 Q0 6 0 12 Q11 12 11 6 Q11 0 0 0 Z",
    "7": "M0 0 Q4 0 9 0 Q6 6 3 12 Z",
}
ACCOUNTS = {"0309999999": "pw1", "0101234567": "pw2"}
STATE = {"tthai": 1, "k_invoice": False, "timeout_months": False}


def shift(d, dx, dy):
    out, i = [], 0
    for tok in d.replace("M", " M ").replace("L", " L ").replace("Q", " Q ").replace("Z", " Z ").split():
        if tok.isalpha():
            out.append(tok)
            continue
        out.append(str(round(float(tok) + (dx if i % 2 == 0 else dy), 2)))
        i += 1
    return " ".join(out)


def make_captcha(text):
    parts = ['<svg xmlns="http://www.w3.org/2000/svg" width="200" height="50" viewBox="0,0,200,50">',
             '<path d="M10 40 C60 5 120 45 190 10" stroke="#888" fill="none"/>']
    x = 12
    for ch in text:
        parts.append('<path fill="none" stroke="#%06x" d="%s"/>' % (random.randint(0, 0x777777),
                                                      shift(GLYPHS[ch], x + random.random() * 4, 15 + random.random() * 15)))
        x += 30
    parts.append("</svg>")
    return "".join(parts)


def make_invoice(i, day, buyer="0309999999"):
    return {"nbmst": "0101234567", "nbten": "Cong ty Ban %d" % i, "nmmst": buyer, "nmten": "Cong ty Mua",
            "khmshdon": 1, "khhdon": "C26TAA", "shdon": i, "tdlap": "%sT08:00:00.000+0000" % day,
            "tgtcthue": 1000000 * i, "tgtthue": 100000 * i, "tgtttbso": 1100000 * i,
            "tthai": STATE["tthai"] if i == 1 else 1, "ttxly": 5, "dvtte": "VND", "mhdon": "M%d" % i}


XML = ('<?xml version="1.0" encoding="UTF-8"?><HDon><DLHDon><NDHDon>'
       '<NBan><Ten>Cong ty Ban</Ten><MST>0101234567</MST><DChi>So 1 Le Loi</DChi></NBan>'
       '<NMua><Ten>Cong ty Mua</Ten><DChi>So 2 Hai Ba Trung</DChi></NMua><DSHHDVu>'
       '<HHDVu><TChat>1</TChat><STT>1</STT><THHDVu>Giay A4</THHDVu><DVTinh>Ram</DVTinh><SLuong>2</SLuong>'
       '<DGia>50000</DGia><ThTien>100000</ThTien><TSuat>10%%</TSuat></HHDVu>'
       '<HHDVu><TChat>1</TChat><STT>2</STT><THHDVu>But bi</THHDVu><DVTinh>Cay</DVTinh><SLuong>10</SLuong>'
       '<DGia>5000</DGia><ThTien>50000</ThTien><TSuat>8%%</TSuat></HHDVu>'
       '<HHDVu><TChat>4</TChat><STT>3</STT><THHDVu>Ghi chu giao hang</THHDVu></HHDVu>'
       '</DSHHDVu><TToan><THTTLTSuat><LTSuat><TSuat>10%%</TSuat><ThTien>100000</ThTien><TThue>10000</TThue></LTSuat>'
       '<LTSuat><TSuat>8%%</TSuat><ThTien>50000</ThTien><TThue>4000</TThue></LTSuat></THTTLTSuat></TToan>'
       '</NDHDon><TTChung><THDon>HÓA ĐƠN GIÁ TRỊ GIA TĂNG</THDon><HTTToan>TM/CK</HTTToan><MSTTCGP>0101243150</MSTTCGP><TTKhac><TTin><TTruong>TransactionID</TTruong><KDLieu>string</KDLieu><DLieu>MISA%sABC</DLieu></TTin></TTKhac></TTChung></DLHDon><SHDon>%s</SHDon></HDon>')


class FakePortal(BaseHTTPRequestHandler):
    calls = []
    captchas = {}
    logins = []
    headers = []

    def log_message(self, *a):
        pass

    def _json(self, code, data):
        body = json.dumps(data).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        FakePortal.headers.append(("POST", self.path, dict(self.headers)))
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        FakePortal.logins.append(body["username"])
        if FakePortal.captchas.pop(body["ckey"], None) != body["cvalue"]:
            return self._json(400, {"message": "Mã captcha không đúng"})
        if ACCOUNTS.get(body["username"]) != body["password"]:
            return self._json(400, {"message": "Tên đăng nhập hoặc mật khẩu không đúng"})
        self._json(200, {"token": "TOKEN-" + body["username"]})

    def do_GET(self):
        u = urllib.parse.urlparse(self.path)
        q = dict(urllib.parse.parse_qsl(u.query))
        FakePortal.calls.append((u.path, q))
        FakePortal.headers.append(("GET", u.path, dict(self.headers)))
        if u.path == "/":
            body = b"<html>portal</html>"
            self.send_response(200)
            self.send_header("Set-Cookie", "TS01=abc; Path=/")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            return self.wfile.write(body)
        if u.path.startswith("/category/public/dsdkts/"):
            mst = u.path.split("/")[4]
            if mst == "0309999999":
                return self._json(200, {"mst": mst, "tennnt": "CÔNG TY TNHH MUA", "tthai": "00", "tencqt": "Thuế TP HCM",
                                        "dctsdchi": "So 2 Hai Ba Trung", "dctsxaten": "Phường Bến Nghé", "dctstinhten": "TP HCM"})
            return self._json(400, {"message": "Không tìm thấy thông tin MST"})
        if u.path == "/captcha":
            text = "".join(random.sample(list(GLYPHS), len(GLYPHS)))
            key = "K%d" % len(FakePortal.calls)
            FakePortal.captchas[key] = text
            return self._json(200, {"key": key, "content": make_captcha(text)})
        auth = self.headers.get("Authorization", "")
        if not auth.startswith("Bearer TOKEN-"):
            return self._json(401, {"message": "unauthorized"})
        mst = auth[len("Bearer TOKEN-"):]
        if u.path.endswith(("/invoices/purchase", "/invoices/sold")) and "," in q.get("sort", ""):
            return self._json(400, {"message": "Không hỗ trợ sắp xếp theo nhiều trường"})
        if STATE["timeout_months"] and u.path.endswith("/invoices/purchase"):
            a, b = re.findall(r"(\d\d)/(\d\d)/(\d{4})", q["search"])[:2]
            if (date(int(b[2]), int(b[1]), int(b[0])) - date(int(a[2]), int(a[1]), int(a[0]))).days > 7:
                return self._json(500, {"message": "Row retrieval response timeout of 10000, missing responses from nodes"})
        if u.path == "/query/invoices/purchase":
            if "ttxly==5" not in q["search"]:
                return self._json(200, {"datas": [], "total": 0})
            if "01/09/2026" in q["search"]:
                all_ = [make_invoice(i, "2026-09-%02d" % (i % 28 + 1), mst) for i in range(1, 61)]
                page = all_[50:] if q.get("state") == "S2" else all_[:50]
                return self._json(200, {"datas": page, "total": 60, "state": None if q.get("state") else "S2"})
            datas = [make_invoice(99, "2026-10-02", mst)]
            if STATE["k_invoice"]:  # HĐ không mã: cổng không có XML gốc
                k = make_invoice(100, "2026-10-01", mst)
                k.update(khhdon="K26THA", tdlap="2026-10-01T17:00:00Z", ttxly=6)
                datas.append(k)
            return self._json(200, {"datas": datas, "total": len(datas)})
        if u.path in ("/query/invoices/sold", "/sco-query/invoices/purchase", "/sco-query/invoices/sold"):
            return self._json(200, {"datas": [], "total": 0})
        if u.path == "/query/invoices/detail":  # dữ liệu "Xem hoá đơn" của cổng
            return self._json(200, {"nbten": "Cong ty Ban K", "nbmst": q.get("nbmst"), "nbdchi": "So 2 Le Loi",
                                    "nmten": "CONG TY A", "nmmst": "0309999999", "khmshdon": 1, "khhdon": q.get("khhdon"),
                                    "shdon": int(q.get("shdon") or 0), "tdlap": "2026-10-01T17:00:00Z", "thtttoan": "TM",
                                    "tgtcthue": 200000, "tgtthue": 16000, "tgtttbso": 216000,
                                    "tgtttbchu": "Hai trăm mười sáu nghìn đồng",
                                    "hdhhdvu": [{"stt": 1, "ten": "Nước suối Lavie", "dvtinh": "Thùng", "sluong": 2,
                                                 "dgia": 100000, "thtien": 200000, "tsuat": "8%"}],
                                    "thttltsuat": [{"tsuat": "8%", "thtien": 200000, "tthue": 16000}]})
        if u.path == "/query/invoices/export-xml" and q.get("khhdon", "").startswith("K"):
            return self._json(500, {"message": "Không tồn tại hồ sơ gốc của hóa đơn."})
        if u.path == "/query/invoices/export-xml":
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as z:
                z.writestr("invoice.xml", XML % (q["shdon"], q["shdon"]))
                # bản HTML của cổng: nội dung do details.js dựng (cần chạy JavaScript)
                z.writestr("invoice.html", '<!doctype html><html><head><meta charset="utf-8"><script src="details.js">'
                                           '</script></head><body><h1>HÓA ĐƠN CỦA CỔNG THUẾ</h1><p id="so"></p></body></html>')
                z.writestr("details.js", "document.addEventListener('DOMContentLoaded',function(){"
                                         "document.getElementById('so').textContent='Số hóa đơn: %s';});" % q["shdon"])
            data = buf.getvalue()
            self.send_response(200)
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            return self.wfile.write(data)
        self._json(404, {"message": "not found"})


class AutoAnswer(threading.Thread):
    """Đóng vai người dùng: thấy captcha thì gõ đáp án đúng."""

    def __init__(self, job):
        super().__init__(daemon=True)
        self.job = job
        self.asked = 0

    def run(self):
        while self.job.running:
            need = self.job.snapshot()["need_captcha"]
            if need:
                for text in list(FakePortal.captchas.values()):
                    if make_captcha_text_matches(need["svg"], text):
                        self.asked += 1
                        self.job.answer(text.lower())  # người dùng gõ chữ thường vẫn được
                        break
            time.sleep(0.02)


def make_captcha_text_matches(svg, text):
    glyphs = t.captcha_glyphs(svg)
    if len(glyphs) != len(text):
        return False
    ref = {ch: t.captcha_glyphs(make_captcha(ch))[0][1] for ch in set(text)}
    return all(t._dist(g[1], ref[ch]) < 0.5 for g, ch in zip(glyphs, text))



class mock_answer:
    """Trả lời captcha cho lượt đồng bộ của một người dùng trên App (như AutoAnswer)."""
    def __init__(self, app, username):
        self.app, self.username, self.stop = app, username, False

    def __enter__(self):
        def loop():
            while not self.stop:
                job = self.app.jobs.get(self.username)
                cap = job and job.need_captcha
                if cap:
                    for text in list(FakePortal.captchas.values()):
                        if make_captcha_text_matches(cap["svg"], text):
                            job.answer(text)
                            break
                time.sleep(0.02)
        self.th = threading.Thread(target=loop, daemon=True)
        self.th.start()
        return self

    def __exit__(self, *a):
        self.stop = True
        self.th.join(1)

class Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = ThreadingHTTPServer(("127.0.0.1", 0), FakePortal)
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()
        cls.base = "http://127.0.0.1:%d" % cls.srv.server_port

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        FakePortal.calls, FakePortal.captchas, FakePortal.logins, FakePortal.headers = [], {}, [], []
        STATE["tthai"] = 1
        STATE["k_invoice"] = False
        STATE["timeout_months"] = False
        self.app = t.App(self.tmp, lambda: t.HoaDonClient(self.base, delay=0))

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def run_job(self, msts, kinds=("purchase", "sold"), test=False, start=date(2026, 9, 1), end=date(2026, 10, 31)):
        job = t.Job()
        job.running = True
        helper = AutoAnswer(job)
        helper.start()
        t.run_batch(job, self.app.store, self.app.solver, msts, list(kinds), start, end, True, True,
                    self.tmp, test, self.app.client_factory)
        helper.join(1)
        return job, helper

    def test_month_ranges(self):
        r = t.month_ranges(date(2026, 1, 15), date(2026, 3, 10))
        self.assertEqual(r, [(date(2026, 1, 15), date(2026, 1, 31)), (date(2026, 2, 1), date(2026, 2, 28)),
                             (date(2026, 3, 1), date(2026, 3, 10))])
        self.assertEqual(t.parse_date("05/09/2026"), date(2026, 9, 5))
        self.assertEqual(t.parse_date("2026-09-05"), date(2026, 9, 5))

    def test_captcha_learning(self):
        solver = t.CaptchaSolver(os.path.join(self.tmp, "c.json"))
        svg = make_captcha("ABCD7")
        self.assertIsNone(solver.solve(svg))
        self.assertTrue(solver.learn(svg, "ABCD7"))
        for _ in range(20):
            text = "".join(random.choice(list(GLYPHS)) for _ in range(6))
            self.assertEqual(solver.solve(make_captcha(text)), text)
        # Đọc lại từ file
        self.assertEqual(t.CaptchaSolver(os.path.join(self.tmp, "c.json")).solve(make_captcha("7DCBA")), "7DCBA")

    def test_real_captcha_builtin_table(self):
        """Captcha thật của cổng: giải ngay bằng bảng chữ ký có sẵn, không cần học."""
        solver = t.CaptchaSolver(os.path.join(self.tmp, "c.json"))
        folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_data")
        for name in sorted(os.listdir(folder)):
            if name.startswith("captcha_"):
                svg = open(os.path.join(folder, name), encoding="utf-8").read()
                self.assertEqual(solver.solve(svg), name[8:-4])

    def test_mst_lookup(self):
        c = t.HoaDonClient(self.base, delay=0)
        d = c.lookup_company("0309999999")
        self.assertEqual((d["ten"], d["tthai_text"]), ("CÔNG TY TNHH MUA", "Đang hoạt động"))
        self.assertIn("Phường Bến Nghé", d["dia_chi"])
        with self.assertRaisesRegex(t.PortalError, "Không tìm thấy"):
            c.lookup_company("0101234567")
        with self.assertRaises(ValueError):
            c.lookup_company("123")
        # Nhập danh sách bỏ trống tên → tự lấy tên
        added, errors = self.app.store.import_text("0309999999\t\tpw1\n")
        self.assertEqual(self.app.store.get("0309999999")["ten"], "")

    def test_store_and_import(self):
        added, errors = self.app.store.import_text("0309999999\tCONG TY A\tpw1\n0101234567 | CONG TY B | pw2\nabc\tX\n")
        self.assertEqual(added, 2)
        self.assertEqual(len(errors), 1)
        c = self.app.store.get("0309999999")
        self.assertEqual(t.unprotect(c["pw"]), "pw1")
        self.assertNotIn("pw1", open(self.app.store.path, encoding="utf-8").read())
        self.assertTrue(all("pw" not in x for x in self.app.store.public()))

    def test_batch_sync(self):
        self.app.store.import_text("0309999999\tA\tpw1\n0101234567\tB\tSAI\n")
        job, helper = self.run_job(["0309999999", "0101234567"])
        log = "\n".join(job.log)
        a, b = self.app.store.get("0309999999"), self.app.store.get("0101234567")
        # DN A: đủ 61 HĐ mua vào, XML, chi tiết hàng hoá
        self.assertEqual(a["so_hd"]["purchase"], 61, log)
        self.assertEqual(a["lich_su"]["purchase"]["moi"], 61)
        self.assertEqual(a["loi"], "")
        folder = os.path.join(self.tmp, "0309999999", "mua-vao_20260901_20261031")
        self.assertEqual(len([f for f in os.listdir(os.path.join(folder, "xml")) if f.endswith(".xml")]), 61)
        from openpyxl import load_workbook
        xlsx = os.path.join(folder, "MUA_VAO_0309999999_20260901_20261031.xlsx")
        wb = load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["HoaDon_TongQuat", "Smart_KTSC_OK", "BangKe_MuaVao", "BangKe_HoanThue_OK"])
        tq = wb["HoaDon_TongQuat"]
        self.assertEqual(tq["A1"].value, "HÓA ĐƠN MUA VÀO")
        self.assertEqual(tq["A3"].value, "Tên DN: A")
        self.assertEqual((tq["A5"].value, tq["D5"].value, tq["I5"].value, tq["L5"].value, tq["M5"].value),
                         ("V", "So 1 Le Loi", "TM/CK", "HĐ Mới", "Đã cấp MST"))
        self.assertEqual(tq["A66"].value, "Total")             # 61 HĐ ở dòng 5..65
        self.assertEqual(tq["R66"].value, "=SUBTOTAL(109,R5:R65)")
        sm = wb["Smart_KTSC_OK"]
        self.assertEqual(sm.max_row, 1 + 61 * 2)                # bỏ dòng ghi chú (TChat 4)
        row = {h.value: c.value for h, c in zip(sm[1], sm[2])}
        self.assertEqual((row["MATHANG"], row["DONVI"], row["LUONG"], row["TTVND"], row["TS_GTGT"], row["THUEVND"],
                          row["TKTHUE"], row["MADTPNCO"]), ("Giay A4", "Ram", 2, 100000, "10", 10000, "1331", "0101234567"))
        bk = wb["BangKe_MuaVao"]
        self.assertEqual(bk["B7"].value, "Kỳ tính thuế: Từ ngày 01/09/2026 đến ngày 31/10/2026")
        self.assertEqual([bk.cell(18, c).value for c in (8, 9, 10, 11, 12)],
                         ["0101234567", "Giay A4", 100000, "10", 10000])   # gom theo HĐ + thuế suất
        self.assertEqual(bk.cell(19, 11).value, "8")
        self.assertEqual(bk["B140"].value, "Tổng")              # 122 dòng: 18..139
        self.assertEqual(bk["L140"].value, "=SUM(L18:L139)")
        ht = wb["BangKe_HoanThue_OK"]
        self.assertEqual([ht.cell(17, c).value for c in (9, 10, 11, 12, 13)], ["Giay A4", "Ram", 2, 50000, 100000])
        # DN B: sai mật khẩu → ghi lỗi, chỉ thử đăng nhập đúng 1 lần, không làm hỏng DN khác
        self.assertIn("mật khẩu", b["loi"])
        self.assertEqual(FakePortal.logins.count("0101234567"), 1)
        # Chỉ DN đầu tiên phải gõ captcha; DN thứ hai đã tự giải
        self.assertEqual(helper.asked, 1)

        # Đồng bộ lại: HĐ số 1 bị huỷ → phát hiện đổi trạng thái, tải lại XML của HĐ đó thôi
        STATE["tthai"] = 6
        n_xml = sum(1 for p, _ in FakePortal.calls if p.endswith("export-xml"))
        job2, helper2 = self.run_job(["0309999999"], kinds=("purchase",))
        self.assertEqual(helper2.asked, 0, "lần 2 phải tự giải captcha")
        self.assertEqual(sum(1 for p, _ in FakePortal.calls if p.endswith("export-xml")), n_xml + 1)
        a = self.app.store.get("0309999999")
        self.assertEqual(a["lich_su"]["purchase"]["moi"], 0)
        self.assertEqual(job2.results[0]["doi"], 1)
        wb = load_workbook(xlsx)
        self.assertEqual(wb["Doi_TrangThai"].cell(2, 7).value, "Đã bị huỷ")
        self.assertEqual(wb["Smart_KTSC_CAN_XEM_XET"]["A2"].value, "HĐ Đã bị hủy")  # HĐ huỷ → cần xem xét
        self.assertEqual(wb["BangKe_MuaVao"]["B138"].value, "Tổng")                # chỉ còn 60 HĐ × 2

    def test_request_profiles(self):
        """Đăng nhập: trang chủ → captcha → authenticate, chung cookie, header tối giản.
        Tra cứu/XML: header trang tra cứu có Action, không cookie."""
        self.app.store.import_text("0309999999\tA\tpw1\n")
        self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        h = FakePortal.headers
        self.assertEqual([x[1] for x in h[:3]], ["/", "/captcha", "/security-taxpayer/authenticate"])
        for _, _, hd in h[1:3]:
            self.assertIn("TS01=abc", hd.get("Cookie", ""))
            self.assertNotIn("Origin", hd)
            self.assertNotIn("Referer", hd)
            self.assertTrue(hd.get("Request-Id") or hd.get("request-id"))
        query = next(hd for m, p, hd in h if p == "/query/invoices/purchase")
        self.assertEqual(urllib.parse.unquote(query["Action"]), "Tìm kiếm")
        self.assertEqual(query["End-Point"], "/tra-cuu/tra-cuu-hoa-don")
        self.assertNotIn("Cookie", query)
        xml = next(hd for m, p, hd in h if p == "/query/invoices/export-xml")
        self.assertEqual(urllib.parse.unquote(xml["Action"]), "Xuất xml (hóa đơn mua vào)")

    def test_invoice_store(self):
        self.app.store.import_text("0309999999\tA\tpw1\n")
        self.run_job(["0309999999"], kinds=("purchase",))
        f = {"kind": "purchase", "from": "2026-09-01", "to": "2026-09-30"}
        rows = t.query_invoices(self.tmp, "0309999999", f)
        self.assertEqual(len(rows), 60)
        r = rows[0]
        self.assertEqual((r["mst"], r["tthai"], r["kq"], r["duyet"], r["mat_hang"]),
                         ("0101234567", "HĐ Mới", "Đã cấp MST", "Chờ duyệt", "Giay A4; But bi"))
        self.assertTrue(r["xml"].endswith(".xml"))
        self.assertEqual(len(t.query_invoices(self.tmp, "0309999999", dict(f, shdon="5"))), 1)
        self.assertEqual(len(t.query_invoices(self.tmp, "0309999999", dict(f, file="no_xml"))), 0)
        # Ghi chú / duyệt / HĐ dịch vụ được giữ lại sau khi đồng bộ lại
        t.update_invoice(self.tmp, "0309999999", "purchase", r["key"], note="da doi chieu", duyet="Đã duyệt", dv=True)
        self.run_job(["0309999999"], kinds=("purchase",))
        dv = t.query_invoices(self.tmp, "0309999999", dict(f, kind="purchase_dv"))
        self.assertEqual([(x["key"], x["note"], x["duyet"]) for x in dv], [(r["key"], "da doi chieu", "Đã duyệt")])
        self.assertEqual(len(t.query_invoices(self.tmp, "0309999999", dict(f, q="doi chieu"))), 1)
        # Kết xuất theo bộ lọc
        x = t.export_invoices(self.tmp, "0309999999", "A", dict(f, kind="purchase_dv"), "xlsx")
        from openpyxl import load_workbook
        tq = load_workbook(x)["HoaDon_TongQuat"]
        self.assertEqual((tq["S5"].value, tq["T5"].value, tq["A6"].value), ("Đã duyệt", "da doi chieu", "Total"))
        two = [x["key"] for x in rows[:2]]
        self.assertEqual(len(t.query_invoices(self.tmp, "0309999999", dict(f, keys=two))), 2)  # chỉ các dòng đã chọn
        self.assertEqual(len(zipfile.ZipFile(t.export_invoices(self.tmp, "0309999999", "A", dict(f, keys=two), "xml")).namelist()), 2)
        z = t.export_invoices(self.tmp, "0309999999", "A", f, "xml")
        self.assertEqual(len(zipfile.ZipFile(z).namelist()), 60)
        with self.assertRaises(ValueError):
            t.export_invoices(self.tmp, "0309999999", "A", f, "pdf")

    def test_lookup_pdf_view(self):
        self.app.store.import_text("0309999999\tA\tpw1\n")
        self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        f = {"kind": "purchase", "from": "2026-10-01", "to": "2026-10-31"}
        r = t.query_invoices(self.tmp, "0309999999", f)[0]
        self.assertEqual(r["tra_cuu"], {"ncc": "MISA meInvoice", "url": "https://www.meinvoice.vn/tra-cuu/?sc=MISA99ABC",
                                        "field": "TransactionID", "code": "MISA99ABC", "ncc_mst": "0101243150"})
        # Viettel: "Mã số bí mật"
        self.assertEqual(t.lookup_info({"msttcgp": "0100109106", "ttkhac": {"Mã số bí mật": "X1Y2"}})["code"], "X1Y2")
        # Gắn PDF gốc
        with self.assertRaises(ValueError):
            t.attach_pdf(self.tmp, "0309999999", "purchase", r["key"], b"not a pdf")
        t.attach_pdf(self.tmp, "0309999999", "purchase", r["key"], b"%PDF-1.4 test")
        self.assertEqual(len(t.query_invoices(self.tmp, "0309999999", dict(f, file="pdf"))), 1)
        z = t.export_invoices(self.tmp, "0309999999", "A", f, "pdf")
        self.assertEqual(len(zipfile.ZipFile(z).namelist()), 1)
        # Bản in dựng từ XML
        page = t.render_invoice_html(self.tmp, "0309999999", "purchase", r["key"])
        for txt in ("HÓA ĐƠN GIÁ TRỊ GIA TĂNG", "So 1 Le Loi", "Giay A4", "But bi", "MISA99ABC"):
            self.assertIn(txt, page)
        self.assertIn("Ghi chu giao hang", page)  # dòng ghi chú hiển thị như hoá đơn gốc

    def test_session_reuse_and_progress(self):
        """Đồng bộ lần 2 dùng lại phiên đăng nhập: không hỏi captcha, không gọi authenticate."""
        self.app.store.import_text("0309999999\tA\tpw1\n")
        clients = {}

        def run():
            job = t.Job()
            job.running = True
            helper = AutoAnswer(job)
            helper.start()
            t.run_batch(job, self.app.store, self.app.solver, ["0309999999"], ["purchase"], date(2026, 10, 1),
                        date(2026, 10, 31), True, True, self.tmp, False, self.app.client_factory, clients)
            return job, helper
        job1, h1 = run()
        snap = job1.snapshot()
        self.assertEqual((snap["phase"], snap["found"], len(snap["rows"])), ("đồng bộ xong", {"purchase": 1}, 1))
        self.assertEqual(snap["rows"][0]["dongbo"], "Mới")
        self.assertTrue(snap["auth"].startswith("thành công"))
        logins = len(FakePortal.logins)
        job2, h2 = run()
        self.assertEqual((h2.asked, len(FakePortal.logins)), (0, logins))
        self.assertEqual(job2.snapshot()["rows"][0]["dongbo"], "Đã có")
        self.assertIn("dùng lại phiên", job2.snapshot()["auth"])
        # Token hết hạn → đăng nhập lại
        clients["0309999999"].login_at = 0
        run()
        self.assertEqual(len(FakePortal.logins), logins + 1)

    def test_no_original_xml_and_vn_date(self):
        self.assertEqual(t.invoice_date({"tdlap": "2026-10-01T17:00:00Z"}), date(2026, 10, 2))
        self.assertEqual(t.invoice_date({"tdlap": "2026-10-01T16:59:00Z"}), date(2026, 10, 1))
        self.assertEqual(t.invoice_date({"tdlap": "2026-10-01"}), date(2026, 10, 1))
        STATE["k_invoice"] = True
        self.app.store.import_text("0309999999\tA\tpw1\n")
        start = time.time()
        job, _ = self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        self.assertLess(time.time() - start, 5, "không được thử lại khi cổng báo không có XML gốc")
        rows = {r["so"]: r for r in job.snapshot()["rows"]}
        self.assertEqual(rows["100"]["ketqua"], "OK (HĐ không có XML gốc)")
        self.assertEqual(rows["100"]["ngay"], "02/10/2026")
        self.assertEqual(rows["99"]["ketqua"], "OK")
        xml_calls = lambda: sum(1 for p, q in FakePortal.calls if p.endswith("export-xml") and q["khhdon"].startswith("K"))
        self.assertEqual(xml_calls(), 1)
        self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        self.assertEqual(xml_calls(), 1, "lần sau không hỏi lại XML của HĐ không mã")
        detail_calls = sum(1 for p, q in FakePortal.calls if p.endswith("/detail"))
        self.assertEqual(detail_calls, 1, "HĐ không có XML: lấy dữ liệu chi tiết một lần")
        # PDF của thuế: HĐ không mã dựng từ dữ liệu chi tiết, HĐ có XML dựng từ XML
        import pypdf
        rows = {r["shdon"]: r for r in t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})}
        self.assertEqual(rows["100"]["mat_hang"], "Nước suối Lavie")
        self.assertEqual(rows["100"]["cqt"], "")
        import unittest.mock as mock
        no_browser = mock.patch.object(t, "find_browser", return_value=None)  # kiểm bản dựng dự phòng
        no_browser.start()
        self.addCleanup(no_browser.stop)
        for so, want in (("100", ("K26THA", "Nước suối Lavie", "216.000", "Hai trăm mười sáu nghìn đồng", "Cong ty Ban K")),
                         ("99", ("C26TAA", "Giay A4", "108.900.000", "Cong ty Ban"))):
            rel = t.tax_pdf(self.tmp, "0309999999", "purchase", rows[so]["key"])
            data = open(os.path.join(self.tmp, "0309999999", rel), "rb").read()
            text = " ".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(data)).pages)
            for w in want:
                self.assertIn(w, text, (so, w))
        self.assertTrue(t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})[0]["cqt"].endswith("_CQT.pdf"))
        zp = t.export_invoices(self.tmp, "0309999999", "A", {"kind": "purchase"}, "cqt")
        self.assertEqual(len(zipfile.ZipFile(zp).namelist()), 2)

    def test_portal_timeout_splits_range(self):
        """Cổng timeout khi tra cả tháng → tự chia theo tuần, vẫn lấy đủ hoá đơn."""
        STATE["timeout_months"] = True
        c = t.HoaDonClient(self.base, delay=0)
        c.token = "TOKEN-0309999999"
        import unittest.mock as m
        with m.patch("time.sleep"):
            invs = c.list_invoices("purchase", date(2026, 10, 1), date(2026, 10, 31), include_mtt=False)
        self.assertEqual([i["shdon"] for i in invs], [99])

    def test_tools_xml_view(self):
        page = t.render_xml_invoice((XML % (5, 5)).encode("utf-8"))
        for txt in ("HÓA ĐƠN GIÁ TRỊ GIA TĂNG", "Cong ty Ban", "So 1 Le Loi", "Cong ty Mua", "Giay A4", "But bi",
                    "100,000", "TM/CK", "MISA meInvoice", "MISA5ABC", "Không có chữ ký số"):
            self.assertIn(txt, page)
        with self.assertRaises(ValueError):
            t.render_xml_invoice(b"khong phai xml")
        with self.assertRaises(ValueError):
            t.render_xml_invoice(b"<a><b/></a>")

    def test_tools_pdf(self):
        import pypdf

        def pdf(n):
            w = pypdf.PdfWriter()
            for _ in range(n):
                w.add_blank_page(width=200, height=200)
            b = io.BytesIO()
            w.write(b)
            return b.getvalue()
        merged = t.merge_pdfs([("a.pdf", pdf(2)), ("b.pdf", pdf(3))])
        self.assertEqual(len(pypdf.PdfReader(io.BytesIO(merged)).pages), 5)
        with self.assertRaises(ValueError):
            t.merge_pdfs([("a.pdf", pdf(1))])
        with self.assertRaisesRegex(ValueError, "b.pdf"):
            t.merge_pdfs([("a.pdf", pdf(1)), ("b.pdf", b"hong")])
        self.assertEqual(t.parse_page_ranges("1-3, 5", 6), [[1, 2, 3], [5]])
        self.assertEqual(t.parse_page_ranges("", 2), [[1], [2]])
        for bad in ("0-2", "4-9", "a"):
            with self.assertRaises(ValueError):
                t.parse_page_ranges(bad, 5)
        z = zipfile.ZipFile(io.BytesIO(t.split_pdf("hoa don.pdf", merged, "1-2, 5")))
        self.assertEqual(sorted(z.namelist()), ["hoa don_trang_1-2.pdf", "hoa don_trang_5.pdf"])
        self.assertEqual(len(pypdf.PdfReader(io.BytesIO(z.read("hoa don_trang_1-2.pdf"))).pages), 2)

    def test_misa_original_pdf(self):
        """MISA giả mô phỏng trang tra cứu thật: ext có dấu '_' (J1V4E6D_), link DownloadHandler trong trang."""
        state = {"link_in_page": False, "calls": [], "direct": False,
                 "xml": b'<?xml version="1.0"?><HDon><DLHDon>misa</DLHDon></HDon>'}

        class FakeMisa(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_GET(self):
                u = urllib.parse.urlparse(self.path)
                q = dict(urllib.parse.parse_qsl(u.query))
                state["calls"].append(self.path)
                body, ctype = b"<html>not found</html>", "text/html"
                if u.path == "/tra-cuu/" and state["link_in_page"]:
                    body = ('<iframe src="tra-cuu/DownloadHandler.ashx?Type=pdf&amp;Viewer=1&amp;ext=PAGE1234&amp;Code=%s">'
                            % q.get("sc")).encode()
                elif u.path == "/tra-cuu/DownloadHandler.ashx" and state["direct"] and q.get("Type") == "pdf" \
                        and q.get("Code") == "MISA99ABC" and "ext" not in q:
                    body, ctype = b"%PDF-1.4 truc tiep", "application/pdf"
                elif u.path == "/tra-cuu/GetRequestTimeEnCode":
                    body = b'"J1V4E6D_"'
                elif u.path == "/tra-cuu/tra-cuu/DownloadHandler.ashx" and q.get("Code") == "MISA99ABC" \
                        and q.get("ext") in ("J1V4E6D_", "PAGE1234"):
                    body, ctype = b"%PDF-1.4 misa goc", "application/pdf"
                    if q.get("Type") == "xml":
                        body, ctype = state["xml"], "text/xml"
                self.send_response(200)
                self.send_header("Content-Type", ctype)
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
        srv = ThreadingHTTPServer(("127.0.0.1", 0), FakeMisa)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        base = "http://127.0.0.1:%d" % srv.server_port
        import unittest.mock as m
        try:
            with m.patch.multiple(t, MISA_WWW=base, MISA_APEX=base, MISA_DL=base):
                state["direct"] = True
                self.assertEqual(t.misa_pdf("MISA99ABC"), b"%PDF-1.4 truc tiep")      # link nút "Tải PDF" của MISA
                self.assertTrue(state["calls"][-1].startswith("/tra-cuu/DownloadHandler.ashx?Type=pdf&Code=MISA99ABC"))
                state["direct"] = False
                self.assertEqual(t.misa_pdf("MISA99ABC"), b"%PDF-1.4 misa goc")       # qua GetRequestTimeEnCode
                state["link_in_page"] = True
                self.assertEqual(t.misa_pdf("MISA99ABC"), b"%PDF-1.4 misa goc")       # qua link có sẵn trong trang
                self.assertIn("ext=PAGE1234", state["calls"][-1])
                with self.assertRaisesRegex(t.PortalError, "GetRequestTimeEnCode trả"):
                    t.misa_pdf("SAIMA")
                # Đồng bộ kèm tải PDF gốc
                self.app.store.import_text("0309999999\tA\tpw1\n")
                job = t.Job()
                job.running = True
                job.want_pdf = True
                helper = AutoAnswer(job)
                helper.start()
                with m.patch("time.sleep"):
                    t.run_batch(job, self.app.store, self.app.solver, ["0309999999"], ["purchase"], date(2026, 10, 1),
                                date(2026, 10, 31), True, True, self.tmp, False, self.app.client_factory)
                # XML tải về không đúng hoá đơn → không gắn, ghi lý do
                self.assertEqual(job.snapshot()["rows"][0]["ketqua"], "OK + PDF gốc")
                self.assertRegex("\n".join(job.log), "chưa lấy được XML gốc .*(không khớp|không đọc được)")
                r = t.query_invoices(self.tmp, "0309999999", {"kind": "purchase", "file": "pdf"})
                self.assertEqual(len(r), 1)
                self.assertEqual(r[0]["xml_goc"], "")
                nb, mau, kh, so = r[0]["key"].split("|")
                state["xml"] = ('<?xml version="1.0"?><HDon><DLHDon><TTChung><KHMSHDon>%s</KHMSHDon><KHHDon>%s</KHHDon>'
                                '<SHDon>%s</SHDon></TTChung><NDHDon><NBan><MST>%s</MST></NBan></NDHDon></DLHDon></HDon>'
                                % (mau, kh, so, nb)).encode()
                got = t.fetch_original(self.tmp, "0309999999", "purchase", r[0]["key"])
            self.assertTrue(got["xml"].endswith("_goc-ncc.xml"), got)
            r = t.query_invoices(self.tmp, "0309999999", {"kind": "purchase", "file": "pdf"})
            self.assertFalse(r[0]["need_goc"])
            with open(os.path.join(t.company_dir(self.tmp, "0309999999"), r[0]["xml_goc"]), "rb") as fh:
                self.assertEqual(fh.read(), state["xml"])
        finally:
            srv.shutdown()
            srv.server_close()
        self.assertFalse(t.can_fetch_pdf({"ncc": "Viettel S-Invoice", "ncc_mst": "0100109106", "code": "A"})[0])

    def test_match_pdf_text(self):
        """Gắn PDF có sẵn: các bẫy gặp trên PDF thật (số tiền 308.448, địa chỉ 'Số 412', HĐ điều chỉnh nhắc số HĐ cũ)."""
        buyer = "0309999999"

        def ent(nb, kh, so):
            return {"%s|1|%s|%d" % (nb, kh, so): {"inv": {"nbmst": nb, "nmmst": buyer}}}
        entries = {}
        for nb, kh, so in [("0101111111", "C26TYY", 178), ("0101111111", "C26TYY", 308), ("0101111111", "C26TYY", 288),
                           ("0102222222", "C26TSG", 7647), ("0102222222", "C26TSG", 412), ("0103333333", "K26TAB", 51680635),
                           ("0103333333", "K26TAB", 414)]:
            entries.update(ent(nb, kh, so))
        m = lambda text: [k.split("|")[3] for _, k in t.match_pdf_text(text, {"purchase": entries})]
        misa = "CÔNG TY A 0101111111 (Sign): 1C26TYY 00000178(No): Ký hiệu ... Cộng tiền hàng 308.448 Mã số thuế %s" % buyer
        self.assertEqual(m(misa), ["178"])
        dc = ("0101111111 (Sign): 1C26TYY 00000308(No): Ký hiệu ... Điều chỉnh cho hoá đơn 1C26TYY số 00000288 %s" % buyer)
        self.assertEqual(m(dc), ["308"])
        viettel = "Địa chỉ: Số 412 Nguyễn Thị Minh Khai Mã số thuế: 0 1 0 2 2 2 2 2 2 2 Ký hiệu (Serial): 1C26TSG Số (No.): 7647 MST %s" % buyer
        self.assertEqual(m(viettel), ["7647"])
        bhx = "Mã số thuế: 0103333333 HÓA ĐƠN 51680635GIÁ TRỊ GIA TĂNG Số: Ký hiệu: 1K26TAB CỬA HÀNG BHX SỐ 414 %s" % buyer
        self.assertEqual(m(bhx), ["51680635"])
        self.assertEqual(m("Ký hiệu 1C26TYY Số: 999 MST 0101111111"), [])  # số không có trong kho
        self.assertEqual(m("Ký hiệu 1C26TYY Số: 178 MST 0109999999"), [])  # MST không khớp

    def test_scan_downloads(self):
        """Trình duyệt tải PDF gốc về Downloads → phần mềm tự gắn vào đúng hoá đơn."""
        import unittest.mock as m
        self.app.store.import_text("0309999999\tA\tpw1\n")
        self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        row = t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})[0]
        self.assertTrue(row["pdf_url"].endswith("/tra-cuu/DownloadHandler.ashx?Type=pdf&Code=MISA99ABC"))
        dl = os.path.join(self.tmp, "Downloads")
        os.makedirs(dl)
        since = time.time()
        old = os.path.join(dl, "cu.pdf")
        open(old, "wb").write(b"%PDF cu")
        os.utime(old, (since - 3600, since - 3600))                    # file cũ: bỏ qua
        open(os.path.join(dl, "MISA99ABC.pdf"), "wb").write(b"%PDF moi")
        open(os.path.join(dl, "ghi chu.txt"), "w").write("x")
        texts = {b"%PDF moi": "Ký hiệu 1C26TAA Số: 99 Mã số thuế 0101234567", b"%PDF cu": "Số: 99 1C26TAA 0101234567"}
        seen = set()
        with m.patch.object(t, "_pdf_text", lambda data: texts[data]):
            res = t.scan_downloads(self.tmp, "0309999999", dl, since, seen)
            self.assertEqual([(r["file"], r["ok"]) for r in res], [("MISA99ABC.pdf", True)])
            self.assertEqual(t.scan_downloads(self.tmp, "0309999999", dl, since, seen), [])   # không đọc lại
        self.assertTrue(t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})[0]["pdf"])

    def test_browser_download_settings(self):
        base = os.path.join(self.tmp, "LocalAppData")
        chrome = os.path.join(base, "Google", "Chrome", "User Data")
        for prof, prefs in (("Default", {"download": {"default_directory": "C:\\Users\\A\\Desktop", "prompt_for_download": True}}),
                            ("Profile 1", {"download": {"default_directory": "D:\\TaiVe"}}), ("System Profile", {})):
            os.makedirs(os.path.join(chrome, prof))
            t._save_json(os.path.join(chrome, prof, "Preferences"), prefs)
        st = t.browser_download_settings(base)
        self.assertEqual(st, {"dirs": ["C:\\Users\\A\\Desktop", "D:\\TaiVe"], "prompt": True})
        self.assertEqual(t.browser_download_settings(os.path.join(self.tmp, "khong-co")), {"dirs": [], "prompt": False})

    def test_easyinvoice_lookup(self):
        lk = t.lookup_info({"msttcgp": "0105987432", "ttkhac": {"Mã tra cứu": "8D3VYYQYD"}, "nb": {"MST": "0308783233"}})
        self.assertEqual((lk["ncc"], lk["url"], lk["code"]), ("EasyInvoice (SoftDreams)", "http://0308783233hd.easyinvoice.com.vn", "8D3VYYQYD"))

    def test_expired_token(self):
        c = t.HoaDonClient(self.base, delay=0)
        c.token = "bad"
        with self.assertRaisesRegex(t.PortalError, "hết hạn"):
            c.list_invoices("sold", date(2026, 9, 1), date(2026, 9, 30))

    def test_local_ui(self):
        srv = ThreadingHTTPServer(("127.0.0.1", 0), t.make_handler(self.app, 0))
        port = srv.server_port
        srv.RequestHandlerClass = t.make_handler(self.app, port)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        base = "http://127.0.0.1:%d" % port

        def call(path, body=None, key=True):
            req = urllib.request.Request(base + path, data=json.dumps(body).encode() if body is not None else None,
                                         headers={"X-App-Key": self.app.key} if key else {})
            return json.loads(urllib.request.urlopen(req).read())
        try:
            self.assertIn(self.app.key, urllib.request.urlopen(base + "/").read().decode())
            with self.assertRaises(urllib.error.HTTPError) as cm:
                call("/api/state", key=False)
            self.assertEqual(cm.exception.code, 403)
            d = call("/api/company/import", {"text": "0309999999\t\tpw1\n0101234567\t\tpw2"})
            self.assertEqual(self.app.store.get("0309999999")["ten"], "CÔNG TY TNHH MUA")   # tự lấy tên theo MST
            self.assertEqual(len(d["errors"]), 1)                                          # MST không tra được
            self.assertEqual(call("/api/mst-lookup", {"mst": "0309999999"})["ten"], "CÔNG TY TNHH MUA")
            rows = call("/api/tool/mst", {"text": "0309999999\n0101234567"})["rows"]
            self.assertEqual([(r["ten"], bool(r["loi"])) for r in rows], [("CÔNG TY TNHH MUA", False), ("", True)])
            import base64
            v = call("/api/tool/xml-view", {"data": base64.b64encode((XML % (7, 7)).encode()).decode()})
            self.assertIn("Giay A4", v["html"])
            call("/api/company/save", {"mst": "0309999999", "ten": "A", "password": "pw1", "vao": True, "ra": False})
            call("/api/company/save", {"mst": "0309999999", "ghichu": "goi lai"})
            st = call("/api/state")
            self.assertEqual(st["companies"][0]["ghichu"], "goi lai")
            self.assertTrue(st["companies"][0]["co_mk"])
            self.assertFalse(st["companies"][0]["ra"])
            self.assertNotIn("pw1", json.dumps(st))
            # Mở file: cần khoá phiên, không cho đi ra ngoài thư mục DN
            d = os.path.join(self.tmp, "0309999999")
            os.makedirs(d, exist_ok=True)
            open(os.path.join(d, "a.xml"), "w").write("<x/>")
            q = lambda p, k=self.app.key: base + "/file?" + urllib.parse.urlencode({"k": k, "mst": "0309999999", "p": p})
            self.assertEqual(urllib.request.urlopen(q("a.xml")).read(), b"<x/>")
            zr = urllib.request.urlopen(q("a.xml") + "&zip=1")  # XML nén .zip để trình duyệt không chặn
            self.assertIn('filename="a.zip"', zr.headers["Content-Disposition"])
            self.assertEqual(zipfile.ZipFile(io.BytesIO(zr.read())).read("a.xml"), b"<x/>")
            for bad in (q("a.xml", "sai"), q("../_cau-hinh/doanh-nghiep.json")):
                with self.assertRaises(urllib.error.HTTPError):
                    urllib.request.urlopen(bad)
        finally:
            srv.shutdown()
            srv.server_close()

    def test_web_server_mode(self):
        """Bản web: đăng nhập, CSRF, nhân viên chỉ thấy DN được giao, quản trị quản lý người dùng."""
        t.set_secret_key(os.path.join(self.tmp, "_cau-hinh", "khoa.key"))
        try:
            app = t.App(self.tmp, lambda: t.HoaDonClient(self.base, delay=0), server=True)
            app.store.import_text("0309999999\tA\tpw1\n0101234567\tB\tpw2\n")
            self.assertTrue(app.store.get("0309999999")["pw"].startswith("fernet:"))
            self.assertEqual(t.unprotect(app.store.get("0309999999")["pw"]), "pw1")
            app.users.upsert("boss", ten="Sếp", role="admin", password="matkhau-boss")
            app.users.upsert("nv1", ten="Lan", role="staff", password="matkhau-nv1", msts=["0309999999"])
            self.assertNotIn("matkhau", open(app.users.path, encoding="utf-8").read())
            with self.assertRaises(ValueError):          # không được mất quản trị cuối cùng
                app.users.upsert("boss", role="staff")
            srv = ThreadingHTTPServer(("127.0.0.1", 0), t.make_handler(app, 0))
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            base = "http://127.0.0.1:%d" % srv.server_port
            import http.cookiejar

            class S:
                def __init__(s):
                    s.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
                    s.key = ""

                def login(s, u, pw):
                    s.call("/api/login", {"username": u, "password": pw})
                    page = s.op.open(base + "/").read().decode()
                    s.key = re.search(r'const KEY = "([^"]+)"', page).group(1)

                def call(s, path, body=None, key=None):
                    req = urllib.request.Request(base + path, data=json.dumps(body).encode() if body is not None else None,
                                                 headers={"X-App-Key": s.key if key is None else key,
                                                          "Content-Type": "application/json"})
                    return json.loads(s.op.open(req).read())

                def code(s, *a, **k):
                    try:
                        s.call(*a, **k)
                        return 200
                    except urllib.error.HTTPError as e:
                        return e.code
            try:
                anon = S()
                self.assertIn("Đăng nhập", anon.op.open(base + "/").read().decode())   # chưa đăng nhập → trang đăng nhập
                self.assertEqual(anon.code("/api/state"), 401)
                self.assertEqual(anon.code("/api/login", {"username": "nv1", "password": "sai-mat-khau"}), 401)
                nv, boss = S(), S()
                nv.login("nv1", "matkhau-nv1")
                boss.login("boss", "matkhau-boss")
                self.assertEqual(nv.code("/api/state", key="sai"), 403)                # thiếu khoá CSRF
                st = nv.call("/api/state")
                self.assertEqual([c["mst"] for c in st["companies"]], ["0309999999"])
                self.assertEqual(st["me"]["role"], "staff")
                self.assertEqual(len(boss.call("/api/state")["companies"]), 2)
                self.assertEqual(nv.code("/api/invoices", {"mst": "0101234567"}), 403)
                self.assertEqual(nv.code("/api/company/delete", {"mst": "0309999999"}), 403)
                self.assertEqual(nv.code("/api/company/save", {"mst": "0101234567", "ghichu": "x"}), 403)
                self.assertEqual(nv.code("/api/users", {}), 403)
                self.assertEqual(nv.code("/api/company/save", {"mst": "0309999999", "ghichu": "goi lai"}), 200)
                d = os.path.join(self.tmp, "0101234567")
                os.makedirs(d, exist_ok=True)
                open(os.path.join(d, "a.xml"), "w").write("<x/>")
                f = base + "/file?" + urllib.parse.urlencode({"mst": "0101234567", "p": "a.xml"})
                self.assertEqual(boss.op.open(f).read(), b"<x/>")
                for who in (nv, anon):
                    with self.assertRaises(urllib.error.HTTPError):
                        who.op.open(f)
                # Đồng bộ: nhân viên chỉ chạy được DN của mình
                self.assertEqual(nv.code("/api/run", {"msts": ["0101234567"], "test": True}), 400)
                with mock_answer(app, "nv1"):
                    nv.call("/api/run", {"msts": ["0309999999"], "test": True})
                    for _ in range(100):
                        if not nv.call("/api/state")["job"]["running"]:
                            break
                        time.sleep(0.05)
                self.assertIn("đăng nhập thành công", "\n".join(nv.call("/api/state")["job"]["log"]))
                self.assertFalse(boss.call("/api/state")["job"]["log"])                # mỗi người một lượt riêng
                # Quản trị giao thêm DN cho nhân viên, khoá tài khoản thì phiên bị huỷ
                boss.call("/api/user/save", {"username": "nv1", "msts": ["0309999999", "0101234567"]})
                self.assertEqual(len(nv.call("/api/state")["companies"]), 2)
                self.assertEqual(boss.code("/api/user/save", {"username": "nv1", "password": "x", "create": True}), 400)
                boss.call("/api/user/save", {"username": "nv1", "active": False})
                self.assertEqual(nv.code("/api/state"), 401)
                self.assertEqual(S().code("/api/login", {"username": "nv1", "password": "matkhau-nv1"}), 401)
                # Đổi mật khẩu
                self.assertEqual(boss.code("/api/me/password", {"old": "sai", "new": "matkhau-moi-1"}), 400)
                boss.call("/api/me/password", {"old": "matkhau-boss", "new": "matkhau-moi-1"})
                S().login("boss", "matkhau-moi-1")
                # Đăng nhập sai nhiều lần → tạm khoá
                bad = S()
                codes = [bad.code("/api/login", {"username": "boss", "password": "sai"}) for _ in range(t.LOGIN_MAX + 1)]
                self.assertEqual(codes[-1], 429)
                boss.call("/api/logout", {})
                self.assertEqual(boss.code("/api/state"), 401)
            finally:
                srv.shutdown()
                srv.server_close()
        finally:
            t._FERNET = None

    def test_auto_sync_and_notices(self):
        self.assertEqual(t.auto_range("auto", date(2026, 10, 9)), (date(2026, 9, 1), date(2026, 10, 9)))
        self.assertEqual(t.auto_range("auto", date(2026, 10, 25)), (date(2026, 10, 1), date(2026, 10, 25)))
        self.assertEqual(t.auto_range("2month", date(2026, 1, 25)), (date(2025, 12, 1), date(2026, 1, 25)))
        self.assertEqual(t.auto_range("7", date(2026, 10, 9)), (date(2026, 10, 2), date(2026, 10, 9)))
        self.app.store.import_text("0309999999\tA\tpw1\n0101234567\tB\tSAI\n")
        # Captcha của cổng giả không có trong bảng có sẵn → dạy trước để lượt tự động tự giải
        self.run_job(["0309999999"], kinds=("purchase",), test=True)
        app = t.App(self.tmp, lambda: t.HoaDonClient(self.base, delay=0), server=True)
        app.users.upsert("boss", role="admin", password="matkhau-boss")
        nv = app.users.upsert("nv1", role="staff", password="matkhau-nv1", msts=["0101234567"])
        job = app.run_auto(wait=True)
        log = "\n".join(job.log)
        self.assertIn("tự giải captcha", log)
        self.assertNotIn("ask", log)
        self.assertTrue(app.settings()["auto"]["lan_cuoi"])
        boss = app.users.get("boss")
        texts = [n["text"] for n in app.notices_for(boss)["items"]]
        self.assertTrue(any("HĐ mới" in x and "0309999999" in x for x in texts), texts)
        self.assertTrue(any("lỗi" in x and "0101234567" in x for x in texts), texts)
        mine = app.notices_for(nv)                                             # nhân viên chỉ thấy DN của mình
        self.assertTrue(mine["items"] and all(n["mst"] == "0101234567" for n in mine["items"]))
        app.mark_read(nv)
        self.assertEqual(app.notices_for(nv)["unread"], 0)
        self.assertGreater(app.notices_for(boss)["unread"], 0)
        # HĐ đổi trạng thái → thông báo cho cả lượt đồng bộ tay
        STATE["tthai"] = 6
        job, _ = app.start_job(boss, ["0309999999"], ["purchase"], date(2026, 9, 1), date(2026, 10, 31))
        helper = AutoAnswer(job)
        helper.start()
        _.join(30)
        self.assertTrue(any("đổi trạng thái" in n["text"] for n in app.notices_for(boss)["items"]))

    def test_mst_busy(self):
        a, b = t.Job(), t.Job()
        self.assertTrue(t.claim_mst("0309999999", a))
        self.assertFalse(t.claim_mst("0309999999", b))
        self.app.store.import_text("0309999999\tA\tpw1\n")
        b.running = True
        t.run_batch(b, self.app.store, self.app.solver, ["0309999999"], [], None, None, True, True, self.tmp, True,
                    self.app.client_factory)
        self.assertIn("đang được người khác đồng bộ", "\n".join(b.log))
        t.release_mst("0309999999", a)
        self.assertNotIn("0309999999", t.busy_msts())

    def test_web_behind_nginx(self):
        """FASTPANEL/nginx đổi Host thành 127.0.0.1, tên miền thật ở X-Forwarded-Host."""
        t.set_secret_key(os.path.join(self.tmp, "_cau-hinh", "khoa.key"))
        try:
            app = t.App(self.tmp, lambda: t.HoaDonClient(self.base, delay=0), server=True)
            app.allowed_hosts, app.trust_proxy = {"hoadon.congty.vn"}, True
            app.users.upsert("boss", role="admin", password="matkhau-boss")
            srv = ThreadingHTTPServer(("127.0.0.1", 0), t.make_handler(app, 0))
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            base = "http://127.0.0.1:%d" % srv.server_port

            def login(fwd_host, origin):
                req = urllib.request.Request(base + "/api/login", data=json.dumps(
                    {"username": "boss", "password": "matkhau-boss"}).encode(), headers={
                    "X-Forwarded-Host": fwd_host, "X-Forwarded-Proto": "https", "Origin": origin,
                    "X-Forwarded-For": "1.2.3.4", "Content-Type": "application/json"})
                try:
                    r = urllib.request.urlopen(req)
                    return r.status, r.headers.get("Set-Cookie", "")
                except urllib.error.HTTPError as e:
                    return e.code, ""
            try:
                code, cookie = login("hoadon.congty.vn", "https://hoadon.congty.vn")
                self.assertEqual(code, 200)
                self.assertIn("Secure", cookie)
                self.assertIn("HttpOnly", cookie)
                self.assertEqual(login("evil.com", "https://evil.com")[0], 403)                 # tên miền lạ
                self.assertEqual(login("hoadon.congty.vn", "https://evil.com")[0], 403)         # trang khác gọi vào
                # FASTPANEL gửi Host 127.0.0.1:cổng, không có X-Forwarded-Host → vẫn đăng nhập được từ đúng tên miền
                req = urllib.request.Request(base + "/api/login", data=json.dumps(
                    {"username": "boss", "password": "matkhau-boss"}).encode(),
                    headers={"Origin": "https://hoadon.congty.vn", "Content-Type": "application/json"})
                self.assertEqual(urllib.request.urlopen(req).status, 200)
                req.add_header("Origin", "https://evil.com")
                with self.assertRaises(urllib.error.HTTPError):
                    urllib.request.urlopen(req)
            finally:
                srv.shutdown()
                srv.server_close()
        finally:
            t._FERNET = None

    def test_easyinvoice_captcha(self):
        d = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_data")
        for name in ("1588", "3493", "2360", "4891", "8531"):
            with open(os.path.join(d, "easy_%s.png" % name), "rb") as f:
                self.assertEqual(t.solve_easy_captcha(f.read()), name)
        self.assertIsNone(t.solve_easy_captcha(b"khong phai anh"))

    def test_easyinvoice_pdf(self):
        import unittest.mock as mock
        """Trang tra cứu EasyInvoice giả: captcha 4 số, /Search/Search, tải PDF qua fileGuid."""
        import html as H
        d = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_data")
        caps = [open(os.path.join(d, "easy_%s.png" % n), "rb").read() for n in ("1588", "3493")]
        seen = {"posts": [], "n": 0}

        class Easy(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _out(self, body, ctype="text/html; charset=utf-8", cookie=None):
                self.send_response(200)
                self.send_header("Content-Type", ctype)
                if cookie:
                    self.send_header("Set-Cookie", cookie)
                self.end_headers()
                self.wfile.write(body)

            def do_GET(self):
                if self.path == "/":  # trang chủ không có ô tra cứu → phải sang /Search/Index
                    return self._out(b"<html>EasyInvoice</html>", cookie="ASP.NET_SessionId=abc; path=/")
                if self.path == "/Search/Index":
                    return self._out(b'<form action="/Search/Search" id="Search" method="post">'
                                     b'<input id="typeSearch" name="typeSearch" type="hidden" value="fKeySearch">'
                                     b'<input id="iFkey" name="FKey" value=""><input id="Capcha" name="Capcha" value="">'
                                     b'</form>', cookie="ASP.NET_SessionId=abc; path=/")
                if self.path == "/Captcha/Show":
                    seen["cap"] = ("1588", "3493")[seen["n"] % 2]
                    seen["n"] += 1
                    return self._out(caps[(seen["n"] - 1) % 2], "image/png")
                if self.path.startswith("/Invoice/Download?fileGuid=g1"):
                    return self._out(b"%PDF-1.4 easy goc", "application/pdf")
                if self.path == "/Search/DownloadXML?token=XT%2F9%2B":
                    return self._out(b'<?xml version="1.0"?><HDon>easy</HDon>', "text/xml")
                self.send_error(404)

            def do_POST(self):
                form = dict(urllib.parse.parse_qsl(self.rfile.read(int(self.headers["Content-Length"])).decode()))
                seen["posts"].append((self.path, form))
                if "ASP.NET_SessionId=abc" not in (self.headers.get("Cookie") or ""):
                    return self._out(b'<input id="msg" name="msg" value="Phi&#234;n h&#7871;t h&#7841;n">')
                if self.path == "/Search/Search":
                    if form.get("Capcha") != seen["cap"]:
                        return self._out('<input id="msg" name="msg" value="Mã xác thực không đúng">'.encode())
                    if form.get("FKey") != "HIUOVGNMC":
                        return self._out('<input id="msg" name="msg" value="Không tìm thấy hóa đơn">'.encode())
                    inv = H.escape(json.dumps({"str": "<?xml version=\"1.0\"?><html><body>HOA DON</body></html>"}))
                    return self._out(('<input id="InvData" name="InvData" type="hidden" value="%s">'
                                      '<button onclick="downloadPdfAndFileAttachFromAvailableHtml(\'TOK/123+\');">'
                                      '<button onclick="downloadXML(\'XT/9+\')">Tải tệp XML</button>'
                                      % inv).encode())
                if self.path == "/Invoice/DownloadPdfAndFileAttachFromAvailableHtml":
                    import base64
                    ok = form.get("token") == "TOK/123+" and base64.b64decode(form["html"]).decode() == \
                        "<html><body>HOA DON</body></html>"
                    return self._out(json.dumps({"fileGuid": "g1" if ok else "", "fileName": "a.pdf",
                                                 "msg": "" if ok else "sai"}).encode(), "application/json")
                self.send_error(404)
        srv = ThreadingHTTPServer(("127.0.0.1", 0), Easy)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        old = t.EASY_BASE_RE
        t.EASY_BASE_RE = r"http://127\.0\.0\.1:\d+"
        try:
            url = "http://127.0.0.1:%d" % srv.server_port
            self.assertEqual(t.easyinvoice_pdf({"url": url, "code": "HIUOVGNMC"}), b"%PDF-1.4 easy goc")
            self.assertEqual(seen["posts"][0][1]["typeSearch"], "fKeySearch")
            got = t.easyinvoice_files({"url": url, "code": "HIUOVGNMC"})  # chưa biết đường dẫn XML thật → không đoán
            self.assertEqual((got["pdf"], got["xml"]), (b"%PDF-1.4 easy goc", None))
            self.assertIn("tải tay", got["xml_err"])
            with mock.patch.object(t, "EASY_XML_PATHS", ("/Search/DownloadXML?token=",)):
                got = t.easyinvoice_files({"url": url, "code": "HIUOVGNMC"})
            self.assertEqual(got["xml"], b'<?xml version="1.0"?><HDon>easy</HDon>')
            with self.assertRaisesRegex(t.PortalError, "Không tìm thấy"):
                t.easyinvoice_pdf({"url": url, "code": "SAI"})
            self.assertTrue(t.can_fetch_pdf({"ncc_mst": "0105987432", "code": "X"})[0])
        finally:
            t.EASY_BASE_RE = old
            srv.shutdown()
            srv.server_close()

    def test_match_pdf_layouts(self):
        """Các kiểu trình bày PDF của nhiều nhà cung cấp (chữ giả lập theo PDF thật)."""
        def entries(nb, nm, kh, so):
            return {"%s|1|%s|%d" % (nb, kh, s): {"inv": {"nmmst": nm}} for s in (so, so + 1, 1, 2, 99, 100)}
        cases = [
            # chữ PDF bị đảo thứ tự: số trước nhãn, ký hiệu in thành C25TAD1
            ("0309999999:Mã số thuế CONG TY A:Đơn vị bán hàng 00000165:Số C25TAD1:Ký hiệu HOÁ ĐƠN 0101234567:Mã số thuế",
             "0309999999", "0101234567", "C25TAD", 165),
            # số đứng ngay sau ký hiệu, không có nhãn
            ("Mã số thuế 0309999999 Ngày 24 tháng 02 năm 2025 1C25TQP 106 Ký hiệu Số 0099 Mã số thuế 0101234567",
             "0309999999", "0101234567", "C25TQP", 106),
            # MISA: "Ký hiệu: Số: 1C25TLV 00004279" – không đọc nhầm số 1 của ký hiệu
            ("Mã số thuế: 0309999999 Ký hiệu: Số: 1C25TLV 00004279 09 tháng 4 2025 Mã số thuế: 0101234567",
             "0309999999", "0101234567", "C25TLV", 4279),
            # ký hiệu không có số 1 đầu, MST in cách chữ số, nhãn "Số (No):"
            ("Ký hiệu (Serial No): C26TMS Số (No): 268 Mã số thuế (Tax code): 0 3 0 9 9 9 9 9 9 9 Mã số thuế: 0101234567",
             "0309999999", "0101234567", "C26TMS", 268),
        ]
        for text, nb, nm, kh, so in cases:
            hits = t.match_pdf_text(text, {"purchase": entries(nb, nm, kh, so)})
            self.assertEqual(hits, [("purchase", "%s|1|%s|%d" % (nb, kh, so))], text[:40])

    def _vietin_xlsx(self, rows, opening=1000000):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LICH SU GIAO DICH"
        ws.append(["NGÂN HÀNG TMCP CÔNG THƯƠNG VIỆT NAM"])
        ws.append([None, "  VIETINBANK"])
        ws.append([None, "Công ty/Company name:", "CONG TY A"])
        ws.append([None, "Số tài khoản/ Account No.:", "117000000001"])
        tin = sum(r[3] for r in rows)
        tout = sum(r[2] for r in rows)
        ws.append([None, "Số dư đầu kỳ/Opening Balance:", opening])
        ws.append([None, "Số dư cuối kỳ/Closing Balance:", opening + tin - tout])
        ws.append([None, "Tổng giá trị ghi có/ Total credits:", tin, "Tổng giá trị ghi nợ/ Total debits:", tout])
        ws.append([None, "Tổng số giao dịch ghi có/ Total number of credit :", "2", "Tổng số giao dịch ghi nợ/ Total number of debits:", "1"])
        ws.append(["STT/No.", "Ngày hạch toán/Accounting date", "Mô tả giao dịch/ Transaction description", "Nợ/ Debit",
                   "Có / Credit", "Số dư TK/ Account Balance", "Số giao dịch/ Transaction number",
                   "Số tài khoản đối ứng/ Corresponsive account", "Tên tài khoản đối ứng/ Corresponsive name",
                   "MTID/ CITAD", "Mã định danh TK thụ hưởng/ To virtual account", "Ngày phát sinh giao dịch/ Transaction date"])
        for i, (d, desc, out, inn, ref, name) in enumerate(rows, 1):
            ws.append([i, d, desc, out, inn, 0, ref, "123", name, "", "", d])
        ws.append(["Từ ngày 21/03/2025, VietinBank thực hiện cập nhật lại báo cáo sao kê"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_bank_statement(self):
        rows = [("31-01-2026 14:34:15", "CT DI:928K NEW SUN THANH TOAN CHO CONG TY B HOA DON SO 967", 20850300, 0, "928K2611", "CONG TY B"),
                ("31-01-2026 01:57:49", "Tra lai tai khoan DDA", 0, 67629, "2937", ""),
                ("30-01-2026 20:09:36", "CT DEN:2009  TT tien in", 0, 25128080, "928S2611", "CONG TY C")]
        info = t.parse_bank_statement("a.xlsx", self._vietin_xlsx(rows))
        self.assertEqual((info["bank"], info["account"], info["company"]), ("VIETIN", "117000000001", "CONG TY A"))
        self.assertEqual((info["from"], info["to"], len(info["rows"]), info["warnings"]), ("2026-01-30", "2026-01-31", 3, []))
        self.assertEqual(info["rows"][0], {"date": "2026-01-31", "time": "14:34:15", "ref": "928K2611", "in": 0.0,
                                           "out": 20850300.0, "name": "CONG TY B",
                                           "desc": "CT DI:928K NEW SUN THANH TOAN CHO CONG TY B HOA DON SO 967"})
        self.assertEqual(info["rows"][2]["desc"], "CT DEN:2009  TT tien in")   # giữ nguyên như sao kê
        k = t.bank_ktsc_rows(info["rows"], {"ghi_chu": "NIBOT IMPORT"})
        self.assertEqual([(r["TKNO"], r["TKCO"], r["TTVND"], r["TTVND_TT"]) for r in k],
                         [("", "1121", 20850300, 20850300), ("1121", "", 67629, 67629), ("1121", "", 25128080, 25128080)])
        self.assertEqual((k[0]["LCTG"], k[0]["SOCT"], k[0]["TENKH"], k[0]["ID_NGHIEPVU"], k[0]["GHICHU"]),
                         ("CTNH", "928K2611", "CONG TY B", "TIENHANG", "NIBOT IMPORT"))
        # TK đối ứng mặc định + riêng từng dòng, mã đối tượng ngân hàng nằm bên TK ngân hàng
        info["rows"][1]["tkdu"] = "515"
        k = t.bank_ktsc_rows(info["rows"], {"tk_thu": "131", "tk_chi": "331", "ma_dt": "112_VIETTIN"})
        self.assertEqual([(r["TKNO"], r["MADTPNNO"], r["TKCO"], r["MADTPNCO"]) for r in k],
                         [("331", "", "1121", "112_VIETTIN"), ("1121", "112_VIETTIN", "515", ""),
                          ("1121", "112_VIETTIN", "131", "")])
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(t.write_bank_ktsc(info["rows"])))["KTSC"]
        self.assertEqual([c.value for c in ws[1]], list(t.KTSC_BANK_COLS))
        self.assertEqual(ws.max_row, 4)
        # Sao kê tháng sau lặp giao dịch cuối tháng trước → bỏ trùng khi gộp
        feb = t.parse_bank_statement("b.xlsx", self._vietin_xlsx([rows[1], ("02-02-2026", "Phi", 3500, 0, "2990", "")]))
        merged, dups = t.merge_bank_statements([info, feb])
        self.assertEqual((len(merged), dups), (4, 1))

    def test_bank_statement_other_formats(self):
        csv_data = ("Ngày giao dịch;Số tham chiếu;Nội dung;Ghi nợ;Ghi có;Số dư\n"
                    "05/03/2026;FT123;Chuyen tien hang;1.500.000;;9.000.000\n"
                    "06/03/2026;FT124;Nhan tien;;2.000.000,50;11.000.000\n").encode("utf-8")
        info = t.parse_bank_statement("vcb.csv", csv_data)
        self.assertEqual([(r["date"], r["ref"], r["in"], r["out"]) for r in info["rows"]],
                         [("2026-03-05", "FT123", 0.0, 1500000.0), ("2026-03-06", "FT124", 2000000.5, 0.0)])
        html_data = ("<html><body><p>TECHCOMBANK</p><table><tr><th>Ngày</th><th>Diễn giải</th><th>Số tiền</th>"
                     "<th>Tên người chuyển</th></tr><tr><td>2026-03-07</td><td>Thu tien</td><td>3,000,000</td><td>NGUYEN A</td></tr>"
                     "<tr><td>2026-03-08</td><td>Tra tien</td><td>-500,000</td><td></td></tr></table></body></html>").encode()
        info = t.parse_bank_statement("tcb.xls", html_data)
        self.assertEqual(info["bank"], "TCB")
        self.assertEqual([(r["in"], r["out"], r["name"]) for r in info["rows"]], [(3000000.0, 0.0, "NGUYEN A"), (0.0, 500000.0, "")])
        with self.assertRaisesRegex(ValueError, "không tìm thấy bảng giao dịch"):
            t.parse_bank_statement("x.csv", b"a;b;c\n1;2;3\n")
        self.assertEqual(t._bank_amount("(1.234.567)"), -1234567.0)
        self.assertEqual(t._bank_amount("1,234.5"), 1234.5)

    def test_update_invoices_bulk(self):
        self.app.store.import_text("0309999999\tA\tpw1\n")
        self.run_job(["0309999999"], kinds=("purchase",))
        keys = list(t._load_json(t.index_file(self.tmp, "0309999999"), {})["purchase"])
        self.assertEqual(t.update_invoices(self.tmp, "0309999999", "purchase", keys + ["khong|co"], duyet="Đã duyệt", dv=True),
                         len(keys))
        rows = t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})
        self.assertTrue(rows and all(r["duyet"] == "Đã duyệt" and r["dv"] for r in rows))
        self.assertEqual(len(t.query_invoices(self.tmp, "0309999999", {"kind": "purchase_dv"})), len(rows))

    @staticmethod
    def _pdf(pages):
        """PDF tối giản: mỗi trang là danh sách (x, y, chữ ASCII)."""
        objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
        kids = []
        for frags in pages:
            body = " ".join("BT 1 0 0 1 %s %s Tm /F1 7 Tf (%s) Tj 0 g ET" % (x, y, t.replace("(", "\\(").replace(")", "\\)"))
                            for x, y, t in frags)  # mỗi ô một khối chữ như sổ phụ thật
            objs.append("<< /Length %d >>\nstream\n%s\nendstream" % (len(body), body))
            objs.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] /Contents %d 0 R "
                        "/Resources << /Font << /F1 3 0 R >> >> >>" % len(objs))
            kids.append(len(objs))
        objs[1] = "<< /Type /Pages /Kids [%s] /Count %d >>" % (" ".join("%d 0 R" % k for k in kids), len(kids))
        out, offs = "%PDF-1.4\n", []
        for i, o in enumerate(objs, 1):
            offs.append(len(out))
            out += "%d 0 obj\n%s\nendobj\n" % (i, o)
        x = len(out)
        out += "xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + "".join("%010d 00000 n \n" % o for o in offs)
        out += "trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)
        return out.encode("latin-1")

    def test_bank_statement_pdf(self):
        """Sổ phụ PDF (bố cục như MB): cột theo toạ độ, số bút toán 2 dòng, tổng in 2 dòng, nội dung cắt 35 ký tự."""
        hdr = [(23, 547, "Ngay giao dich"), (95, 547, "Ngay hach toan"), (170, 552, "So but toan"), (186, 531, "No"),
               (221, 547, "Phat sinh no"), (283, 547, "Phat sinh co"), (428, 547, "Noi dung"),
               (560, 547, "Don vi thu huong/Don vi chuyen"), (706, 547, "Tai khoan")]
        top = [(250, 592, "SO PHU CHI TIET NGAN HANG TMCP QUAN DOI"), (20, 568, "Tai khoan/Account No: 2018000001"),
               (20, 580 - 24, "Ten khach hang/Customer name: CONG TY A"), (40, 580 - 36, "So du dau ky/ Opening Balance: 2,000,000 VND")]

        def row(y, d, ref1, ref2, debit, credit, desc, name):
            return [(26, y + 5, d + " 10:"), (41, y - 4, "01:30"), (106, y, d), (172, y + 5, ref1), (172, y - 4, ref2),
                    (269, y, debit), (300, y, credit)] + [(344, y + 5 - 9 * i, line) for i, line in enumerate(desc)] + \
                [(559, y, name), (697, y, "0601607")]
        p1 = top + hdr + row(500, "03/09/2026", "FT26246090", "960093", "0", "1,090,000",
                             ["IBFT Strawberry Tao chuyen khoan nh  anh qua Zalo", "W2LQ5EU7/722239"], "KHA BINH LUONG") + \
            row(470, "04/09/2026", "FT26247479", "933602", "2,340,000", "0", ["MBCT thanh toan"], "CT TNHH TM FNB") + \
            [(20, 30, "Chung tu nay duoc xuat tu dong tu he thong")]
        p2 = hdr + row(500, "05/09/2026", "FT26248550", "198084", "0", "500", ["Lai"], "") + \
            [(60, 440, "Tong phat sinh trong ky/Total"), (228, 446, "2,340,00"), (262, 434, "0"),
             (292, 446, "1,090,50"), (326, 434, "0"), (20, 420, "So du cuoi ky/ Closing Balance: 750,500 VND")]
        info = t.parse_bank_statement("sp.pdf", self._pdf([p1, p2]))
        self.assertEqual((info["bank"], info["account"], info["company"]), ("MB", "2018000001", "CONG TY A"))
        self.assertEqual((info["opening"], info["closing"], info["total_out"], info["total_in"]),
                         (2000000.0, 750500.0, 2340000.0, 1090500.0))
        self.assertEqual(info["warnings"], [])
        self.assertEqual([(r["date"], r["ref"], r["in"], r["out"], r["name"]) for r in info["rows"]],
                         [("2026-09-03", "FT26246090960093", 1090000.0, 0.0, "KHA BINH LUONG"),
                          ("2026-09-04", "FT26247479933602", 0.0, 2340000.0, "CT TNHH TM FNB"),
                          ("2026-09-05", "FT26248550198084", 500.0, 0.0, "")])
        self.assertEqual(info["rows"][0]["desc"], "IBFT Strawberry Tao chuyen khoan nhanh qua Zalo W2LQ5EU7/722239")
        with self.assertRaisesRegex(ValueError, "scan"):
            t.parse_bank_statement("scan.pdf", self._pdf([[(10, 10, "x")]]))

    @unittest.skipUnless(t.find_browser(), "máy không có Chrome/Edge/Chromium")
    def test_tax_pdf_from_portal_html(self):
        """PDF thuế = bản HTML của cổng in qua trình duyệt chạy ngầm (chạy cả details.js đi kèm)."""
        self.app.store.import_text("0309999999\tA\tpw1\n")
        self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        row = t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})[0]
        base = os.path.join(self.tmp, "0309999999")
        goc = os.path.join(base, os.path.splitext(row["xml"])[0] + "_goc")
        self.assertEqual(sorted(os.listdir(goc)), ["details.js", "invoice.html", "invoice.xml"])
        rel = t.tax_pdf(self.tmp, "0309999999", "purchase", row["key"])
        import pypdf
        text = pypdf.PdfReader(os.path.join(base, rel)).pages[0].extract_text()
        self.assertIn("HÓA ĐƠN CỦA CỔNG THUẾ", text)
        self.assertIn("Số hóa đơn: 99", text)
        e = t._load_json(t.index_file(self.tmp, "0309999999"), {})["purchase"][row["key"]]
        self.assertEqual(e["cqt_src"], "html")
        # Bản cũ chỉ lưu .html (thiếu details.js) → lần đồng bộ sau tải lại gói đầy đủ
        import shutil
        shutil.rmtree(goc)
        n = sum(1 for p_, _ in FakePortal.calls if p_.endswith("export-xml"))
        self.run_job(["0309999999"], kinds=("purchase",), start=date(2026, 10, 1), end=date(2026, 10, 31))
        self.assertEqual(sum(1 for p_, _ in FakePortal.calls if p_.endswith("export-xml")), n + 1)
        self.assertTrue(os.path.isdir(goc))

    def test_download_everything(self):
        """Đồng bộ kèm PDF gốc (nhà cung cấp lỗi liên tiếp thì tạm bỏ qua) và kết xuất Tất cả: XML + PDF + Excel."""
        import unittest.mock as mock
        self.app.store.import_text("0309999999\tA\tpw1\n")
        calls = []

        def broken(t_, **kw):
            calls.append(t_["code"])
            raise t.PortalError("MISA chặn")
        job = t.Job()
        job.running = True
        job.want_pdf = True
        helper = AutoAnswer(job)
        helper.start()
        with mock.patch.dict(t.ORIGINAL_FETCHERS, {"0101243150": broken}):
            t.run_batch(job, self.app.store, self.app.solver, ["0309999999"], ["purchase"], date(2026, 9, 1),
                        date(2026, 10, 31), True, True, self.tmp, False, self.app.client_factory)
        helper.join(1)
        self.assertEqual(len(calls), 3, "lỗi 3 lần liên tiếp thì không thử tiếp các hoá đơn MISA còn lại")
        log = "\n".join(job.log)
        self.assertIn("bỏ qua PDF gốc của nhà cung cấp này", log)
        self.assertTrue(any("HĐ gốc: bỏ qua" in r["ketqua"] for r in job.rows))
        with mock.patch.object(t, "find_browser", return_value=None):
            zp = t.export_invoices(self.tmp, "0309999999", "A", {"kind": "purchase"}, "all")
        names = zipfile.ZipFile(zp).namelist()
        self.assertEqual(sum(n.startswith("XML/") for n in names), 61)
        self.assertEqual(sum(n.startswith("PDF/") for n in names), 61)
        self.assertEqual(sum(n.endswith(".xlsx") for n in names), 1)

    def test_import_original_xml(self):
        """Gắn XML gốc tải tay từ trang tra cứu (kể cả trong ZIP): khớp theo MST bán, ký hiệu, số hoá đơn."""
        self.app.store.import_text("0309999999\tA\tpw1\n")
        job = t.Job()
        job.running = True
        helper = AutoAnswer(job)
        helper.start()
        t.run_batch(job, self.app.store, self.app.solver, ["0309999999"], ["purchase"], date(2026, 10, 1),
                    date(2026, 10, 31), True, True, self.tmp, False, self.app.client_factory)
        helper.join(1)
        r = t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"})[0]
        nb, mau, kh, so = r["key"].split("|")
        xml = ('<?xml version="1.0" encoding="UTF-8"?><HDon><DLHDon><TTChung><KHMSHDon>%s</KHMSHDon><KHHDon>%s</KHHDon>'
               '<SHDon>%s</SHDon></TTChung><NDHDon><NBan><Ten>X</Ten><MST>%s</MST></NBan></NDHDon></DLHDon></HDon>'
               % (mau, kh, so, nb)).encode()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("hoa-don.xml", xml)
        res = t.import_pdfs(self.tmp, "0309999999", [("x.zip", buf.getvalue()),
                                                    ("sai.xml", xml.replace(b"<SHDon>", b"<SHDon>9"))])
        self.assertTrue(res[0]["ok"], res)
        self.assertFalse(res[1]["ok"])
        r = [x for x in t.query_invoices(self.tmp, "0309999999", {"kind": "purchase"}) if x["key"] == r["key"]][0]
        self.assertTrue(r["xml_goc"])


if __name__ == "__main__":
    unittest.main()
