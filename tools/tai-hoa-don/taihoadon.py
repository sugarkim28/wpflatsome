#!/usr/bin/env python3
"""Phần mềm tải hoá đơn điện tử từ cổng hoadondientu.gdt.gov.vn – quản lý nhiều doanh nghiệp.

Chạy: python taihoadon.py  → trình duyệt mở http://127.0.0.1:8765
- Danh sách doanh nghiệp (MST, tên, mật khẩu cổng HĐĐT mã hoá trên máy, ghi chú, bật/tắt đồng bộ vào/ra, ẩn).
- Xử lý hàng loạt: đăng nhập từng DN, tải hoá đơn mua vào / bán ra / máy tính tiền theo khoảng ngày
  (tự chia theo tháng), XML gốc, bảng kê Excel + chi tiết hàng hoá, phát hiện HĐ đổi trạng thái (huỷ, thay thế...).
- Captcha tự học: vài lần đầu người dùng gõ, sau đó phần mềm tự giải.

Chỉ dùng thư viện chuẩn Python 3.8+. openpyxl là tuỳ chọn (pip install openpyxl) để xuất .xlsx.
"""

import argparse
import calendar
import http.cookiejar
import csv
import io
import json
import os
import re
import secrets
import ssl
import sys
import threading
import time
import uuid
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
import zipfile
from datetime import date, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

__version__ = "3.5.0"

BASE_URL = os.environ.get("HDDT_BASE_URL", "https://hoadondientu.gdt.gov.vn/api")
PAGE_SIZE = 50
XML_WORKERS = 4  # số file XML tải cùng lúc
MST_STATUS = {"00": "Đang hoạt động"}  # mã khác: hiện nguyên mã
NO_XML_MSG = "khongtontaihosogoc"  # cổng báo HĐ không có XML gốc (HĐ không mã), dạng đã bỏ dấu
# Cổng có tường lửa nhận dạng hành vi: đăng nhập chỉ mang header tối giản như trình duyệt gọi qua proxy
# của trang; tra cứu/tải XML mang header của trang tra cứu (Action, End-Point). Gửi sai bộ → 403
# "Hệ thống phát hiện hành vi không hợp lệ".
UA_LOGIN = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
UA_QUERY = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/106.0.0.0 Safari/537.36"

# Trạng thái xử lý (ttxly) của CQT dùng khi tra hoá đơn mua vào.
TTXLY_LABELS = {
    5: "Đã cấp mã hoá đơn",
    6: "CQT đã nhận, không mã",
    8: "CQT đã nhận, máy tính tiền",
}
TTHAI_LABELS = {
    1: "Hoá đơn mới",
    2: "Hoá đơn thay thế",
    3: "Hoá đơn điều chỉnh",
    4: "Đã bị thay thế",
    5: "Đã bị điều chỉnh",
    6: "Đã bị huỷ",
}

# (cột, khoá dữ liệu) cho bảng kê.
COLUMNS = [
    ("STT", "_stt"),
    ("Ngày lập", "_ngay"),
    ("Ký hiệu mẫu số", "khmshdon"),
    ("Ký hiệu hoá đơn", "khhdon"),
    ("Số hoá đơn", "shdon"),
    ("MST người bán", "nbmst"),
    ("Tên người bán", "nbten"),
    ("MST người mua", "nmmst"),
    ("Tên người mua", "nmten"),
    ("Tiền chưa thuế", "tgtcthue"),
    ("Tiền thuế", "tgtthue"),
    ("Tổng thanh toán", "tgtttbso"),
    ("Tiền tệ", "dvtte"),
    ("Trạng thái hoá đơn", "_tthai"),
    ("Kết quả xử lý", "_ttxly"),
    ("Nguồn", "_nguon"),
    ("Mã CQT", "mhdon"),
    ("File XML", "_xml"),
]
MONEY_KEYS = {"tgtcthue", "tgtthue", "tgtttbso"}


class PortalError(Exception):
    """Lỗi trả về từ cổng hoá đơn điện tử (thông báo hiển thị cho người dùng)."""


# ---------------------------------------------------------------------------
# Tiện ích ngày tháng
# ---------------------------------------------------------------------------

def parse_date(text):
    """Nhận 'dd/mm/yyyy' hoặc 'yyyy-mm-dd'."""
    text = (text or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError("Ngày không hợp lệ: %r (dùng dd/mm/yyyy)" % text)


def month_ranges(start, end):
    """Chia [start, end] thành các đoạn không vượt quá một tháng dương lịch."""
    if end < start:
        raise ValueError("Ngày kết thúc phải sau ngày bắt đầu")
    ranges = []
    cur = start
    while cur <= end:
        last = date(cur.year, cur.month, calendar.monthrange(cur.year, cur.month)[1])
        stop = min(last, end)
        ranges.append((cur, stop))
        cur = stop + timedelta(days=1)
    return ranges


def search_query(start, end, ttxly=None):
    q = "tdlap=ge=%sT00:00:00;tdlap=le=%sT23:59:59" % (start.strftime("%d/%m/%Y"), end.strftime("%d/%m/%Y"))
    if ttxly is not None:
        q += ";ttxly==%d" % ttxly
    return q


def invoice_date(inv):
    """Ngày lập theo giờ Việt Nam (cổng trả giờ UTC, vd 2026-10-01T17:00:00Z = 02/10/2026)."""
    raw = str(inv.get("tdlap") or "")
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2}))?", raw)
    if m:
        d = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4) or 0), int(m.group(5) or 0))
        if m.group(4) and re.search(r"(Z|[+-]00:?00)$", raw):
            d += timedelta(hours=7)
        return d.date()
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", raw)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def safe_name(text):
    return re.sub(r"[^0-9A-Za-z._-]+", "_", str(text)).strip("_") or "x"


# ---------------------------------------------------------------------------
# Client cổng hoá đơn điện tử
# ---------------------------------------------------------------------------

class HoaDonClient:
    def __init__(self, base_url=BASE_URL, timeout=60, verify_ssl=True, delay=0.15):
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout
        self.delay = delay
        self.token = None
        self.username = None
        # Trang chủ của cổng = base_url bỏ đuôi /api
        self.site_url = self.base_url[:-4] if self.base_url.endswith("/api") else self.base_url
        ctx = None if verify_ssl else ssl._create_unverified_context()
        https = urllib.request.HTTPSHandler(context=ctx)
        # Đăng nhập dùng chung cookie với trang chủ + captcha; tra cứu không gửi cookie.
        self._jar = http.cookiejar.CookieJar()
        self._login_opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self._jar), https)
        self._query_opener = urllib.request.build_opener(https)
        self._visited = False

    # -- HTTP ---------------------------------------------------------------
    def _headers(self, profile, action=None, json_body=False):
        if profile == "portal":
            h = {"User-Agent": UA_LOGIN, "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                 "Accept-Language": "vi"}
        elif profile == "login":
            h = {"User-Agent": UA_LOGIN, "Accept": "application/json, text/plain, */*"}
            if json_body:
                h["Content-Type"] = "application/json"
        else:
            h = {"Authorization": "Bearer " + (self.token or ""), "Accept": "application/json, text/plain, */*",
                 "Accept-Language": "vi", "End-Point": "/tra-cuu/tra-cuu-hoa-don", "Origin": self.site_url,
                 "Referer": self.site_url + "/", "User-Agent": UA_QUERY,
                 "Action": urllib.parse.quote(action or "Tìm kiếm", safe="()")}
        h["request-id"] = str(uuid.uuid4())
        return h

    def _request(self, method, path, params=None, body=None, raw=False, retries=3, profile="query", action=None,
                 url=None):
        url = url or self.base_url + path
        if params:
            url += "?" + urllib.parse.urlencode(params, safe=":,;=/")
        data = json.dumps(body).encode("utf-8") if body is not None else None
        opener = self._query_opener if profile == "query" else self._login_opener

        for attempt in range(retries + 1):
            headers = self._headers(profile, action, body is not None)
            req = urllib.request.Request(url, data=data, headers=headers, method=method)
            try:
                with opener.open(req, timeout=self.timeout) as resp:
                    content = resp.read()
                break
            except urllib.error.HTTPError as e:
                content = e.read()
                # Chỉ thử lại khi cổng quá tải; cổng đã trả thông báo cụ thể (vd "Không tồn tại hồ sơ gốc") thì không.
                if e.code in (429, 500, 502, 503, 504) and attempt < retries and \
                        NO_XML_MSG not in _plain(self._error_message(content) or ""):
                    time.sleep(2 ** attempt * 2)
                    continue
                if e.code == 401:
                    self.token = None
                    raise PortalError("Phiên đăng nhập đã hết hạn, vui lòng đăng nhập lại.")
                msg = self._error_message(content) or "Cổng hoá đơn trả lỗi HTTP %d" % e.code
                if "hành vi không hợp lệ" in msg:
                    msg += " (tường lửa của cổng chặn; chờ vài phút rồi thử lại, nếu vẫn bị hãy báo để cập nhật phần mềm)"
                raise PortalError(msg)
            except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
                if attempt < retries:
                    time.sleep(2 ** attempt * 2)
                    continue
                reason = getattr(e, "reason", e)
                raise PortalError("Không kết nối được cổng hoá đơn điện tử: %s" % reason)

        if raw:
            return content
        try:
            return json.loads(content.decode("utf-8")) if content else {}
        except ValueError:
            raise PortalError("Cổng hoá đơn trả dữ liệu không đọc được.")

    @staticmethod
    def _error_message(content):
        try:
            data = json.loads(content.decode("utf-8"))
        except (ValueError, UnicodeDecodeError, AttributeError):
            return None
        if isinstance(data, dict):
            return data.get("message") or data.get("error_description") or data.get("error")
        return None

    # -- Đăng nhập ----------------------------------------------------------
    def _visit_portal(self):
        """Mở trang chủ như trình duyệt để nhận cookie trước khi lấy captcha/đăng nhập."""
        if self._visited:
            return
        self._visited = True
        try:
            self._request("GET", "", raw=True, retries=1, profile="portal", url=self.site_url + "/")
        except PortalError:
            pass

    def lookup_company(self, mst):
        """Tra thông tin người nộp thuế theo MST (API công khai của cổng, không cần đăng nhập)."""
        mst = re.sub(r"\s+", "", str(mst or ""))
        if not re.fullmatch(r"\d{10}(-\d{3})?", mst):
            raise ValueError("MST không hợp lệ")
        self._visit_portal()
        d = self._request("GET", "/category/public/dsdkts/%s/manager" % mst, profile="login", retries=1)
        if not isinstance(d, dict) or not d.get("tennnt"):
            msg = d.get("message") if isinstance(d, dict) else None
            raise PortalError(msg or "Không tìm thấy MST %s" % mst)
        dia_chi = ", ".join(x for x in (d.get("dctsdchi"), d.get("dctsxaten"), d.get("dctshuyenten"),
                                        d.get("dctstinhten")) if x)
        tthai = str(d.get("tthai") or "")
        return {"mst": mst, "ten": d.get("tennnt", "").strip(), "dia_chi": dia_chi, "cqt": d.get("tencqt", ""),
                "tthai": tthai, "tthai_text": MST_STATUS.get(tthai, "mã " + tthai if tthai else "")}

    def get_captcha(self):
        """Trả về {'key': ..., 'content': '<svg ...>'}."""
        self._visit_portal()
        data = self._request("GET", "/captcha", profile="login")
        if not data.get("key") or not data.get("content"):
            raise PortalError("Không lấy được captcha.")
        return {"key": data["key"], "content": data["content"]}

    def login(self, username, password, captcha_value, captcha_key):
        self.token = None
        data = self._request(
            "POST",
            "/security-taxpayer/authenticate",
            body={"username": username, "password": password, "ckey": captcha_key, "cvalue": captcha_value},
            retries=0, profile="login",
        )
        token = data.get("token")
        if not token:
            raise PortalError(data.get("message") or "Đăng nhập không thành công.")
        self.token = token
        self.username = username
        self.login_at = time.time()
        return token

    def token_valid(self, margin=120):
        """Token còn hạn? Đọc 'exp' trong JWT; không đọc được thì coi hạn 30 phút kể từ lúc đăng nhập."""
        if not self.token:
            return False
        import base64
        try:
            part = self.token.split(".")[1]
            exp = json.loads(base64.urlsafe_b64decode(part + "=" * (-len(part) % 4))).get("exp")
        except (IndexError, ValueError, AttributeError):
            exp = None
        if not exp:
            exp = getattr(self, "login_at", 0) + 1800
        return time.time() < exp - margin

    # -- Tra cứu -------------------------------------------------------------
    def _sources(self, kind, include_mtt):
        """Danh sách (đường dẫn gốc, ttxly, nhãn nguồn) cần tra cho một loại hoá đơn."""
        if kind == "purchase":
            out = [("/query", 5, "Có mã / không mã"), ("/query", 6, "Có mã / không mã")]
            if include_mtt:
                out.append(("/sco-query", 8, "Máy tính tiền"))
        else:
            out = [("/query", None, "Có mã / không mã")]
            if include_mtt:
                out.append(("/sco-query", None, "Máy tính tiền"))
        return out

    def list_invoices(self, kind, start, end, include_mtt=True, progress=None):
        """kind: 'purchase' (mua vào) hoặc 'sold' (bán ra). Trả về list dict hoá đơn."""
        if kind not in ("purchase", "sold"):
            raise ValueError("kind phải là 'purchase' hoặc 'sold'")
        results, seen = [], set()
        for prefix, ttxly, label in self._sources(kind, include_mtt):
            for a, b in month_ranges(start, end):
                if progress:
                    progress("Tra %s %s → %s (%s)" % (
                        "mua vào" if kind == "purchase" else "bán ra",
                        a.strftime("%d/%m/%Y"), b.strftime("%d/%m/%Y"), label))
                for inv in self._fetch_range(prefix, kind, a, b, ttxly, progress):
                    key = (inv.get("nbmst"), inv.get("khmshdon"), inv.get("khhdon"), inv.get("shdon"))
                    if key in seen:
                        continue
                    seen.add(key)
                    inv["_prefix"] = prefix
                    inv["_nguon"] = label
                    inv["_kind"] = kind
                    results.append(inv)
                time.sleep(self.delay)
        results.sort(key=lambda i: (str(i.get("tdlap") or ""), str(i.get("khhdon") or ""), _num(i.get("shdon"))))
        return results

    def _fetch_range(self, prefix, kind, a, b, ttxly, progress=None):
        """Tra một khoảng ngày (có lật trang). Cổng quá tải (timeout) thì chia nhỏ theo tuần rồi tra lại."""
        try:
            out, state = [], None
            while True:
                params = {"sort": "tdlap:desc", "size": PAGE_SIZE, "search": search_query(a, b, ttxly)}
                if state:
                    params["state"] = state
                data = self._request("GET", "%s/invoices/%s" % (prefix, kind), params=params, retries=2)
                page = data.get("datas") or []
                out.extend(page)
                state = data.get("state")
                if not page or not state or len(page) < PAGE_SIZE:
                    return out
                time.sleep(self.delay)
        except PortalError as e:
            busy = any(w in str(e).lower() for w in ("timeout", "quá tải", "http 5"))
            if not busy or (b - a).days < 1:
                if busy:
                    raise PortalError("Cổng thuế đang quá tải, không trả được dữ liệu ngày %s – %s. Vui lòng đồng bộ "
                                      "lại sau ít phút. (%s)" % (a.strftime("%d/%m/%Y"), b.strftime("%d/%m/%Y"), e))
                raise
            step = 7 if (b - a).days > 7 else 1
            if progress:
                progress("Cổng thuế quá tải, chia nhỏ %s – %s theo %s để tra lại" % (
                    a.strftime("%d/%m/%Y"), b.strftime("%d/%m/%Y"), "tuần" if step == 7 else "ngày"))
            out, cur = [], a
            while cur <= b:
                stop = min(cur + timedelta(days=step - 1), b)
                out.extend(self._fetch_range(prefix, kind, cur, stop, ttxly, progress))
                cur = stop + timedelta(days=1)
                time.sleep(self.delay)
            return out

    def _invoice_params(self, inv):
        return {"nbmst": inv.get("nbmst"), "khhdon": inv.get("khhdon"),
                "shdon": inv.get("shdon"), "khmshdon": inv.get("khmshdon")}

    def get_detail(self, inv):
        return self._request("GET", "%s/invoices/detail" % inv.get("_prefix", "/query"),
                             params=self._invoice_params(inv), action="Xem hóa đơn (%s)" % self._kind_label(inv))

    def export_xml(self, inv):
        """Trả về bytes file zip (invoice.xml, details.js, ...) do cổng cung cấp."""
        return self._request("GET", "%s/invoices/export-xml" % inv.get("_prefix", "/query"),
                             params=self._invoice_params(inv), raw=True,
                             action="Xuất xml (%s)" % self._kind_label(inv))

    @staticmethod
    def _kind_label(inv):
        return "hóa đơn bán ra" if inv.get("_kind") == "sold" else "hóa đơn mua vào"


def _num(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def _money(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ---------------------------------------------------------------------------
# Lưu file
def invoice_key(inv):
    return "%s|%s|%s|%s" % (inv.get("nbmst"), inv.get("khmshdon"), inv.get("khhdon"), inv.get("shdon"))


def invoice_basename(inv):
    d = invoice_date(inv)
    return "_".join([
        d.strftime("%Y%m%d") if d else "nodate",
        safe_name(inv.get("nbmst")),
        safe_name("%s%s" % (inv.get("khmshdon") or "", inv.get("khhdon") or "")),
        safe_name(inv.get("shdon")),
    ])


def save_xml(zip_bytes, folder, basename):
    """Giải nén XML (và HTML nếu có) từ zip của cổng. Trả về đường dẫn file XML."""
    os.makedirs(folder, exist_ok=True)
    try:
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    except zipfile.BadZipFile:
        # Một số trường hợp cổng trả thẳng XML.
        if zip_bytes.lstrip()[:1] == b"<":
            path = os.path.join(folder, basename + ".xml")
            with open(path, "wb") as f:
                f.write(zip_bytes)
            return path
        raise PortalError("File XML tải về không hợp lệ.")
    xml_path = None
    with zf:
        for name in zf.namelist():
            ext = os.path.splitext(name)[1].lower()
            if ext not in (".xml", ".html", ".htm", ".pdf"):
                continue
            target = os.path.join(folder, basename + ext)
            if ext == ".xml" and xml_path:
                target = os.path.join(folder, "%s_%s" % (basename, safe_name(os.path.basename(name))))
            with open(target, "wb") as f:
                f.write(zf.read(name))
            if ext == ".xml" and not xml_path:
                xml_path = target
    if not xml_path:
        raise PortalError("File tải về không chứa XML hoá đơn.")
    return xml_path


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def table_rows(invoices):
    rows = []
    for idx, inv in enumerate(invoices, 1):
        d = invoice_date(inv)
        extra = {
            "_stt": idx,
            "_ngay": d.strftime("%d/%m/%Y") if d else inv.get("tdlap"),
            "_tthai": TTHAI_LABELS.get(_num(inv.get("tthai")), inv.get("tthai")),
            "_ttxly": TTXLY_LABELS.get(_num(inv.get("ttxly")), inv.get("ttxly")),
            "_nguon": inv.get("_nguon"),
            "_xml": inv.get("_xml", ""),
        }
        row = []
        for _, key in COLUMNS:
            val = extra[key] if key in extra else inv.get(key)
            if key in MONEY_KEYS:
                val = _money(val)
            elif key != "_stt":
                val = "" if val is None else str(val)  # số HĐ, MST giữ dạng chữ (không thành 1.234)
            row.append("" if val is None else val)
        rows.append(row)
    return rows


ITEM_HEAD = ["Ngày lập", "Ký hiệu", "Số HĐ", "MST người bán", "Tên người bán", "MST người mua", "Tên người mua"]
CHANGE_HEAD = ["Ngày lập", "Ký hiệu", "Số HĐ", "MST người bán", "Tên người bán", "Trạng thái cũ", "Trạng thái mới"]


def _inv_head(inv):
    d = invoice_date(inv)
    return [d.strftime("%d/%m/%Y") if d else "", "%s%s" % (inv.get("khmshdon") or "", inv.get("khhdon") or ""),
            str(inv.get("shdon") or ""), str(inv.get("nbmst") or ""), inv.get("nbten") or "",
            str(inv.get("nmmst") or ""), inv.get("nmten") or ""]


def write_reports(invoices, folder, basename, csv_only=True):
    """Ghi bảng kê dạng .csv (mở bằng Excel, đúng tiếng Việt). Trả về list đường dẫn."""
    os.makedirs(folder, exist_ok=True)
    csv_path = os.path.join(folder, basename + ".csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in COLUMNS])
        w.writerows(table_rows(invoices))
    return [csv_path]


# ---------------------------------------------------------------------------
# Xuất Excel theo mẫu Nibot: HoaDon_TongQuat, Smart_KTSC (import Smart Pro), BangKe_MuaVao,
# BangKe_MuaVao_KCT_HDBH, BangKe_HoanThue. Số liệu chi tiết đọc từ XML (chuẩn TT78).
# ---------------------------------------------------------------------------

LOAI_HD = {1: "V", 2: "B"}
TTHAI_NIBOT = {1: "HĐ Mới", 2: "HĐ Thay thế", 3: "HĐ Điều chỉnh", 4: "HĐ Đã bị thay thế",
               5: "HĐ Đã bị điều chỉnh", 6: "HĐ Đã bị hủy"}
TTXLY_NIBOT = {0: "TCT đã nhận", 1: "Đang k.tra", 2: "CQT t.chối HĐ", 3: "HĐ đủ đ.kiện", 4: "HĐ k đủ đ.kiện",
               5: "Đã cấp MST", 6: "TCT k nhận mã", 7: "Đã k.tra định kỳ", 8: "HĐ có mã từ máy tính tiền"}
VAT_RATES = {"0", "5", "8", "10"}
SMART_HEAD = ["NIBOT_GHICHU", "LCTG", "SR_HD", "SOCT", "NGAY_KY", "NGAYCT", "SO_HD", "NGAY_HD", "DIENGIAI",
              "HTTT", "TKNO", "MADTPNNO", "TKCO", "MADTPNCO", "MADMNO", "MADMCO", "TENDM", "MATHANG",
              "LUONG_CTU", "DONVI_CTU", "DONVI", "LUONG", "DGUSD", "TTUSD", "TYGIA", "DGVND", "TTVND", "PT_CK",
              "CHIETKHAU", "HDVAT", "TKTHUE", "TS_GTGT", "THUEUSD", "THUEVND", "TTUSD_TT", "TTVND_TT", "MAKH",
              "TENKH", "KHACHHANG", "DIACHI_NGD", "MS_DN", "DIACHI", "TK_XUATKHO", "ID_NGHIEPVU", "GHICHU", "GUID"]


def _float(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def norm_rate(v):
    """'8%' → '8'; 'KCT', 'KKKNT', 'KHAC:5.26%' giữ nguyên dạng chữ hoa."""
    s = str(v or "").strip().upper().replace(" ", "")
    if s.endswith("%"):
        s = s[:-1]
    try:
        f = float(s)
        return ("%g" % f)
    except ValueError:
        return s


def _to_dt(v):
    """Chuỗi ISO/dd-mm-yyyy → datetime (bỏ múi giờ)."""
    if not v:
        return None
    s = str(v)
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2}):(\d{2}))?", s)
    if m:
        y, mo, d, hh, mm, ss = m.groups()
        dt = datetime(int(y), int(mo), int(d), int(hh or 0), int(mm or 0), int(ss or 0))
        if s.endswith("Z") or "+0000" in s or "+00:00" in s:
            dt += timedelta(hours=7)  # giờ Việt Nam
        return dt
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", s)
    if m:
        return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


# Nhà cung cấp giải pháp HĐĐT (MSTTCGP trong XML) → (tên, trang tra cứu hoá đơn gốc).
PROVIDERS = {
    "0101243150": ("MISA meInvoice", "https://www.meinvoice.vn/tra-cuu/?sc="),  # ?sc=<mã> mở thẳng hoá đơn
    "0105987432": ("EasyInvoice (SoftDreams)", "http://{nbmst}hd.easyinvoice.com.vn"),  # trang riêng từng người bán
    "0100109106": ("Viettel S-Invoice", "https://vinvoice.viettel.vn/utilities/invoice-search"),
    "0100684378": ("VNPT Invoice", ""),
    "0101360697": ("BKAV eHoadon", "https://van.ehoadon.vn/Lookup?InvoiceGUID="),
    "0100686209": ("MobiFone Invoice", "http://tracuuhoadon.mobifoneinvoice.vn/trang-chu"),
    "0101300842": ("Thái Sơn E-invoice", "https://einvoice.vn/tra-cuu"),
    "0100727825": ("FAST e-Invoice", "https://einvoice.fast.com.vn/"),
}
# Tên trường "mã tra cứu" trong phần thông tin khác (TTKhac) của XML, đã bỏ dấu, chữ thường, bỏ khoảng trắng.
LOOKUP_FIELDS = ("transactionid", "masobimat", "reservationcode", "matracuu", "fkey", "keysearch", "matc",
                 "searchkey", "magiaodich")


def _plain(text):
    import unicodedata
    t = unicodedata.normalize("NFD", str(text or "")).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in t if unicodedata.category(c) != "Mn").lower().replace(" ", "").replace("_", "")


def lookup_info(info):
    """{'ncc', 'url', 'field', 'code'} để tra cứu hoá đơn gốc trên trang của nhà cung cấp."""
    mst = info.get("msttcgp", "")
    name, url = PROVIDERS.get(mst.split("-")[0], ("", ""))
    field = code = ""
    for want in LOOKUP_FIELDS:
        for k, v in info.get("ttkhac", {}).items():
            if _plain(k) == want and v:
                field, code = k, v
                break
        if code:
            break
    if mst.startswith("0101360697") and not code:
        field, code = "InvoiceGUID", info.get("dlhdon_id", "")
    if url.endswith("=") and code:
        url += urllib.parse.quote(code)
    elif url.endswith("=") or "{nbmst}" in url:
        url = url.split("?")[0] if url.endswith("=") else url
    url = url.replace("{nbmst}", (info.get("nb") or {}).get("MST", "").split("-")[0])
    return {"ncc": name or (("MST " + mst) if mst else ""), "url": url, "field": field, "code": code,
            "ncc_mst": mst.split("-")[0]}


# ---------------------------------------------------------------------------
# Tải PDF gốc từ trang tra cứu của nhà cung cấp (bằng mã tra cứu trong XML)
# ---------------------------------------------------------------------------

def _http_get(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": UA_LOGIN, "Accept": "*/*"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
        raise PortalError("Không kết nối được trang của nhà cung cấp: %s" % getattr(e, "reason", e))


MISA_WWW, MISA_APEX, MISA_DL = "https://www.meinvoice.vn", "https://meinvoice.vn", "https://download.meinvoice.vn"


def misa_pdf(code):
    """MISA meInvoice – tải PDF gốc bằng mã tra cứu, không cần tài khoản. Thử lần lượt:
    0) đúng link nút "Tải hóa đơn dạng PDF" của trang tra cứu (Default.js của MISA):
       /tra-cuu/DownloadHandler.ashx?Type=pdf&Code=<mã> – không cần mã ext;
    1) mở trang tra cứu ?sc=<mã> như trình duyệt (giữ cookie) và lấy link DownloadHandler.ashx có sẵn trong trang;
    2) lấy mã thời gian 'ext' từ GetRequestTimeEnCode (theo tài liệu MISA) rồi dựng link
       www.meinvoice.vn/tra-cuu/tra-cuu/DownloadHandler.ashx?Type=pdf&Viewer=1&ext=<ext>&Code=<mã>
       (dạng thấy trong trang tra cứu thật, ext có thể chứa '_' – vd J1V4E6D_) và download.meinvoice.vn."""
    import html as H
    jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    lookup = MISA_WWW + "/tra-cuu/?sc=" + urllib.parse.quote(code)

    def get(url, referer=lookup):
        req = urllib.request.Request(url, headers={"User-Agent": UA_LOGIN, "Accept": "*/*", "Referer": referer,
                                                   "Accept-Language": "vi-VN,vi;q=0.9"})
        try:
            with opener.open(req, timeout=30) as resp:
                return resp.read()
        except urllib.error.HTTPError as e:
            return e.read() or b""
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            raise PortalError("Không kết nối được meinvoice.vn: %s" % getattr(e, "reason", e))

    urls, diag = [], []
    page = get(lookup, referer=MISA_WWW + "/tra-cuu/").decode("utf-8", "replace")  # lấy cookie như trình duyệt
    direct = get(MISA_WWW + "/tra-cuu/DownloadHandler.ashx?Type=pdf&Code=" + urllib.parse.quote(code))
    if direct[:4] == b"%PDF":
        return direct
    for link in re.findall(r"DownloadHandler\.ashx\?[^\"'<>\s]+", page, re.I):
        urls.append(urllib.parse.urljoin(MISA_WWW + "/tra-cuu/", "tra-cuu/" + H.unescape(link)))
    exts = []
    for src in (MISA_WWW + "/tra-cuu/GetRequestTimeEnCode", MISA_APEX + "/tra-cuu/GetRequestTimeEnCode"):
        raw = get(src).decode("utf-8", "replace")
        diag.append(re.sub(r"\s+", " ", raw)[:80])
        plain = re.sub(r"<[^>]+>", "", raw).strip().strip('"')
        if re.fullmatch(r"[\w-]{8}", plain):
            exts.append(plain)
        if len(raw) >= 63:
            exts.append(raw[55:63])  # cách lấy trong mã mẫu của MISA (Substring(55, 8))
        exts += re.findall(r"(?<![\w-])[\w-]{8}(?![\w-])", plain)[:3]
        if exts:
            break
    q = urllib.parse.quote(code)
    for ext in dict.fromkeys(exts):
        e = urllib.parse.quote(ext)
        urls.append(MISA_WWW + "/tra-cuu/tra-cuu/DownloadHandler.ashx?Type=pdf&Viewer=1&ext=%s&Code=%s" % (e, q))
        urls.append(MISA_DL + "/downloadhandler.ashx?type=pdf&code=%s&viewer=1&ext=%s" % (q, e))
    got = []
    for url in dict.fromkeys(urls):
        data = get(url)
        if data[:4] == b"%PDF":
            return data
        got.append(re.sub(r"\s+", " ", re.sub(rb"<[^>]+>", b" ", data[:600]).decode("utf-8", "replace")).strip()[:70])
    raise PortalError("MISA không trả file PDF cho mã tra cứu %s (đã thử %d đường dẫn; GetRequestTimeEnCode trả: %s; "
                      "trang tải trả: %s)" % (code, len(urls), " | ".join(diag) or "rỗng", " | ".join(got) or "rỗng"))


# ---- EasyInvoice (SoftDreams): trang tra cứu riêng của từng người bán, captcha 4 chữ số ----

def png_pixels(data):
    """Giải mã PNG 8-bit (xám / RGB / RGBA, không xen kẽ) chỉ bằng thư viện chuẩn → (rộng, cao, hàng điểm ảnh RGB)."""
    import struct
    import zlib
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError("không phải ảnh PNG")
    pos, idat, w, h, ctype = 8, b"", 0, 0, 0
    while pos + 8 <= len(data):
        n, typ = struct.unpack(">I4s", data[pos:pos + 8])
        chunk = data[pos + 8:pos + 8 + n]
        pos += 12 + n
        if typ == b"IHDR":
            w, h, depth, ctype, _, _, inter = struct.unpack(">IIBBBBB", chunk)
            if depth != 8 or inter or ctype not in (0, 2, 4, 6):
                raise ValueError("định dạng PNG chưa hỗ trợ")
        elif typ == b"IDAT":
            idat += chunk
        elif typ == b"IEND":
            break
    bpp = {0: 1, 2: 3, 4: 2, 6: 4}[ctype]
    raw = zlib.decompress(idat)
    stride = w * bpp
    rows, prev, i = [], bytearray(stride), 0
    for _ in range(h):
        f, line = raw[i], bytearray(raw[i + 1:i + 1 + stride])
        i += 1 + stride
        for x in range(stride):
            a = line[x - bpp] if x >= bpp else 0
            b = prev[x]
            c = prev[x - bpp] if x >= bpp else 0
            if f == 1:
                line[x] = (line[x] + a) & 255
            elif f == 2:
                line[x] = (line[x] + b) & 255
            elif f == 3:
                line[x] = (line[x] + ((a + b) >> 1)) & 255
            elif f == 4:
                p = a + b - c
                pa, pb, pc = abs(p - a), abs(p - b), abs(p - c)
                line[x] = (line[x] + (a if pa <= pb and pa <= pc else b if pb <= pc else c)) & 255
        prev = line
        if ctype in (0, 4):
            rows.append([(line[k * bpp],) * 3 for k in range(w)])
        else:
            rows.append([tuple(line[k * bpp:k * bpp + 3]) for k in range(w)])
    return w, h, rows


# Mẫu chữ số của captcha EasyInvoice (chữ trắng font có chân trên nền xám): lưới 10x14 điểm, dạng hex.
# Tạo từ font Liberation Serif (cùng số đo với Times New Roman) và từ captcha thật đã kiểm chứng.
EASY_DIGITS = {
    "0": "1e0cce1f87c0f03c0f03c0f03e1f87738fc 1e18c61b06c1f07c1f07c1f07c1b86639fc 1f0ce63f87c1f07c1f07c1f07e1d8f73cfe 1f1cf61f87e1f03c0f03c0f87e1d8773cfe 1f1ff61d03c0701c0701c0701c0d877fcfe 3f1ce60b83e0f83e0f83e0f83e0d82738fc",
    "1": "060f8ee0380e0380e0380e0380e0380e3ff 0e0f8ee0380e0380e0380e0380e0380e3ff 0e0f8fe0781e0781e0781e0781e0781e3ff 0e1f86e0380e0380e0380e0380e0380e3ff 0f1fcef03c0f03c0f03c0f03c0f03c0f3ff 1e3f806018060180601806018060181e1ff 1e3f81e018060180601806018060181e1ff",
    "2": "1e0fe63c0300c020080c00000000007ffff 1f1ff41c0300c0701c0e0f0781c3c0fffff 7e10e41906018060380c06070381c0fffff 7f18e41d0701c070180c06070381c0fffff 7f19f43d0701c070381e0f078181c0fffff 7f38ec380f0380e0381c0e07038180fffff 7f38ec3c0300c0f0381e0f0783c1e0fffff",
    "3": "0f0fe01c010180e0f80f0040100401e1ffe 1f1fe01c030180e1f80f00c0300c03c1ff6 1f1ff61c0301c7f1f80700c0300d07ffffe 1f1ff61c0301c7f1f80700c0300f07ffffe 3e3fec1c030187e1f00600c0300e07ffbfc 7e10e419060180e1f00e01c0380e07c3ffe 7e30c818060380e1f07e0380781e07c3bfe 7f19e43d0f03c1e1f81e01c0781e07c3ffe 7f19f43d0701c0f1f87f03c0700f07c3fff 7f38ec380e0381e1f00e00c03c0f03e3bfc",
    "4": "02018020080200802008023ffffc0802008 0301c0f034390c471104fffffffc0401004 0301c0f03c0f04c331cc633ffffc0c0300c 0380e0783e1b8ee3b98e43bffffc0e0380e 0380e0783e1b8ee3b98e63bffffc0e0380e 0380e0783e1f86e338ce63bfffffff0380e 0381e0783e1b86e1388e43bffffc0e0380e",
    "5": "0fcfe601c0fe3fc0f8060040100000c03e0 7f9fc40100401f87f80e01c0781e07c3bfe 7f9fe40100401007f00601c0380f07c3ffe 7f9fe40100401007f00f01c0781f07c3ffe 7f9fe60180601807f00e00c03c0f03e39fc 7fdff40100401007fc0f01c0700f0763dff",
    "6": "0f8c260b80e037cffb03c0f83e0d8371cfe 0f9fe781c0c037cffb07c0f03c0dcf7f9fc 1f8c660980c037cffb87c1f03e0f87718fe 1f8e3707c0e0380ffb8fe0f83e0dc370c7c 1fce370d80e0380fffcfe1f8360dc779cff 1fce770d80e03fefff07c1f87e1d8773cff",
    "7": "fffff80e060180c060180c030180c030180 fffff81a060300c060100c060180c030180 fffffc0a060300c060180c03018060300c0 fffffc0c020380c070180e0301c060380c0 fffffc0f0e0381c070380c070380e0781c0",
    "8": "1f1cee3f87e1dcf7f9cee1f07c1f87e3dff 1f1cfe1f8761dcf7fdcfe1f07c0f87f3dff 3e18c61b06e198e7f1fee3b07c1f07e19fe 3f186c0701f1dfe7e0fc63b07c0701e05fc 3f186c0701f1dfe7e0fc63b07c0701e1dfe 3f18ee1f87e1d8e7f986c0f03c0f03e1dfe 3f1cee0f83e0dce3f1cee0f83e0f83f38fc",
    "9": "1e08441b02c0b03c0d861f806030181c1e0 1f1ffe1f0380701e0dff1ec0300c1f7f9fc 1f3cfe1f07c1f07f3dff3fc0701f0fc7bfe 3e18cc1b06c1f07e1dff3fc0601a0ec73fc 3e18cc1f07c0f03c0dc73ec0701e06c3bfc 3f1dee3f07c1f07e1dcf3fc0703e0fc7bfc 3f3cee0b83e0f83f0dfb00c0303f0ee73f8",
}


def _easy_glyphs(data, thr=200):
    """Tách các chữ số (điểm ảnh gần trắng) theo cột; đường nhiễu màu xám bị loại."""
    w, h, rows = png_pixels(data)
    on = [[min(p) > thr for p in r] for r in rows]
    cols = [any(on[y][x] for y in range(h)) for x in range(w)]
    out, x = [], 0
    while x < w:
        if not cols[x]:
            x += 1
            continue
        s = x
        while x < w and cols[x]:
            x += 1
        ys = [y for y in range(h) if any(on[y][s:x])]
        if ys and (x - s) * (ys[-1] - ys[0] + 1) >= 20:
            out.append([r[s:x] for r in on[ys[0]:ys[-1] + 1]])
    return out


def _easy_norm(g, W=10, H=14):
    gh, gw = len(g), len(g[0])
    return int("".join("1" if any(g[min(gh - 1, int((y + dy / 3) * gh / H))][min(gw - 1, int((x + dx / 3) * gw / W))]
                                  for dy in range(3) for dx in range(3)) else "0"
                       for y in range(H) for x in range(W)), 2)


def solve_easy_captcha(data, max_dist=34):
    """Đọc captcha EasyInvoice → chuỗi 4 chữ số, hoặc None nếu không chắc chắn (khi đó lấy captcha khác)."""
    try:
        glyphs = _easy_glyphs(data)
    except (ValueError, KeyError):
        return None
    if len(glyphs) != 4:
        return None
    templates = [(int(hx, 16), d) for d, hs in EASY_DIGITS.items() for hx in hs.split()]
    text = ""
    for g in glyphs:
        v = _easy_norm(g)
        dist, d = min((bin(v ^ t).count("1"), d) for t, d in templates)
        if dist > max_dist:
            return None
        text += d
    return text


EASY_BASE_RE = r"https?://[\w-]+\.easyinvoice\.(com\.)?vn"


def easyinvoice_pdf(tra_cuu, tries=8):
    """EasyInvoice: mở trang tra cứu của người bán, tự giải captcha, tra bằng mã tra cứu rồi bấm
    "Tải PDF & đính kèm" như trình duyệt: gửi HTML hoá đơn lên /Invoice/DownloadPdfAndFileAttachFromAvailableHtml,
    nhận fileGuid và tải /Invoice/Download (PDF, hoặc ZIP gồm PDF + file đính kèm)."""
    import base64
    import html as H
    base = (tra_cuu.get("url") or "").rstrip("/")
    code = tra_cuu.get("code") or ""
    if not re.fullmatch(EASY_BASE_RE, base):
        raise PortalError("Không xác định được trang tra cứu EasyInvoice của người bán")
    jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))

    def call(path, form=None, ajax=False):
        headers = {"User-Agent": UA_LOGIN, "Accept": "*/*", "Accept-Language": "vi-VN,vi;q=0.9", "Referer": base + "/"}
        if ajax:
            headers["X-Requested-With"] = "XMLHttpRequest"
        body = urllib.parse.urlencode(form).encode() if form is not None else None
        if body is not None:
            headers["Content-Type"] = "application/x-www-form-urlencoded; charset=UTF-8"
            headers["Origin"] = base
        try:
            with opener.open(urllib.request.Request(base + path, data=body, headers=headers), timeout=60) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            return e.read() or b""
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            raise PortalError("Không kết nối được %s: %s" % (base, getattr(e, "reason", e)))

    last = ""
    for _ in range(tries):
        page = call("/").decode("utf-8", "replace")
        form = re.search(r'<form[^>]*id="Search"[^>]*>(.*?)</form>', page, re.S | re.I)
        fields = {}
        for tag in re.findall(r"<input[^>]*>", form.group(1) if form else page):
            name = re.search(r'name="([^"]*)"', tag)
            if name:
                val = re.search(r'value="([^"]*)"', tag)
                fields[name.group(1)] = H.unescape(val.group(1)) if val else ""
        cap = solve_easy_captcha(call("/Captcha/Show"))
        if not cap:
            last = "không đọc được captcha"
            continue
        fields.update(FKey=code, Capcha=cap)
        res = call("/Search/Search", fields).decode("utf-8", "replace")
        token = re.search(r"downloadPdfAndFileAttachFromAvailableHtml\('([^']+)'\)", res)
        if not token:
            msg = re.search(r'id="msg"[^>]*value="([^"]*)"', res) or re.search(r'value="([^"]*)"[^>]*id="msg"', res)
            last = H.unescape(msg.group(1)) if msg else "trang tra cứu không trả kết quả"
            if re.search(r"x[aá]c th[ựu]c|captcha|capcha", _plain(last) + last.lower()):
                continue  # đọc sai captcha → thử captcha khác
            raise PortalError("EasyInvoice: %s" % last)
        inv = re.search(r'id="InvData"[^>]*value="([^"]*)"', res) or re.search(r'value="([^"]*)"[^>]*id="InvData"', res)
        try:
            html_doc = json.loads(H.unescape(inv.group(1))).get("str", "") if inv else ""
        except ValueError:
            html_doc = ""
        html_doc = re.sub(r"^\s*<\?xml[^>]*\?>", "", html_doc)
        js = call("/Invoice/DownloadPdfAndFileAttachFromAvailableHtml",
                  {"token": token.group(1), "html": base64.b64encode(html_doc.encode("utf-8")).decode()}, ajax=True)
        try:
            j = json.loads(js.decode("utf-8", "replace"))
        except ValueError:
            raise PortalError("EasyInvoice không trả mã file PDF")
        if j.get("msg"):
            raise PortalError("EasyInvoice: %s" % j["msg"])
        data = call("/Invoice/Download?" + urllib.parse.urlencode({"fileGuid": j.get("fileGuid", ""),
                                                                   "fileName": j.get("fileName", "")}))
        if data[:4] == b"%PDF":
            return data
        if data[:2] == b"PK":
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for n in z.namelist():
                    if n.lower().endswith(".pdf"):
                        return z.read(n)
        raise PortalError("EasyInvoice trả file không phải PDF")
    raise PortalError("EasyInvoice: thử %d lần chưa được (%s)" % (tries, last))


# MST nhà cung cấp → hàm tải PDF gốc (nhận thông tin tra cứu). Viettel, VNPT… chưa tự động được.
PDF_FETCHERS = {"0101243150": lambda t: misa_pdf(t["code"]), "0105987432": easyinvoice_pdf}


# Link tải PDF gốc mà trình duyệt của người dùng mở được (MISA chặn chương trình tự động nhưng không chặn trình duyệt).
PDF_BROWSER_URLS = {"0101243150": lambda code: MISA_WWW + "/tra-cuu/DownloadHandler.ashx?Type=pdf&Code=" +
                    urllib.parse.quote(code)}


def browser_pdf_url(tra_cuu):
    ok, prov = can_fetch_pdf(tra_cuu)
    return PDF_BROWSER_URLS[prov](tra_cuu["code"]) if ok and prov in PDF_BROWSER_URLS else ""


def default_downloads_dir():
    """Thư mục Downloads của người dùng (Windows: theo cài đặt của Windows nếu đọc được)."""
    if os.name == "nt":
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                                r"Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders") as k:
                return winreg.QueryValueEx(k, "{374DE290-123F-4565-9164-39C4925E467B}")[0]
        except OSError:
            pass
    return os.path.join(os.path.expanduser("~"), "Downloads")


def browser_download_settings(base=None):
    """Đọc cài đặt tải xuống của Chrome / Edge (mọi hồ sơ người dùng): thư mục lưu và có bật
    'Hỏi vị trí lưu từng tệp' không. base: thư mục LOCALAPPDATA (để kiểm thử)."""
    base = base or os.environ.get("LOCALAPPDATA") or ""
    dirs, prompt = [], False
    for app_dir in (os.path.join(base, "Google", "Chrome", "User Data"), os.path.join(base, "Microsoft", "Edge", "User Data")):
        try:
            profiles = [d for d in os.listdir(app_dir) if d == "Default" or d.startswith("Profile ")]
        except OSError:
            continue
        for prof in profiles:
            prefs = _load_json(os.path.join(app_dir, prof, "Preferences"), {})
            dl = prefs.get("download") or {}
            if dl.get("default_directory"):
                dirs.append(dl["default_directory"])
            if dl.get("prompt_for_download"):
                prompt = True
    return {"dirs": list(dict.fromkeys(dirs)), "prompt": prompt}


def watch_dirs(app):
    """Các thư mục có thể chứa PDF trình duyệt vừa tải: thư mục cấu hình, thư mục tải của Chrome/Edge, Downloads, Desktop."""
    st = browser_download_settings()
    cands = [app.downloads] + st["dirs"] + [default_downloads_dir(), os.path.join(os.path.expanduser("~"), "Desktop")]
    return [d for d in dict.fromkeys(os.path.abspath(c) for c in cands if c) if os.path.isdir(d)], st["prompt"]


def scan_downloads(out_root, mst, folders, since, seen):
    """Gắn các PDF mới xuất hiện trong các thư mục tải về (sau thời điểm 'since') vào đúng hoá đơn.
    seen: tập (đường dẫn, mtime) đã xử lý để không đọc lại."""
    files = []
    if isinstance(folders, str):
        folders = [folders]
    for folder in folders:
        try:
            names = os.listdir(folder)
        except OSError:
            continue
        files += _new_pdfs(folder, names, since, seen)
    return import_pdfs(out_root, mst, files) if files else []


def _new_pdfs(folder, names, since, seen):
    files = []
    for name in names:
        path = os.path.join(folder, name)
        if not name.lower().endswith(".pdf") or not os.path.isfile(path):
            continue
        mt = os.path.getmtime(path)
        if mt < since - 5 or (path, mt) in seen:
            continue
        seen.add((path, mt))
        try:
            with open(path, "rb") as fh:
                files.append((name, fh.read()))
        except OSError:
            continue
    return files


def can_fetch_pdf(tra_cuu):
    t = tra_cuu or {}
    mst = t.get("ncc_mst") or ("0101243150" if str(t.get("ncc", "")).startswith("MISA") else "")
    return bool(t.get("code")) and mst in PDF_FETCHERS, mst


def fetch_original_pdf(out_root, mst, kind, key):
    """Tải PDF gốc của một hoá đơn trong kho rồi gắn vào hoá đơn. Trả về đường dẫn tương đối."""
    e = _load_json(index_file(out_root, mst), {}).get(kind, {}).get(key)
    if e is None:
        raise ValueError("Không tìm thấy hoá đơn")
    ok, prov = can_fetch_pdf(e.get("tra_cuu"))
    if not ok:
        raise ValueError("Chưa hỗ trợ tự tải PDF gốc của nhà cung cấp này – bấm 'Tra cứu' để tải tay rồi '+PDF'")
    return attach_pdf(out_root, mst, kind, key, PDF_FETCHERS[prov](e["tra_cuu"]))


def parse_invoice_xml(xml_path):
    """Đọc XML hoá đơn → {'items': [...], 'rates': [...], 'httt', 'nb_dchi', 'nm_dchi', 'nky'}."""
    import xml.etree.ElementTree as ET
    info = {"items": [], "rates": [], "httt": "", "nb_dchi": "", "nm_dchi": "", "nky": None, "msttcgp": "",
            "ttkhac": {}, "dlhdon_id": "", "nb": {}, "nm": {}, "ttchung": {}, "tong": {}}
    try:
        root = ET.parse(xml_path).getroot()
    except (ET.ParseError, OSError):
        return info

    def child(el, name):
        for c in el:
            if _local(c.tag) == name:
                return c
        return None

    def text(el, name):
        c = child(el, name) if el is not None else None
        return (c.text or "").strip() if c is not None and c.text else ""

    for el in root.iter():
        tag = _local(el.tag)
        if tag == "HTTToan" and not info["httt"]:
            info["httt"] = (el.text or "").strip()
        elif tag == "NBan":
            info["nb_dchi"] = info["nb_dchi"] or text(el, "DChi")
        elif tag == "NMua":
            info["nm_dchi"] = info["nm_dchi"] or text(el, "DChi")
        elif tag == "SigningTime" and not info["nky"]:
            info["nky"] = _to_dt(el.text)
        elif tag == "HHDVu":
            nature = text(el, "TChat")
            if nature == "4":  # dòng ghi chú/diễn giải
                continue
            amount = _float(text(el, "ThTien"))
            if nature == "3":  # chiết khấu thương mại giảm trừ
                amount = -abs(amount)
            rate = norm_rate(text(el, "TSuat"))
            tax_text = text(el, "TThue")
            if tax_text:
                tax = _float(tax_text)
            else:
                tax = round(amount * float(rate) / 100, 2) if rate in VAT_RATES else 0.0
            info["items"].append({
                "name": text(el, "THHDVu"), "unit": text(el, "DVTinh"), "qty": _float(text(el, "SLuong")),
                "price": _float(text(el, "DGia")), "amount": amount, "rate": rate, "tax": tax,
                "discount": _float(text(el, "STCKhau")), "nature": nature})
        elif tag == "LTSuat":
            info["rates"].append({"rate": norm_rate(text(el, "TSuat")), "base": _float(text(el, "ThTien")),
                                  "tax": _float(text(el, "TThue"))})
        elif tag == "MSTTCGP" and not info["msttcgp"]:
            info["msttcgp"] = (el.text or "").strip()
        elif tag == "TTin":
            k, v = text(el, "TTruong"), text(el, "DLieu")
            if k and v:
                info["ttkhac"].setdefault(k, v)
        elif tag == "DLHDon" and not info["dlhdon_id"]:
            info["dlhdon_id"] = el.get("Id", "")
        if tag in ("NBan", "NMua"):
            info["nb" if tag == "NBan" else "nm"] = {_local(c.tag): (c.text or "").strip() for c in el if len(c) == 0}
        elif tag == "TTChung":
            info["ttchung"] = {_local(c.tag): (c.text or "").strip() for c in el if len(c) == 0}
        elif tag == "TToan":
            info["tong"] = {_local(c.tag): (c.text or "").strip() for c in el if len(c) == 0}
    return info


def invoice_lines(inv):
    """Danh sách dòng hàng của một hoá đơn; không có XML thì dựng từ tổng tiền của cổng."""
    info = inv.get("_xmlinfo") or {}
    if info.get("items"):
        return info["items"]
    rates = [{"rate": norm_rate(r.get("tsuat")), "base": _float(r.get("thtien")), "tax": _float(r.get("tthue"))}
             for r in (inv.get("thttltsuat") or []) if isinstance(r, dict)]
    if not rates:
        rates = [{"rate": "", "base": _float(inv.get("tgtcthue")), "tax": _float(inv.get("tgtthue"))}]
    return [{"name": "", "unit": "", "qty": 0.0, "price": 0.0, "amount": r["base"], "rate": r["rate"],
             "tax": r["tax"], "discount": 0.0, "nature": ""} for r in rates]


def rate_groups(inv):
    """Gom dòng hàng theo thuế suất (đúng cách lập bảng kê): [{rate, names, base, tax, items}]."""
    groups = {}
    for it in invoice_lines(inv):
        g = groups.setdefault(it["rate"], {"rate": it["rate"], "names": [], "base": 0.0, "tax": 0.0, "items": []})
        if it["name"]:
            g["names"].append(it["name"])
        g["base"] += it["amount"]
        g["tax"] += it["tax"]
        g["items"].append(it)
    # Tổng theo thuế suất trong XML (LTSuat) là số chính thức – ưu tiên dùng.
    official = {r["rate"]: r for r in (inv.get("_xmlinfo") or {}).get("rates", [])}
    for rate, g in groups.items():
        if rate in official:
            g["base"], g["tax"] = official[rate]["base"], official[rate]["tax"]
    if len(groups) == 1:
        g = next(iter(groups.values()))
        if not official and inv.get("tgtthue") is not None:
            g["base"], g["tax"] = _float(inv.get("tgtcthue")), _float(inv.get("tgtthue"))
    return list(groups.values())


def _status_note(inv):
    parts = []
    tthai = _num(inv.get("tthai"))
    if tthai != 1:
        parts.append(TTHAI_NIBOT.get(tthai, ""))
    if _num(inv.get("ttxly")) == 8 or inv.get("_prefix") == "/sco-query":
        parts.append(TTXLY_NIBOT[8])
    return " - ".join(p for p in parts if p)


def _is_bad(inv):
    """HĐ bị thay thế / bị huỷ: không đưa vào bảng kê, xếp vào nhóm cần xem xét."""
    return _num(inv.get("tthai")) in (4, 6)


def write_nibot_workbook(path, kind, invoices, mst, company_name, start, end, changes=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    wrap_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    money = "#,##0"
    dfmt = "dd/mm/yyyy"
    purchase = kind == "purchase"
    period = ("Tháng %d năm %d" % (start.month, start.year)
              if start.day == 1 and end == date(start.year, start.month, calendar.monthrange(start.year, start.month)[1])
              else "Từ ngày %s đến ngày %s" % (start.strftime("%d/%m/%Y"), end.strftime("%d/%m/%Y")))
    today = date.today()
    sign_date = "..............., ngày %02d tháng %02d năm %d" % (today.day, today.month, today.year)

    def header_row(ws, row, values, start_col=1):
        for i, v in enumerate(values):
            c = ws.cell(row=row, column=start_col + i, value=v)
            c.font, c.fill, c.alignment, c.border = bold, head_fill, wrap_c, box

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # 1. HoaDon_TongQuat ---------------------------------------------------
    ws = wb.active
    ws.title = "HoaDon_TongQuat"
    ws["A1"] = "HÓA ĐƠN MUA VÀO" if purchase else "HÓA ĐƠN BÁN RA"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws["A2"] = "MST: %s" % mst
    ws["A3"] = "Tên DN: %s" % (company_name or "")
    heads = ["Loại HĐ", "MST người bán", "Người bán", "Địa chỉ người bán", "MST người mua", "Người Mua",
             "Địa chỉ người mua", "Ngày", "HTTT", "Ký hiệu", "Số", "Trạng thái HĐ", "Kết quả kiểm tra",
             "Tiền Chưa thuế", "Tiền Thuế", "Tiền CK TM", "Tiền Phí", "Tiền Thanh toán", "Duyệt nội bộ", "Ghi chú",
             "File XML"]
    header_row(ws, 4, heads)
    r = 5
    for inv in invoices:
        info = inv.get("_xmlinfo") or {}
        d = invoice_date(inv)
        row = [LOAI_HD.get(_num(inv.get("khmshdon")), str(inv.get("khmshdon") or "")),
               str(inv.get("nbmst") or ""), inv.get("nbten") or "", inv.get("nbdchi") or info.get("nb_dchi", ""),
               str(inv.get("nmmst") or ""), inv.get("nmten") or "", inv.get("nmdchi") or info.get("nm_dchi", ""),
               datetime(d.year, d.month, d.day) if d else None, inv.get("thtttoan") or info.get("httt", ""),
               "%s%s" % (inv.get("khmshdon") or "", inv.get("khhdon") or ""), str(inv.get("shdon") or ""),
               TTHAI_NIBOT.get(_num(inv.get("tthai")), inv.get("tthai")),
               TTXLY_NIBOT.get(_num(inv.get("ttxly")), inv.get("ttxly")),
               _float(inv.get("tgtcthue")), _float(inv.get("tgtthue")), _float(inv.get("ttcktmai")),
               _float(inv.get("tgtphi")), _float(inv.get("tgtttbso")), inv.get("_duyet") or "Chờ duyệt",
               " - ".join(x for x in (_status_note(inv), inv.get("_note")) if x),
               inv.get("_xml", "")]
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            if i == 8:
                c.number_format = dfmt
            elif 14 <= i <= 18:
                c.number_format = money
        r += 1
    ws.cell(row=r, column=1, value="Total").font = bold
    if r > 5:
        ws.cell(row=r, column=11, value="=SUBTOTAL(103,K5:K%d)" % (r - 1)).font = bold
        for col in "NOPQR":
            c = ws["%s%d" % (col, r)]
            c.value, c.font, c.number_format = "=SUBTOTAL(109,%s5:%s%d)" % (col, col, r - 1), bold, money
    ws.auto_filter.ref = "A4:%s%d" % (get_column_letter(len(heads)), max(r - 1, 4))
    ws.freeze_panes = "A5"
    widths(ws, [7, 14, 36, 36, 14, 36, 36, 11, 14, 11, 9, 16, 18, 14, 13, 11, 10, 15, 11, 30, 40])

    if changes:
        wc = wb.create_sheet("Doi_TrangThai")
        header_row(wc, 1, CHANGE_HEAD)
        for row in changes:
            wc.append(row)
        widths(wc, [12, 12, 10, 14, 40, 20, 20])

    if not purchase:
        wi = wb.create_sheet("ChiTiet_HangHoa")
        header_row(wi, 1, ITEM_HEAD + ["Tên hàng hoá, dịch vụ", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền",
                                       "Thuế suất", "Tiền thuế"])
        for inv in invoices:
            for it in invoice_lines(inv):
                wi.append(_inv_head(inv) + [it["name"], it["unit"], it["qty"], it["price"], it["amount"],
                                            it["rate"], it["tax"]])
        for row in wi.iter_rows(min_row=2):
            for idx in (10, 11, 13):
                row[idx].number_format = money
        widths(wi, [11, 11, 9, 14, 34, 14, 34, 40, 8, 9, 12, 14, 8, 12])
        wb.save(path)
        return path

    good = [i for i in invoices if not _is_bad(i)]
    bad = [i for i in invoices if _is_bad(i)]

    # 2. Smart_KTSC_OK / CAN_XEM_XET (mẫu import phần mềm kế toán Smart Pro) ------
    def smart_rows(inv, bad_inv):
        info = inv.get("_xmlinfo") or {}
        d = invoice_date(inv)
        ngay = datetime(d.year, d.month, d.day) if d else None
        total = _float(inv.get("tgtttbso"))
        big = total > 5000000
        sr = "%s%s" % (inv.get("khmshdon") or "", inv.get("khhdon") or "")
        so = str(inv.get("shdon") or "")
        ngay_txt = d.strftime("%d/%m/%Y") if d else ""
        ghichu = []
        if big:
            ghichu.append("HĐ có giá trị > 5,000,000 đồng")
        note = _status_note(inv)
        if note:
            ghichu.append(note)
        lctg, tkco = ("PKT", "331") if big else ("PC", "1111")
        dien_giai = ("Chi phí theo hđơn ký hiệu số %s - %s - %s" if big else "Chi trả tiền theo hđơn số %s - %s - %s") % (
            sr, so, ngay_txt)
        nbmst, nbten = str(inv.get("nbmst") or ""), inv.get("nbten") or ""
        dchi = inv.get("nbdchi") or info.get("nb_dchi", "")
        guid = str(uuid.uuid5(uuid.NAMESPACE_URL, invoice_key(inv)))
        lines = invoice_lines(inv)
        if bad_inv:
            label = TTHAI_NIBOT.get(_num(inv.get("tthai")), "")
            lines = [{"name": label, "unit": "", "qty": 0.0, "price": 0.0, "amount": 0.0,
                      "rate": (lines[0]["rate"] if lines else ""), "tax": 0.0, "discount": 0.0}]
            ghichu, dien_giai, lctg = [label], "", ""
        out = []
        for it in lines:
            out.append(["; ".join(ghichu), lctg, sr, so, info.get("nky") or _to_dt(inv.get("nky")), ngay,
                        _num(so), ngay, dien_giai, inv.get("thtttoan") or info.get("httt", ""), "", "", tkco,
                        nbmst, "", "", it["name"], it["name"], 0, "", it["unit"], it["qty"], 0, 0, 0, it["price"],
                        it["amount"], 0, it.get("discount", 0.0), "V", "1331", it["rate"], 0, it["tax"], 0,
                        it["amount"] + it["tax"], nbmst, nbten, nbten, dchi, nbmst, dchi, "", "TIENHANG", "", guid])
        return out

    for title, src, is_bad in (("Smart_KTSC_OK", good, False), ("Smart_KTSC_CAN_XEM_XET", bad, True)):
        if is_bad and not src:
            continue
        wsm = wb.create_sheet(title)
        wsm.append(SMART_HEAD)
        for c in wsm[1]:
            c.font = bold
        for inv in src:
            for row in smart_rows(inv, is_bad):
                wsm.append(row)
        for row in wsm.iter_rows(min_row=2):
            for idx in (4, 5, 7):
                row[idx].number_format = dfmt
            for idx in (25, 26, 33, 35):
                row[idx].number_format = money
        widths(wsm, [28, 6, 11, 9, 18, 11, 9, 11, 40, 14, 6, 6, 6, 14, 6, 6, 30, 30] + [8] * 28)

    # 3. Bảng kê mua vào (kèm tờ khai 01/GTGT) ---------------------------------
    def bang_ke(title, heading, groups, with_tax):
        wsb = wb.create_sheet(title)
        last = "M"
        wsb["A4"] = heading
        wsb["A4"].font = Font(bold=True, size=13)
        wsb["A4"].alignment = Alignment(horizontal="center")
        wsb.merge_cells("A4:%s4" % last)
        wsb["B6"] = "(Kèm theo tờ khai thuế GTGT theo mẫu số 01/GTGT)"
        wsb.merge_cells("B6:%s6" % last)
        wsb["B7"] = "Kỳ tính thuế: %s" % period
        wsb.merge_cells("B7:%s7" % last)
        for cell in ("B6", "B7"):
            wsb[cell].alignment = Alignment(horizontal="center")
        wsb["B9"] = "Người nộp thuế: %s" % (company_name or "")
        wsb["B10"] = "Mã số thuế: %s" % mst
        wsb["B12"] = "Đơn vị tiền: đồng Việt Nam"
        header_row(wsb, 13, ["STT", "Hoá đơn, chứng từ, biên lai nộp thuế", "", "", "", "Tên người bán",
                             "Mã số thuế người bán", "Mặt hàng", "Doanh số mua chưa có thuế", "Thuế suất",
                             "Thuế GTGT\nđủ điều kiện khấu trừ thuế", "Ghi chú"], 2)
        header_row(wsb, 15, ["KH mẫu HĐ", "Ký hiệu hoá đơn", "Số hoá đơn", "Ngày, tháng, năm lập hóa đơn"], 3)
        for col in "BGHIJKLM":
            wsb.merge_cells("%s13:%s15" % (col, col))
        wsb.merge_cells("C13:F14")
        header_row(wsb, 16, ["[1]", "[2]", "[3]", "[4]", "[5]", "[6]", "[7]", "[8]", "[9]", "[10]", "[11]", "[12]"], 2)
        wsb["B17"] = ("1. HH, DV dùng riêng cho SXKD chịu thuế GTGT và sử dụng cho các hoạt động cung cấp HH, DV "
                      "không kê khai, nộp thuế GTGT đủ điều kiện khấu trừ thuế: ")
        wsb["B17"].font = bold
        r = 18
        for n, (inv, g) in enumerate(groups, 1):
            d = invoice_date(inv)
            vals = [n, _num(inv.get("khmshdon")), inv.get("khhdon") or "", _num(inv.get("shdon")),
                    datetime(d.year, d.month, d.day) if d else None, inv.get("nbten") or "",
                    str(inv.get("nbmst") or ""), g["names"][0] if g["names"] else "", g["base"],
                    g["rate"], g["tax"] if with_tax else None, _status_note(inv)]
            for i, v in enumerate(vals, 2):
                c = wsb.cell(row=r, column=i, value=v)
                c.border = box
                if i == 6:
                    c.number_format = dfmt
                elif i in (10, 12):
                    c.number_format = money
            r += 1
        wsb.cell(row=r, column=2, value="Tổng").font = bold
        if r > 18:
            for col in ("J", "L") if with_tax else ("J",):
                c = wsb["%s%d" % (col, r)]
                c.value, c.font, c.number_format = "=SUM(%s18:%s%d)" % (col, col, r - 1), bold, money
        total_row = r
        r += 1
        wsb.cell(row=r, column=2, value="2. HH, DV dùng chung cho SXKD chịu thuế và không chịu thuế đủ điều kiện "
                                         "khấu trừ thuế:").font = bold
        wsb.cell(row=r + 2, column=2, value="Tổng").font = bold
        wsb.cell(row=r + 3, column=2, value="3. HH, DV dùng cho dự án đầu tư đủ điều kiện được khấu trừ thuế (*):").font = bold
        wsb.cell(row=r + 5, column=2, value="Tổng").font = bold
        r += 7
        wsb.cell(row=r, column=2, value="Tổng giá trị HHDV mua vào phục vụ SXKD được khấu trừ thuế GTGT (**):")
        c = wsb.cell(row=r, column=10, value="=J%d" % total_row)
        c.number_format, c.font = money, bold
        if with_tax:
            wsb.cell(row=r + 1, column=2, value="Tổng số thuế GTGT của HHDV mua vào đủ điều kiện được khấu trừ (***):")
            c = wsb.cell(row=r + 1, column=12, value="=L%d" % total_row)
            c.number_format, c.font = money, bold
        r += 3
        for i, t in enumerate([sign_date, "NGƯỜI NỘP THUẾ hoặc", "ĐẠI DIỆN HỢP PHÁP CỦA NGƯỜI NỘP THUẾ",
                               " Ký tên, đóng dấu (ghi rõ họ tên và chức vụ)"]):
            c = wsb.cell(row=r + i, column=10, value=t)
            c.alignment = Alignment(horizontal="center")
            c.font = Font(bold=0 < i < 3, italic=i in (0, 3))
            wsb.merge_cells("J%d:M%d" % (r + i, r + i))
        wsb.freeze_panes = "I17"
        widths(wsb, [2, 6, 6, 10, 10, 11, 36, 14, 34, 15, 8, 14, 26])
        return wsb

    vat_groups, kct_groups = [], []
    for inv in good:
        for g in rate_groups(inv):
            if _num(inv.get("khmshdon")) == 1 and g["rate"] in VAT_RATES:
                vat_groups.append((inv, g))
            else:
                kct_groups.append((inv, g))
    bang_ke("BangKe_MuaVao", "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ MUA VÀO", vat_groups, True)
    if kct_groups:
        bang_ke("BangKe_MuaVao_KCT_HDBH", "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ MUA VÀO KHÔNG CHỊU THUẾ, "
                "HÓA ĐƠN BÁN HÀNG", kct_groups, False)

    # 4. Bảng kê hoàn thuế (kèm Giấy đề nghị hoàn trả) ---------------------------
    wsh = wb.create_sheet("BangKe_HoanThue_OK")
    wsh["B1"] = ("BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ MUA VÀO\n(Kèm theo Giấy đề nghị hoàn trả khoản thu "
                 "NSNN số            ngày      tháng     năm    )")
    wsh["B1"].font = Font(bold=True, size=12)
    wsh["B1"].alignment = wrap_c
    wsh.merge_cells("B1:P3")
    wsh["B4"] = "[01] Kỳ đề nghị hoàn thuế: %s" % period
    wsh["B4"].alignment = Alignment(horizontal="center")
    wsh.merge_cells("B4:P4")
    wsh["B7"] = "[02] Tên người nộp thuế: %s" % (company_name or "")
    wsh["B8"] = "[03] Mã số thuế: %s" % mst
    wsh["B9"] = "[04] Tên đại lý thuế (nếu có):"
    wsh["B10"] = "[05] Mã số thuế: "
    wsh["B11"] = "Đơn vị tiền: Đồng Việt Nam"
    header_row(wsh, 13, ["STT", "Hoá đơn, chứng từ nộp thuế", "", "", "", "Tên người bán", "Mã số thuế \nngười bán",
                         "Tên hàng hóa, dịch vụ ", "Đơn \nvị \ntính", "Số \nlượng", "Đơn giá",
                         "Giá trị HHDV\nmua vào chưa có thuế GTGT", "Thuế suất GTGT (%)", "Thuế GTGT", "Ghi chú"], 2)
    header_row(wsh, 15, ["Mẫu \nsố", "Ký hiệu", "Số", "Ngày, tháng, năm"], 3)
    for col in "BGHIJKLMNOP":
        wsh.merge_cells("%s13:%s15" % (col, col))
    wsh.merge_cells("C13:F14")
    header_row(wsh, 16, ["[%d]" % i for i in range(1, 16)], 2)
    r = 17
    for inv, g in vat_groups:
        d = invoice_date(inv)
        single = len(g["items"]) == 1 and g["items"][0]["name"]
        it = g["items"][0]
        vals = [r - 16, str(inv.get("khmshdon") or ""), inv.get("khhdon") or "", str(inv.get("shdon") or ""),
                datetime(d.year, d.month, d.day) if d else None, inv.get("nbten") or "", str(inv.get("nbmst") or ""),
                "| ".join(g["names"]), it["unit"] if single else None, it["qty"] if single else 0,
                it["price"] if single else 0, g["base"], g["rate"], g["tax"], _status_note(inv)]
        for i, v in enumerate(vals, 2):
            c = wsh.cell(row=r, column=i, value=v)
            c.border = box
            if i == 6:
                c.number_format = dfmt
            elif i in (12, 13, 15):
                c.number_format = money
        r += 1
    wsh.cell(row=r, column=5, value="TỔNG CỘNG").font = bold
    if r > 17:
        for col in "MO":
            c = wsh["%s%d" % (col, r)]
            c.value, c.font, c.number_format = "=SUM(%s17:%s%d)" % (col, col, r - 1), bold, money
    r += 2
    wsh.cell(row=r, column=4, value="Tôi cam đoan số liệu khai trên là đúng và chịu trách nhiệm trước pháp luật về "
                                    "số liệu đã khai./.").font = Font(italic=True)
    for i, (left, right) in enumerate([("", sign_date), ("NHÂN VIÊN ĐẠI LÝ THUẾ", "NGƯỜI NỘP THUẾ hoặc "),
                                       ("Họ và tên:.............................", "ĐẠI DIỆN HỢP PHÁP CỦA NGƯỜI NỘP THUẾ"),
                                       ("Chứng chỉ hành nghề số:......",
                                        "(Chữ ký, ghi rõ họ tên; chức vụ và đóng dấu (nếu có)/Ký điện tử)")]):
        row = r + 1 + i
        if left:
            wsh.cell(row=row, column=4, value=left)
        c = wsh.cell(row=row, column=11, value=right)
        c.alignment = Alignment(horizontal="center")
        wsh.merge_cells("K%d:O%d" % (row, row))
    widths(wsh, [2, 6, 6, 10, 10, 11, 36, 14, 40, 8, 8, 12, 15, 9, 13, 24])

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Mật khẩu lưu trên máy: Windows mã hoá bằng DPAPI (chỉ tài khoản Windows này giải được)
# ---------------------------------------------------------------------------

def _dpapi(data, encrypt):
    import ctypes
    from ctypes import wintypes

    class BLOB(ctypes.Structure):
        _fields_ = [("cbData", wintypes.DWORD), ("pbData", ctypes.POINTER(ctypes.c_char))]

    buf = ctypes.create_string_buffer(data, len(data))
    blob_in = BLOB(len(data), ctypes.cast(buf, ctypes.POINTER(ctypes.c_char)))
    blob_out = BLOB()
    crypt32 = ctypes.windll.crypt32
    if encrypt:
        ok = crypt32.CryptProtectData(ctypes.byref(blob_in), "TaiHoaDon", None, None, None, 0, ctypes.byref(blob_out))
    else:
        ok = crypt32.CryptUnprotectData(ctypes.byref(blob_in), None, None, None, None, 0, ctypes.byref(blob_out))
    if not ok:
        raise OSError("DPAPI lỗi")
    try:
        return ctypes.string_at(blob_out.pbData, blob_out.cbData)
    finally:
        ctypes.windll.kernel32.LocalFree(blob_out.pbData)


_FERNET = None  # bản web trên máy chủ: mã hoá mật khẩu cổng thuế bằng khoá bí mật của máy chủ


def set_secret_key(path=None, key=None):
    """Bật mã hoá Fernet (AES) cho mật khẩu cổng thuế. Khoá lấy từ biến môi trường hoặc file (tự tạo, quyền 600)."""
    global _FERNET
    try:
        from cryptography.fernet import Fernet
    except ImportError:
        raise SystemExit("Bản web cần thư viện cryptography: pip install cryptography")
    if not key:
        if path and os.path.exists(path):
            with open(path, "rb") as f:
                key = f.read().strip()
        else:
            key = Fernet.generate_key()
            os.makedirs(os.path.dirname(path), exist_ok=True)
            fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            with os.fdopen(fd, "wb") as f:
                f.write(key)
    _FERNET = Fernet(key if isinstance(key, bytes) else key.encode())


def protect(text):
    import base64
    raw = text.encode("utf-8")
    if _FERNET is not None:
        return "fernet:" + _FERNET.encrypt(raw).decode()
    if os.name == "nt":
        return "dpapi:" + base64.b64encode(_dpapi(raw, True)).decode()
    return "b64:" + base64.b64encode(raw).decode()


def unprotect(value):
    import base64
    if not value:
        return ""
    kind, _, data = value.partition(":")
    if kind == "fernet":
        if _FERNET is None:
            raise PortalError("Mật khẩu được mã hoá bằng khoá máy chủ – chạy bản web để dùng")
        try:
            return _FERNET.decrypt(data.encode()).decode("utf-8")
        except Exception:
            raise PortalError("Không giải mã được mật khẩu (sai khoá bí mật máy chủ?) – nhập lại mật khẩu")
    raw = base64.b64decode(data)
    if kind == "dpapi":
        raw = _dpapi(raw, False)
    return raw.decode("utf-8")


def _load_json(path, default):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return default


def _save_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


# ---------------------------------------------------------------------------
# Danh sách doanh nghiệp
# ---------------------------------------------------------------------------

class Store:
    def __init__(self, path):
        self.path = path
        self.lock = threading.RLock()
        self.items = _load_json(path, [])

    def save(self):
        with self.lock:
            _save_json(self.path, self.items)

    def get(self, mst):
        with self.lock:
            return next((c for c in self.items if c["mst"] == mst), None)

    def upsert(self, mst, ten=None, password=None, **fields):
        mst = re.sub(r"\s+", "", str(mst or ""))
        if not re.fullmatch(r"\d{10}(-\d{3})?", mst):
            raise ValueError("MST không hợp lệ: %r" % mst)
        with self.lock:
            c = self.get(mst)
            if not c:
                c = {"mst": mst, "ten": "", "pw": "", "ghichu": "", "vao": True, "ra": True, "an": False,
                     "loi": "", "lich_su": {}, "so_hd": {}, "tao": date.today().strftime("%d/%m/%y")}
                self.items.append(c)
            if ten is not None:
                c["ten"] = ten.strip()
            if password:
                c["pw"] = protect(password)
                c["loi"] = ""
            for k in ("ghichu", "vao", "ra", "an", "loi"):
                if k in fields and fields[k] is not None:
                    c[k] = fields[k]
            self.save()
            return c

    def delete(self, mst):
        with self.lock:
            self.items = [c for c in self.items if c["mst"] != mst]
            self.save()

    def update(self, mst, fn):
        with self.lock:
            c = self.get(mst)
            if c:
                fn(c)
                self.save()

    def public(self):
        with self.lock:
            out = []
            for c in self.items:
                d = {k: v for k, v in c.items() if k != "pw"}
                d["co_mk"] = bool(c.get("pw"))
                out.append(d)
            return out

    def import_text(self, text):
        """Mỗi dòng: MST <tab hoặc |> Tên <tab hoặc |> Mật khẩu (dán thẳng từ Excel)."""
        added, errors = 0, []
        for n, line in enumerate((text or "").splitlines(), 1):
            if not line.strip():
                continue
            parts = [p.strip() for p in re.split(r"\t|\|", line)]
            try:
                self.upsert(parts[0], parts[1] if len(parts) > 1 and parts[1] else None,
                            parts[2] if len(parts) > 2 else None)
                added += 1
            except ValueError as e:
                errors.append("Dòng %d: %s" % (n, e))
        return added, errors

    def reencrypt(self):
        """Chuyển mật khẩu đang lưu dạng cũ (base64 / DPAPI) sang cách mã hoá hiện tại (khoá máy chủ)."""
        n = 0
        with self.lock:
            for c in self.items:
                pw = c.get("pw") or ""
                if pw and not pw.startswith("fernet:") and _FERNET is not None:
                    try:
                        c["pw"] = protect(unprotect(pw))
                        n += 1
                    except Exception:
                        c["loi"] = "Không đọc được mật khẩu cũ – nhập lại mật khẩu"
            if n:
                self.save()
        return n


# ---------------------------------------------------------------------------
# Người dùng phần mềm (bản web): quản trị (admin) xem tất cả; nhân viên chỉ thấy các DN được giao
# ---------------------------------------------------------------------------

ROLES = {"admin": "Quản trị", "staff": "Nhân viên"}


class UserStore:
    ITER = 200000

    def __init__(self, path):
        self.path = path
        self.lock = threading.RLock()
        self.items = _load_json(path, [])

    def save(self):
        with self.lock:
            _save_json(self.path, self.items)
            try:
                os.chmod(self.path, 0o600)
            except OSError:
                pass

    @classmethod
    def hash(cls, password, salt=None, iterations=None):
        import hashlib
        salt = salt or secrets.token_hex(16)
        iterations = iterations or cls.ITER
        dk = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode(), iterations)
        return "pbkdf2_sha256$%d$%s$%s" % (iterations, salt, dk.hex())

    @classmethod
    def check(cls, password, stored):
        try:
            _, it, salt, _h = stored.split("$")
            return secrets.compare_digest(cls.hash(password, salt, int(it)), stored)
        except (ValueError, AttributeError):
            cls.hash(password or "", "x" * 32)  # tốn thời gian như nhau để không lộ tên đăng nhập có tồn tại
            return False

    def get(self, username):
        username = (username or "").strip().lower()
        with self.lock:
            return next((u for u in self.items if u["username"] == username), None)

    def count(self):
        with self.lock:
            return len(self.items)

    def authenticate(self, username, password):
        u = self.get(username)
        ok = self.check(password or "", u["pw"] if u else "")
        return u if ok and u.get("active", True) else None

    def _active_admins(self):
        return [u for u in self.items if u["role"] == "admin" and u.get("active", True)]

    def upsert(self, username, ten=None, role=None, password=None, msts=None, active=None, create=False):
        username = (username or "").strip().lower()
        if not re.fullmatch(r"[a-z0-9][a-z0-9._@-]{2,39}", username):
            raise ValueError("Tên đăng nhập 3–40 ký tự: chữ thường không dấu, số, . _ - @ (bắt đầu bằng chữ/số)")
        if password is not None and password != "" and len(password) < 8:
            raise ValueError("Mật khẩu phần mềm tối thiểu 8 ký tự")
        if role is not None and role not in ROLES:
            raise ValueError("Vai trò không hợp lệ")
        with self.lock:
            u = self.get(username)
            if u and create:
                raise ValueError("Tên đăng nhập %s đã tồn tại" % username)
            is_new = u is None
            if is_new:
                if not password:
                    raise ValueError("Nhập mật khẩu cho người dùng mới")
                u = {"username": username, "ten": "", "role": "staff", "pw": "", "msts": [], "active": True,
                     "tao": date.today().strftime("%d/%m/%y")}
                self.items.append(u)
            old = dict(u)
            if ten is not None:
                u["ten"] = str(ten).strip()[:100]
            if role is not None:
                u["role"] = role
            if active is not None:
                u["active"] = bool(active)
            if msts is not None:
                u["msts"] = sorted({str(m).strip() for m in msts if str(m).strip()})
            if password:
                u["pw"] = self.hash(password)
                u["doi_mk"] = int(time.time())
            if not self._active_admins():
                if is_new:
                    self.items.remove(u)
                else:
                    u.clear()
                    u.update(old)
                raise ValueError("Phải còn ít nhất một quản trị đang hoạt động")
            self.save()
            return u

    def delete(self, username):
        with self.lock:
            u = self.get(username)
            if not u:
                return
            rest = [x for x in self.items if x is not u]
            if not any(x["role"] == "admin" and x.get("active", True) for x in rest):
                raise ValueError("Không xoá được quản trị cuối cùng")
            self.items = rest
            self.save()

    def remove_mst(self, mst):
        with self.lock:
            for u in self.items:
                if mst in u.get("msts", []):
                    u["msts"].remove(mst)
            self.save()

    def public(self):
        with self.lock:
            return [{k: v for k, v in u.items() if k != "pw"} for u in self.items]

    @staticmethod
    def can(user, mst):
        return bool(user) and (user["role"] == "admin" or mst in (user.get("msts") or []))


# ---------------------------------------------------------------------------
# Captcha tự học: cổng dùng captcha SVG, mỗi ký tự là một <path> có màu tô với hình dạng cố định.
# Mỗi lần người dùng gõ đúng captcha, phần mềm lưu hình từng ký tự; lần sau tự nhận ra.
# ---------------------------------------------------------------------------

_PATH_RE = re.compile(r"<path\b([^>]*)/?>", re.I)
_ATTR_RE = re.compile(r'([\w:-]+)\s*=\s*"([^"]*)"')
_NUM_RE = re.compile(r"-?\d*\.?\d+(?:e-?\d+)?", re.I)


# Chữ ký M/Q/Z của 30 ký tự trong font captcha của cổng (đã kiểm chứng trên captcha thật; captcha không dùng
# I, L, O, U, 0, 1). Bảng giống bảng trong công cụ cộng đồng hddt-downloader-windows. Ký tự lạ vẫn tự học thêm.
CAPTCHA_SIGNATURES = {
    "MQQQQQZMQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQZMQQZ": "A",
    "MQQQQQQQQQZMQQQQQQZMQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQZMQQQQQQQQZMQQQQQQQQZ": "B",
    "MQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQZ": "C",
    "MQQQQQQQQZMQQQQQQQQQQZMQQQQQQQQQQQQQQQZMQQQQQQQZ": "D",
    "MQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "E",
    "MQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQZ": "F",
    "MQQQQQQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "G",
    "MQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQZ": "H",
    "MQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQZ": "J",
    "MQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQZ": "K",
    "MQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQZ": "M",
    "MQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQZ": "N",
    "MQQQQQQZMQQQQQQQQQQZMQQQQQQQQQQQQQQQZMQQQQQQQQZ": "P",
    "MQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQZ": "Q",
    "MQQQQQQZMQQQQQQQQQQQQZMQQQQQQQQQQQQQQQZMQQQQQQQQZ": "R",
    "MQQQQQQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "S",
    "MQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQZ": "T",
    "MQQQQQQQQQQZMQQQQQQQQQQQQQQQQZ": "V",
    "MQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "W",
    "MQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQZ": "X",
    "MQQQQQQQQQZMQQQQQQQQQQQQQZ": "Y",
    "MQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQZ": "Z",
    "MQQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "2",
    "MQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "3",
    "MQQQQZMQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQZMQQQQQZ": "4",
    "MQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQQQQQQQQQZ": "5",
    "MQQQQQQQQQZMQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQZ": "6",
    "MQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQQZ": "7",
    "MQQQQQQQQZMQQQQQQQZMQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQZMQQQQQQQZ": "8",
    "MQQQQQQQQZMQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQQQQQQQQQQZMQQQQQQQQQQQZ": "9",
}


def captcha_glyphs(svg):
    """Trả về [(chữ lệnh, toạ độ tương đối, x trái)] của các ký tự, xếp từ trái sang phải."""
    glyphs = []
    for m in _PATH_RE.finditer(svg or ""):
        attrs = dict(_ATTR_RE.findall(m.group(1)))
        d = attrs.get("d", "")
        letters = "".join(re.findall(r"[A-Za-z]", d)).upper()
        if "Z" not in letters:
            continue  # đường nhiễu là nét cong hở; ký tự là hình khép kín (có Z), dù tô màu hay chỉ vẽ viền
        letters = re.sub(r"[^MQZ]", "", letters)  # bỏ L (nét rung ngẫu nhiên), giữ chữ ký M/Q/Z của ký tự
        nums = [float(x) for x in _NUM_RE.findall(d)]
        if len(nums) < 4:
            continue
        xs, ys = nums[0::2], nums[1::2]
        minx, miny = min(xs), min(ys)
        rel = []
        for i, v in enumerate(nums):
            rel.append(round(v - (minx if i % 2 == 0 else miny), 1))
        glyphs.append((letters, rel, minx))
    glyphs.sort(key=lambda g: g[2])
    return glyphs


def _dist(a, b):
    if len(a) != len(b):
        return float("inf")
    return sum(abs(x - y) for x, y in zip(a, b)) / max(len(a), 1)


class CaptchaSolver:
    THRESHOLD = 1.5

    def __init__(self, path):
        self.path = path
        self.lock = threading.Lock()
        self.templates = _load_json(path, {})  # chữ lệnh → [[ký tự, toạ độ], ...]

    def count(self):
        with self.lock:
            return sum(len(v) for v in self.templates.values()) + len(CAPTCHA_SIGNATURES)

    def chars(self):
        with self.lock:
            return sorted({t[0] for v in self.templates.values() for t in v} | set(CAPTCHA_SIGNATURES.values()))

    def _match(self, letters, rel):
        cands = self.templates.get(letters, [])
        if cands and len({c for c, _ in cands}) == 1:
            return cands[0][0]  # chuỗi lệnh vẽ chỉ ứng với một ký tự → chắc chắn
        best, best_d = None, float("inf")
        for ch, tpl in cands:
            dd = _dist(rel, tpl)
            if dd < best_d:
                best, best_d = ch, dd
        return best if best_d <= self.THRESHOLD else None

    def solve(self, svg):
        glyphs = captcha_glyphs(svg)
        if not glyphs:
            return None
        out = []
        with self.lock:
            for letters, rel, _ in glyphs:
                ch = self._match(letters, rel) or CAPTCHA_SIGNATURES.get(letters)
                if ch is None:
                    return None
                out.append(ch)
        return "".join(out)

    def learn(self, svg, answer):
        glyphs = captcha_glyphs(svg)
        if not answer or len(glyphs) != len(answer):
            return False
        with self.lock:
            for (letters, rel, _), ch in zip(glyphs, answer):
                lst = self.templates.setdefault(letters, [])
                if not any(c == ch and _dist(rel, t) <= 0.3 for c, t in lst):
                    lst.append([ch, rel])
            _save_json(self.path, self.templates)
        return True

    def save_sample(self, svg, answer):
        """Lưu vài captcha gần nhất (ảnh + đáp án) vào _cau-hinh/captcha-mau/ để chẩn đoán khi không tự học được."""
        try:
            folder = os.path.join(os.path.dirname(self.path), "captcha-mau")
            os.makedirs(folder, exist_ok=True)
            files = sorted(os.listdir(folder))
            for old in files[:-19]:
                os.remove(os.path.join(folder, old))
            with open(os.path.join(folder, "%s_%s.svg" % (datetime.now().strftime("%Y%m%d%H%M%S%f"), safe_name(answer))),
                      "w", encoding="utf-8") as fh:
                fh.write(svg)
        except OSError:
            pass

    def forget(self, svg):
        """Captcha tự giải bị cổng báo sai → bỏ các mẫu đã dùng để lần sau hỏi lại người dùng."""
        with self.lock:
            for letters, _, _ in captcha_glyphs(svg):
                self.templates.pop(letters, None)
            _save_json(self.path, self.templates)


# ---------------------------------------------------------------------------
# Tác vụ chạy nền (đồng bộ hàng loạt / test đăng nhập); giao diện hỏi tiến độ
# ---------------------------------------------------------------------------

class Cancelled(Exception):
    pass


class Job:
    def __init__(self):
        self.lock = threading.Lock()
        self.running = False
        self.cancel = False
        self.log = []
        self.done = 0
        self.total = 0
        self.current = None
        self.results = []
        self.files = []
        self.need_captcha = None
        self._answer = None
        self._event = threading.Event()
        self.started = time.time()
        self.finished = None
        self.auth = ""          # trạng thái chứng thực tài khoản
        self.found = {}         # {'purchase': n, 'sold': n}
        self.rows = []          # kết quả đồng bộ từng hoá đơn
        self.phase = "chờ đồng bộ"
        self.unattended = False  # đồng bộ tự động theo lịch: không hỏi captcha
        self.owner = ""

    def say(self, msg):
        with self.lock:
            self.log.append("%s  %s" % (datetime.now().strftime("%H:%M:%S"), msg))
            self.log = self.log[-300:]

    def ask_captcha(self, company, svg, timeout=600):
        with self.lock:
            self._answer = None
            self._event.clear()
            self.need_captcha = {"mst": company["mst"], "ten": company.get("ten"), "svg": svg}
        self._event.wait(timeout)
        with self.lock:
            self.need_captcha = None
            ans = self._answer
        if self.cancel or not ans:
            raise Cancelled()
        return ans

    def answer(self, text):
        with self.lock:
            self._answer = (text or "").strip()
        self._event.set()

    def stop(self):
        self.cancel = True
        self._event.set()

    def snapshot(self):
        with self.lock:
            return {"running": self.running, "log": list(self.log), "done": self.done, "total": self.total,
                    "current": self.current, "results": list(self.results), "files": list(self.files),
                    "need_captcha": self.need_captcha, "auth": self.auth, "found": dict(self.found),
                    "rows": self.rows[-1500:], "phase": self.phase,
                    "elapsed": round((self.finished or time.time()) - self.started, 1)}


def login_company(job, client, company, solver):
    """Đăng nhập một doanh nghiệp: tự giải captcha nếu đã học, không thì hỏi người dùng.
    Sai mật khẩu → dừng ngay (không thử lại để tránh bị khoá tài khoản)."""
    if client.token_valid() and client.username == company["mst"]:
        job.auth = "thành công (dùng lại phiên đăng nhập)"
        job.say("%s: dùng lại phiên đăng nhập còn hạn" % company["mst"])
        return
    password = unprotect(company.get("pw"))
    if not password:
        raise PortalError("Chưa nhập mật khẩu hoadondientu.gdt.gov.vn")
    job.auth = "đang đăng nhập"
    auto_failed = False
    for _ in range(4):
        cap = client.get_captcha()
        text = None if auto_failed and not job.unattended else solver.solve(cap["content"])
        auto = text is not None
        if not auto and job.unattended:  # đồng bộ tự động: không có người gõ captcha → thử captcha khác
            job.say("%s: không tự giải được captcha, thử captcha khác" % company["mst"])
            time.sleep(1)
            continue
        if not auto:
            glyphs = captcha_glyphs(cap["content"])
            unknown = sum(1 for g in glyphs if g[0] not in solver.templates)
            job.say("%s: cần nhập captcha (ảnh có %d ký tự, %d ký tự chưa học)" % (company["mst"], len(glyphs), unknown))
            job.auth = "chờ nhập captcha"
            text = job.ask_captcha(company, cap["content"]).upper()  # captcha chỉ gồm A-Z, 0-9
            solver.save_sample(cap["content"], text)
        try:
            client.login(company["mst"], password, text, cap["key"])
        except PortalError as e:
            if "captcha" in str(e).lower():
                if auto:
                    solver.forget(cap["content"])
                    auto_failed = True  # tự giải sai → lần sau hỏi người dùng
                job.say("%s: captcha sai, thử lại" % company["mst"])
                continue
            raise
        if not auto and not solver.learn(cap["content"], text):
            job.say("Không học được captcha này (số ký tự trong ảnh khác số ký tự đã gõ) – đã lưu mẫu để kiểm tra")
        job.auth = "thành công" + (" (tự giải captcha)" if auto else "")
        job.say("%s: đăng nhập thành công%s" % (company["mst"], " (tự giải captcha)" if auto else ""))
        return
    raise PortalError("Nhập sai captcha quá nhiều lần")


def sync_kind(job, client, company, kind, start, end, include_mtt, want_xml, out_root):
    mst = company["mst"]
    label = "mua-vao" if kind == "purchase" else "ban-ra"
    period = "%s_%s" % (start.strftime("%Y%m%d"), end.strftime("%Y%m%d"))
    base = os.path.join(out_root, safe_name(mst))
    folder = os.path.join(base, "%s_%s" % (label, period))
    known = _load_json(index_file(out_root, mst), {}).get(kind, {})

    job.phase = "đang tra danh sách hoá đơn"
    invoices = client.list_invoices(kind, start, end, include_mtt, progress=lambda m: job.say("%s: %s" % (mst, m)))
    loai = "Mua vào" if kind == "purchase" else "Bán ra"
    with job.lock:
        job.found[kind] = len(invoices)
    job.phase = "đang đồng bộ hoá đơn"
    new, changes, todo, status, detail_todo = 0, [], [], {}, []
    xml_dir = os.path.join(folder, "xml")
    for inv in invoices:
        key = invoice_key(inv)
        old = known.get(key)
        tthai = _num(inv.get("tthai"))
        if old is None:
            new += 1
            status[key] = "Mới"
        elif old.get("tthai") != tthai:
            status[key] = "Đổi trạng thái"
            changes.append(_inv_head(inv)[:5] + [TTHAI_LABELS.get(old.get("tthai"), old.get("tthai")),
                                                 TTHAI_LABELS.get(tthai, tthai)])
            job.say("%s: HĐ %s/%s đổi trạng thái → %s" % (mst, inv.get("khhdon"), inv.get("shdon"),
                                                          TTHAI_LABELS.get(tthai, tthai)))
        else:
            status[key] = "Đã có"
        if want_xml:
            existing = os.path.join(xml_dir, invoice_basename(inv) + ".xml")
            if (old or {}).get("no_xml") and status[key] != "Đổi trạng thái":
                inv["_noxml"] = True  # đã biết cổng không có XML gốc của HĐ này
                if old.get("detail") and os.path.exists(os.path.join(base, old["detail"])):
                    inv["_detail"] = old["detail"]
                else:
                    detail_todo.append(inv)  # bản cũ chưa lấy dữ liệu chi tiết → lấy để có hàng hoá và PDF của thuế
            elif os.path.exists(existing) and status[key] != "Đổi trạng thái":
                inv["_xml"] = os.path.relpath(existing, folder)
                inv["_xmlinfo"] = parse_invoice_xml(existing)
            else:
                todo.append(inv)
    errors = {}
    with job.lock:
        job.total += len(todo)

    def fetch_detail(inv):
        """HĐ không có XML gốc: lấy dữ liệu 'Xem hoá đơn' của cổng (hàng hoá, thuế suất) → PDF của thuế."""
        try:
            detail = client.get_detail(inv)
            rel = os.path.join(os.path.relpath(xml_dir, base), invoice_basename(inv) + ".json")
            os.makedirs(xml_dir, exist_ok=True)
            _save_json(os.path.join(base, rel), detail)
            inv["_detail"] = rel
        except PortalError as e:
            job.say("%s: không lấy được chi tiết HĐ %s/%s: %s" % (mst, inv.get("khhdon"), inv.get("shdon"), e))

    def fetch(inv):
        if job.cancel:
            return
        try:
            path = save_xml(client.export_xml(inv), xml_dir, invoice_basename(inv))
            inv["_xml"] = os.path.relpath(path, folder)
            inv["_xmlinfo"] = parse_invoice_xml(path)
        except PortalError as e:
            if NO_XML_MSG in _plain(str(e)):
                inv["_noxml"] = True  # HĐ không mã: cổng chỉ có dữ liệu tổng hợp, không có file XML gốc
                fetch_detail(inv)
            else:
                inv["_xml"] = "Lỗi: %s" % e
                errors[invoice_key(inv)] = str(e)
        with job.lock:
            job.done += 1

    if todo or detail_todo:
        # Tải XML song song (vừa phải để cổng không chặn).
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(max_workers=XML_WORKERS) as pool:
            list(pool.map(fetch, todo))
            list(pool.map(lambda i: None if job.cancel else fetch_detail(i), detail_todo))
    if job.cancel:
        raise Cancelled()
    if any("hết hạn" in m for m in errors.values()):
        raise PortalError("Phiên đăng nhập đã hết hạn, vui lòng đăng nhập lại.")
    for key, msg in errors.items():
        job.say("%s: không tải được XML %s: %s" % (mst, key.split("|", 2)[-1].replace("|", "/"), msg))

    rows = []
    for inv in invoices:
        key = invoice_key(inv)
        d = invoice_date(inv)
        rows.append({"mst": mst, "tthai": TTHAI_NIBOT.get(_num(inv.get("tthai")), ""), "loai": loai, "ngay": d.strftime("%d/%m/%Y") if d else "", "mau": str(inv.get("khmshdon") or ""),
                     "kh": inv.get("khhdon") or "", "so": str(inv.get("shdon") or ""), "dongbo": status[key],
                     "ketqua": ("Lỗi XML: " + errors[key]) if key in errors else
                     ("OK (HĐ không có XML gốc)" if inv.get("_noxml") else "OK")})
        old = known.get(key)
        tthai = _num(inv.get("tthai"))
        entry = dict(old or {})
        inv["_kind"] = kind
        entry.update(tthai=tthai, tdlap=inv.get("tdlap"),
                     inv={k: inv.get(k) for k in INV_FIELDS if inv.get(k) is not None})
        if inv.get("_noxml"):
            entry["no_xml"] = True
        if inv.get("_detail"):
            entry["detail"] = inv["_detail"]
            if not entry.get("mat_hang"):
                d = _load_json(os.path.join(base, inv["_detail"]), {})
                entry["mat_hang"] = "; ".join(str(it.get("ten") or "") for it in (d.get("hdhhdvu") or [])
                                              if it.get("ten"))[:500]
        if inv.get("_xmlinfo") is not None:
            entry["xml"] = os.path.relpath(os.path.join(folder, inv["_xml"]), base)
            names = [it["name"] for it in inv["_xmlinfo"]["items"] if it["name"]]
            entry["mat_hang"] = "; ".join(names)[:500]
            entry["tra_cuu"] = lookup_info(inv["_xmlinfo"])
        known[key] = entry
    save_index_kind(out_root, mst, kind, known)
    if getattr(job, "want_pdf", False):
        by_key = {invoice_key(i): r for i, r in zip(invoices, rows)}
        base_dir = company_dir(out_root, mst)
        todo = [k for k in by_key if can_fetch_pdf(known[k].get("tra_cuu"))[0] and not (
            known[k].get("pdf") and os.path.exists(os.path.join(base_dir, known[k]["pdf"])))]
        if todo:
            job.say("%s: tải PDF gốc %d hoá đơn" % (mst, len(todo)))
        for k in todo:
            if job.cancel:
                raise Cancelled()
            try:
                fetch_original_pdf(out_root, mst, kind, k)
                by_key[k]["ketqua"] += " + PDF gốc"
            except (PortalError, ValueError) as e:
                by_key[k]["ketqua"] += " (PDF gốc: lỗi)"
                job.say("%s: không tải được PDF gốc %s: %s" % (mst, k.split("|", 2)[-1].replace("|", "/"), e))
            time.sleep(0.3)
    with job.lock:
        job.rows.extend(rows)

    paths = write_reports(invoices, folder, "bang-ke-%s_%s" % (label, period), csv_only=True)
    try:
        xlsx = os.path.join(folder, "%s_%s_%s.xlsx" % ("MUA_VAO" if kind == "purchase" else "BAN_RA", mst, period))
        paths.append(write_nibot_workbook(xlsx, kind, invoices, mst, company.get("ten"), start, end, changes))
    except ImportError:
        job.say("Chưa cài openpyxl nên chỉ xuất CSV (chạy: pip install openpyxl)")
    with job.lock:
        job.files.extend(os.path.abspath(p) for p in paths)
        job.results.append({"mst": mst, "ten": company.get("ten"), "loai": label, "so_hd": len(invoices),
                            "moi": new, "doi": len(changes), "thu_muc": os.path.abspath(folder)})
    job.say("%s: %s %d HĐ (%d mới, %d đổi trạng thái)" % (mst, label, len(invoices), new, len(changes)))
    return len(invoices), new, len(known)


_BUSY_LOCK = threading.Lock()
_BUSY = {}  # MST → Job đang đồng bộ (hai người không đồng bộ cùng một DN một lúc)


def claim_mst(mst, job):
    with _BUSY_LOCK:
        if _BUSY.get(mst) not in (None, job):
            return False
        _BUSY[mst] = job
        return True


def release_mst(mst, job):
    with _BUSY_LOCK:
        if _BUSY.get(mst) is job:
            del _BUSY[mst]


def busy_msts():
    with _BUSY_LOCK:
        return set(_BUSY)


def run_batch(job, store, solver, msts, kinds, start, end, include_mtt, want_xml, out_root,
              test_only=False, client_factory=None, clients=None):
    """clients: dict MST → HoaDonClient để giữ phiên đăng nhập giữa các lần đồng bộ (không phải nhập lại captcha)."""
    client_factory = client_factory or HoaDonClient
    clients = {} if clients is None else clients
    try:
        for i, mst in enumerate(msts):
            if job.cancel:
                break
            job.dn = [i + 1, len(msts)]
            company = store.get(mst)
            if not company:
                continue
            if not claim_mst(mst, job):
                job.say("%s: LỖI doanh nghiệp đang được người khác đồng bộ – bỏ qua, thử lại sau" % mst)
                continue
            with job.lock:
                job.current = mst
            client = clients.get(mst)
            if client is None:
                client = clients[mst] = client_factory()
            try:
                job.phase = "đang chứng thực tài khoản"
                login_company(job, client, company, solver)
                store.update(mst, lambda c: c.update(loi=""))
                if test_only:
                    continue
                for kind in kinds:
                    if not company.get("vao" if kind == "purchase" else "ra", True):
                        continue
                    try:
                        total, new, known = sync_kind(job, client, company, kind, start, end, include_mtt,
                                                      want_xml, out_root)
                    except PortalError as e:
                        if "hết hạn" not in str(e):
                            raise
                        job.say("%s: phiên hết hạn, đăng nhập lại" % mst)
                        client.token = None
                        login_company(job, client, company, solver)
                        total, new, known = sync_kind(job, client, company, kind, start, end, include_mtt,
                                                      want_xml, out_root)
                    stamp = datetime.now().strftime("%d.%m.%y %H:%M")

                    def upd(c, kind=kind, new=new, known=known, stamp=stamp):
                        c.setdefault("lich_su", {})[kind] = {"moi": new, "luc": stamp}
                        c.setdefault("so_hd", {})[kind] = known
                    store.update(mst, upd)
            except Cancelled:
                job.say("Đã dừng theo yêu cầu.")
                break
            except Exception as e:  # ghi lỗi của DN này rồi chuyển sang DN tiếp theo
                msg = str(e)
                if not job.auth.startswith("thành công"):
                    job.auth = "thất bại – " + msg
                job.say("%s: LỖI %s" % (mst, msg))
                store.update(mst, lambda c: c.update(loi=msg))
            finally:
                release_mst(mst, job)
        job.say("Hoàn tất.")
        job.phase = "đồng bộ xong" if not job.cancel else "đã dừng"
    finally:
        with job.lock:
            job.finished = time.time()
            if job.phase not in ("đồng bộ xong", "đã dừng"):
                job.phase = "lỗi"
            job.running = False
            job.current = None
            job.need_captcha = None


# ---------------------------------------------------------------------------
# Kho hoá đơn đã tải (HoaDon/<MST>/_chi-muc.json): xem, lọc, ghi chú, duyệt, kết xuất lại
# ---------------------------------------------------------------------------

INV_FIELDS = ("khmshdon", "khhdon", "shdon", "tdlap", "nky", "nbmst", "nbten", "nbdchi", "nmmst", "nmten", "nmdchi",
              "thtttoan", "tgtcthue", "tgtthue", "ttcktmai", "tgtphi", "tgtttbso", "dvtte", "tthai", "ttxly",
              "mhdon", "thttltsuat", "_prefix", "_nguon", "_kind")
USER_FIELDS = ("note", "duyet", "dv", "pdf")
DUYET_OPTIONS = ("Chờ duyệt", "Đã duyệt", "Không duyệt")
_INDEX_LOCK = threading.Lock()


def company_dir(out_root, mst):
    return os.path.join(out_root, safe_name(mst))


def index_file(out_root, mst):
    return os.path.join(company_dir(out_root, mst), "_chi-muc.json")


def save_index_kind(out_root, mst, kind, entries):
    """Ghi lại một loại HĐ của chỉ mục, giữ ghi chú/duyệt do người dùng sửa trong lúc đang đồng bộ."""
    path = index_file(out_root, mst)
    with _INDEX_LOCK:
        cur = _load_json(path, {})
        old = cur.get(kind, {})
        for key, e in entries.items():
            for f in USER_FIELDS:
                if f in old.get(key, {}):
                    e[f] = old[key][f]
        cur[kind] = entries
        _save_json(path, cur)


def update_invoice(out_root, mst, kind, key, **fields):
    path = index_file(out_root, mst)
    with _INDEX_LOCK:
        cur = _load_json(path, {})
        e = cur.get(kind, {}).get(key)
        if e is None:
            raise ValueError("Không tìm thấy hoá đơn")
        for f in USER_FIELDS:
            if fields.get(f) is not None:
                e[f] = fields[f]
        _save_json(path, cur)


def attach_pdf(out_root, mst, kind, key, data):
    """Gắn file PDF gốc (người dùng tải từ trang tra cứu) vào hoá đơn."""
    if not data.startswith(b"%PDF"):
        raise ValueError("File không phải PDF")
    base = company_dir(out_root, mst)
    path = index_file(out_root, mst)
    with _INDEX_LOCK:
        cur = _load_json(path, {})
        e = cur.get(kind, {}).get(key)
        if e is None:
            raise ValueError("Không tìm thấy hoá đơn")
        if e.get("xml"):
            rel = os.path.splitext(e["xml"])[0] + ".pdf"
        else:
            rel = os.path.join("pdf", invoice_basename(e.get("inv") or {}) + ".pdf")
        os.makedirs(os.path.dirname(os.path.join(base, rel)), exist_ok=True)
        with open(os.path.join(base, rel), "wb") as fh:
            fh.write(data)
        e["pdf"] = rel
        _save_json(path, cur)
    return rel


def _pdf_text(data):
    import pypdf
    try:
        r = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in r.pages[:3])
    except Exception as e:
        raise ValueError("Không đọc được PDF: %s" % e)


def match_pdf_text(text, entries_by_kind):
    """Tìm hoá đơn trong kho khớp với nội dung PDF: cùng ký hiệu, MST người bán/mua có trong PDF, và số hoá đơn có
    trong PDF. Chấm điểm số hoá đơn: đứng sau nhãn "Số:"/"Số (No.):" hoặc trước "(No)" = 4; in dạng 0000xxxx = 3; số dài ≥ 6
    chữ số = 2; số ngắn bất kỳ = 1 (dễ trùng số nhà, số tiền). Chỉ trả về hoá đơn có điểm cao nhất duy nhất."""
    flat = re.sub(r"\s+", " ", text or "")
    # Số có dấu chấm/phẩy ngăn cách (số tiền 308.448) không được tách thành 308 và 448.
    tokens = [t for t in re.findall(r"\d[\d.,]*\d|\d", flat) if t.isdigit() and len(t) <= 10]
    score = {}
    for tkn in tokens:
        n = int(tkn)
        s_ = 3 if (tkn.startswith("0") and len(tkn) >= 7) else (2 if len(tkn) >= 6 else 1)
        score[n] = max(score.get(n, 0), s_)
    labeled = (
        # "Số: 123", "Số (No.): 123" – số không được dính chữ phía sau (vd "Số: 1C25TLV" không phải số 1)
        r"Số\s*(?:\(No\.?\)\s*:?|:)\s*0*(\d{1,10})(?![\d.,]\d)(?![A-Za-z])",
        r"(?<![\d.,])(\d{1,10})\s*\(No\)",
        # số đứng ngay sau ký hiệu: "1C25TQP 106", "Số: 1C25TLV 00004279", "Ký hiệu: 1C26TAA Số: 1614"
        r"(?<![A-Z0-9])[12]?[CK]\d{2}[A-Z]{2,3}\s+(?:Số\s*(?:\([^)]{0,8}\))?\s*:?\s*)?0*(\d{1,10})(?![\d.,]\d)(?![A-Za-z])",
        # chữ PDF bị đảo thứ tự: "00000165:Số"
        r"(?<![\d.,])0*(\d{1,10})\s*:\s*Số(?!\s*(?:tài|TK|điện))",
    )
    for pat in labeled:
        for m in re.finditer(pat, flat):
            score[int(m.group(1))] = 4  # đúng vị trí số hoá đơn – hơn số nhắc tới trong nội dung (vd HĐ bị điều chỉnh)
    msts = set(re.findall(r"(?<!\d)(\d{10}(?:-\d{3})?)(?!\d)", flat))
    msts |= {re.sub(r"\s", "", m) for m in re.findall(r"(?<!\d)((?:\d ){9}\d)(?!\d)", flat)}
    msts |= {m.split("-")[0] for m in msts}
    # Ký hiệu 1C25TAA; PDF đảo thứ tự chữ có thể in thành C25TAA1
    khs = set(re.findall(r"(?<![A-Z0-9])[12]?([CK]\d{2}[A-Z]{2,3})[12]?(?![A-Z0-9])", flat))
    found = []
    for kind, entries in entries_by_kind.items():
        for key in entries:
            nbmst, _, khhdon, shdon = (key.split("|") + ["", "", "", ""])[:4]
            sc = score.get(_num(shdon), 0)
            if str(khhdon)[-6:] not in khs or not sc:
                continue
            inv = entries[key].get("inv") or {}
            parties = {p for p in (str(nbmst), str(inv.get("nmmst") or "")) if p}
            if parties & msts or {p.split("-")[0] for p in parties} & msts:
                found.append((sc, kind, key))
    if not found:
        return []
    best = max(f[0] for f in found)
    return [(k, key) for sc, k, key in found if sc == best]


def import_pdfs(out_root, mst, files):
    """Gắn hàng loạt PDF gốc có sẵn trên máy vào đúng hoá đơn (đọc ký hiệu, số, MST trong PDF)."""
    idx = _load_json(index_file(out_root, mst), {})
    results = []
    for name, data in files:
        try:
            if not data.startswith(b"%PDF"):
                raise ValueError("không phải file PDF")
            text = _pdf_text(data)
            if re.search(r"chưa\s*cấp\s*số", text, re.I):
                raise ValueError("PDF là bản nháp, hoá đơn chưa được cấp số – lấy bản PDF sau khi người bán đã phát hành")
            hits = match_pdf_text(text, {k: idx.get(k, {}) for k in ("purchase", "sold")})
            if not hits:
                raise ValueError("không khớp hoá đơn nào trong kho (đã đồng bộ kỳ này chưa?)")
            if len(hits) > 1:
                raise ValueError("khớp %d hoá đơn, không chắc gắn vào đâu" % len(hits))
            kind, key = hits[0]
            attach_pdf(out_root, mst, kind, key, data)
            nb, _, kh, so = key.split("|")
            results.append({"file": name, "ok": True, "hd": "%s %s – %s" % (
                "Mua vào" if kind == "purchase" else "Bán ra", kh, so)})
        except ValueError as e:
            results.append({"file": name, "ok": False, "loi": str(e)})
    return results


# ---------------------------------------------------------------------------
# PDF bản thể hiện hoá đơn theo dữ liệu của cổng thuế ("PDF của thuế"): dựng từ XML, hoặc từ dữ liệu chi tiết
# trên cổng với hoá đơn không có XML gốc (HĐ không mã, máy tính tiền). Cổng không có API trả PDF.
# ---------------------------------------------------------------------------

INVOICE_TITLES = {"1": "HÓA ĐƠN GIÁ TRỊ GIA TĂNG", "2": "HÓA ĐƠN BÁN HÀNG", "3": "HÓA ĐƠN BÁN TÀI SẢN CÔNG",
                  "4": "HÓA ĐƠN BÁN HÀNG DỰ TRỮ QUỐC GIA", "5": "HÓA ĐƠN KHÁC",
                  "6": "PHIẾU XUẤT KHO KIÊM VẬN CHUYỂN NỘI BỘ"}


def _dget(d, *keys):
    for k in keys:
        v = (d or {}).get(k)
        if v not in (None, "", []):
            return v
    return ""


def _vn_money(v):
    if v in (None, ""):
        return ""
    try:
        f = float(str(v).replace(",", ""))
    except ValueError:
        return str(v)
    s = "{:,.2f}".format(f).rstrip("0").rstrip(".")
    return s.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def _vn_rate(v):
    v = str(v if v is not None else "").strip()
    try:
        f = float(v)
        return "%g%%" % (f * 100 if f < 1 else f)
    except ValueError:
        return v


def invoice_view_from_xml(xml_bytes):
    x = read_invoice_xml(xml_bytes)
    tc, nb, nm, tong = x["ttchung"], x["nb"], x["nm"], x["tong"]
    return {
        "mau": tc.get("KHMSHDon", ""), "kh": tc.get("KHHDon", ""), "so": tc.get("SHDon", ""), "title": tc.get("THDon", ""),
        "ngay": tc.get("NLap", ""), "mccqt": x["mccqt"], "httt": tc.get("HTTToan", ""), "dvtte": tc.get("DVTTe", "VND"),
        "nb": {"ten": nb.get("Ten"), "mst": nb.get("MST"), "dchi": nb.get("DChi"), "sdt": nb.get("SDThoai"),
               "stk": nb.get("STKNHang"), "nh": nb.get("TNHang")},
        "nm": {"ten": nm.get("Ten"), "hoten": nm.get("HVTNMHang"), "mst": nm.get("MST") or nm.get("CCCDan"),
               "dchi": nm.get("DChi"), "stk": nm.get("STKNHang")},
        "items": [{"stt": it["STT"], "ten": it["THHDVu"], "dvt": it["DVTinh"], "sl": it["SLuong"], "dgia": it["DGia"],
                   "ck": it["STCKhau"], "tsuat": it["TSuat"], "thtien": it["ThTien"], "tchat": it["TChat"]}
                  for it in x["items"]],
        "rates": [{"tsuat": r.get("TSuat"), "thtien": r.get("ThTien"), "tthue": r.get("TThue")} for r in x["rates"]],
        "tong": {"cthue": tong.get("TgTCThue"), "thue": tong.get("TgTThue"), "ck": tong.get("TTCKTMai"),
                 "phi": tong.get("TgTPhi"), "tt": tong.get("TgTTTBSo"), "chu": tong.get("TgTTTBChu")},
        "signed": bool(x["so_chu_ky"]), "nky": "", "nguon": "XML hoá đơn",
    }


def invoice_view_from_detail(d):
    """Dữ liệu 'Xem hoá đơn' của cổng (JSON) – dùng cho HĐ không có XML gốc."""
    items = []
    for i, it in enumerate(d.get("hdhhdvu") or [], 1):
        items.append({"stt": _dget(it, "stt") or i, "ten": _dget(it, "ten", "thhdvu"), "dvt": _dget(it, "dvtinh"),
                      "sl": _dget(it, "sluong"), "dgia": _dget(it, "dgia"), "ck": _dget(it, "stckhau"),
                      "tsuat": _dget(it, "tsuat", "ltsuat"), "thtien": _dget(it, "thtien"), "tchat": _dget(it, "tchat")})
    rates = [{"tsuat": _dget(r, "tsuat"), "thtien": _dget(r, "thtien"), "tthue": _dget(r, "tthue")}
             for r in (d.get("thttltsuat") or [])]
    return {
        "mau": str(_dget(d, "khmshdon")), "kh": _dget(d, "khhdon"), "so": str(_dget(d, "shdon")), "title": _dget(d, "thdon"),
        "ngay": _dget(d, "tdlap", "nlap"), "mccqt": _dget(d, "mhdon"), "httt": _dget(d, "thtttoan", "htttoan"),
        "dvtte": _dget(d, "dvtte") or "VND",
        "nb": {"ten": _dget(d, "nbten"), "mst": _dget(d, "nbmst"), "dchi": _dget(d, "nbdchi"), "sdt": _dget(d, "nbsdthoai"),
               "stk": _dget(d, "nbstkhoan"), "nh": _dget(d, "nbtnhang")},
        "nm": {"ten": _dget(d, "nmten"), "hoten": _dget(d, "nmtnmua"), "mst": _dget(d, "nmmst", "nmcccd"),
               "dchi": _dget(d, "nmdchi"), "stk": _dget(d, "nmstkhoan")},
        "items": items, "rates": rates,
        "tong": {"cthue": _dget(d, "tgtcthue"), "thue": _dget(d, "tgtthue"), "ck": _dget(d, "ttcktmai"),
                 "phi": _dget(d, "tgtphi"), "tt": _dget(d, "tgtttbso"), "chu": _dget(d, "tgtttbchu")},
        "signed": bool(_dget(d, "nky")), "nky": _dget(d, "nky"), "nguon": "dữ liệu chi tiết trên cổng thuế",
    }


def _fill_view(v, extra):
    """Thông tin còn trống (ký hiệu, số, ngày, mã CQT, tổng tiền…) lấy bù từ dữ liệu danh sách hoá đơn của cổng."""
    for k, val in extra.items():
        if isinstance(val, dict):
            for kk, vv in val.items():
                if v.setdefault(k, {}).get(kk) in (None, "") and vv not in (None, ""):
                    v[k][kk] = vv
        elif k not in ("items", "rates", "nguon", "signed") and v.get(k) in (None, "") and val not in (None, ""):
            v[k] = val
    return v


def invoice_view_from_summary(inv):
    v = invoice_view_from_detail(inv)
    v["nguon"] = "dữ liệu tổng hợp (chưa có chi tiết hàng hoá – đồng bộ lại để lấy)"
    return v


_FONT = {}


def _vn_fonts():
    """Font có dấu tiếng Việt cho reportlab: thư mục fonts/ kèm phần mềm, font hệ thống (Linux), Arial (Windows)."""
    if _FONT:
        return _FONT["n"], _FONT["b"]
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.fonts import addMapping
    except ImportError:
        raise ValueError("Tạo PDF cần thư viện reportlab: pip install reportlab (chay.bat / bản web tự cài)")
    here = os.path.dirname(os.path.abspath(__file__))
    win = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for reg, bold in ((os.path.join(here, "fonts", "DejaVuSans.ttf"), os.path.join(here, "fonts", "DejaVuSans-Bold.ttf")),
                      ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
                      ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
                      (os.path.join(win, "arial.ttf"), os.path.join(win, "arialbd.ttf")),
                      ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
                       "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")):
        if os.path.exists(reg) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("VN", reg))
            pdfmetrics.registerFont(TTFont("VN-B", bold))
            addMapping("VN", 0, 0, "VN")
            addMapping("VN", 1, 0, "VN-B")
            addMapping("VN", 0, 1, "VN")
            addMapping("VN", 1, 1, "VN-B")
            _FONT.update(n="VN", b="VN-B")
            return "VN", "VN-B"
    raise ValueError("Không tìm thấy font tiếng Việt – chép thư mục fonts/ (DejaVuSans.ttf) cạnh taihoadon.py")


def build_tax_pdf(v, tthai="", kq=""):
    """Bản thể hiện hoá đơn (A4) theo dữ liệu của cổng thuế → bytes PDF."""
    fn, fb = _vn_fonts()
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import html as H
    esc = lambda s: H.escape(str(s if s is not None else ""))
    st = ParagraphStyle("n", fontName=fn, fontSize=8.5, leading=11.5)
    small = ParagraphStyle("s", parent=st, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", parent=st, fontSize=8, leading=10)
    cell_r = ParagraphStyle("cr", parent=cell, alignment=TA_RIGHT)
    cell_c = ParagraphStyle("cc", parent=cell, alignment=TA_CENTER)
    title = ParagraphStyle("t", parent=st, fontName=fb, fontSize=15, leading=19, alignment=TA_CENTER,
                           textColor=colors.HexColor("#b4161b"))
    center = ParagraphStyle("ce", parent=st, alignment=TA_CENTER)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm,
                            bottomMargin=12 * mm, title="%s %s" % (v.get("kh"), v.get("so")), author=str(v["nb"].get("ten") or ""))
    W = A4[0] - 24 * mm
    d = _to_dt(v.get("ngay"))
    ngay = ("Ngày %02d tháng %02d năm %d" % (d.day, d.month, d.year)) if d else esc(v.get("ngay"))
    line = lambda k, val, bold=False: Paragraph("%s: %s" % (k, ("<b>%s</b>" if bold else "%s") % esc(val)), st) if val else None
    nb, nm, tong = v["nb"], v["nm"], v["tong"]
    story = []
    head = Table([[Paragraph("<b>%s</b>" % esc(nb.get("ten")), ParagraphStyle("h", parent=st, fontSize=10, leading=13)),
                   Paragraph("Mẫu số: <b>%s</b><br/>Ký hiệu: <b>%s</b><br/>Số: <b><font color='#b4161b'>%s</font></b>" % (
                       esc(v.get("mau")), esc(v.get("kh")), esc(v.get("so"))), st)]],
                 colWidths=[W - 55 * mm, 55 * mm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (1, 0), (1, 0), 0.6, colors.HexColor("#999999")),
                              ("LEFTPADDING", (0, 0), (0, 0), 0)]))
    story += [head, Spacer(1, 4 * mm),
              Paragraph(esc((v.get("title") or INVOICE_TITLES.get(str(v.get("mau")), "HÓA ĐƠN")).upper()), title),
              Paragraph("(Bản thể hiện của hóa đơn điện tử)", center), Paragraph(ngay, center),
              Paragraph("Mã của cơ quan thuế: <b>%s</b>" % esc(v.get("mccqt") or "(hóa đơn không có mã)"), center),
              Spacer(1, 3 * mm)]

    def block(rows):
        t = Table([[r] for r in rows if r is not None], colWidths=[W])
        t.setStyle(TableStyle([("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor("#bbbbbb")),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 1),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        return t
    story.append(block([line("Đơn vị bán hàng", nb.get("ten"), True), line("Mã số thuế", nb.get("mst"), True),
                        line("Địa chỉ", nb.get("dchi")), line("Điện thoại", nb.get("sdt")),
                        line("Số tài khoản", " – ".join(x for x in (nb.get("stk"), nb.get("nh")) if x))]))
    story.append(Spacer(1, 2 * mm))
    story.append(block([line("Họ tên người mua hàng", nm.get("hoten")), line("Tên đơn vị", nm.get("ten"), True),
                        line("Mã số thuế", nm.get("mst"), True), line("Địa chỉ", nm.get("dchi")),
                        line("Số tài khoản", nm.get("stk")),
                        Paragraph("Hình thức thanh toán: %s &nbsp;&nbsp;&nbsp; Đơn vị tiền tệ: %s" % (
                            esc(v.get("httt")), esc(v.get("dvtte") or "VND")), st)]))
    story.append(Spacer(1, 3 * mm))
    hdr = ["STT", "Tên hàng hóa, dịch vụ", "Đơn vị tính", "Số lượng", "Đơn giá", "Thuế suất", "Thành tiền"]
    data = [[Paragraph("<b>%s</b>" % h, cell_c) for h in hdr]]
    for it in v["items"]:
        note = str(it.get("tchat") or "")
        ten = esc(it.get("ten")) + (" <font color='#777777'>(%s)</font>" % {"2": "khuyến mại", "3": "chiết khấu",
                                                                             "4": "ghi chú"}.get(note, "") if note in ("2", "3", "4") else "")
        data.append([Paragraph(esc(it.get("stt")), cell_c), Paragraph(ten, cell), Paragraph(esc(it.get("dvt")), cell_c),
                     Paragraph(_vn_money(it.get("sl")), cell_r), Paragraph(_vn_money(it.get("dgia")), cell_r),
                     Paragraph(esc(_vn_rate(it.get("tsuat"))), cell_c), Paragraph(_vn_money(it.get("thtien")), cell_r)])
    if not v["items"]:
        data.append(["", Paragraph("<i>(Không có dữ liệu chi tiết hàng hóa, dịch vụ)</i>", cell), "", "", "", "", ""])
    items = Table(data, colWidths=[12 * mm, W - 114 * mm, 18 * mm, 18 * mm, 22 * mm, 16 * mm, 28 * mm], repeatRows=1)
    items.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f6")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story += [items, Spacer(1, 2 * mm)]
    rate_rows = [[Paragraph("<b>Thuế suất</b>", cell_c), Paragraph("<b>Thành tiền chưa thuế</b>", cell_c),
                  Paragraph("<b>Tiền thuế</b>", cell_c)]]
    rate_rows += [[Paragraph(esc(_vn_rate(r.get("tsuat"))), cell_c), Paragraph(_vn_money(r.get("thtien")), cell_r),
                   Paragraph(_vn_money(r.get("tthue")), cell_r)] for r in v["rates"]]
    tot = [("Tổng tiền chưa thuế", tong.get("cthue")), ("Tổng tiền thuế", tong.get("thue")),
           ("Tổng tiền chiết khấu thương mại", tong.get("ck") or 0), ("Tổng tiền phí", tong.get("phi") or 0),
           ("Tổng tiền thanh toán", tong.get("tt"))]
    tot_rows = [[Paragraph(("<b>%s</b>" if k.endswith("thanh toán") else "%s") % k, cell),
                 Paragraph(("<b>%s</b>" if k.endswith("thanh toán") else "%s") % _vn_money(val), cell_r)] for k, val in tot]
    rt = Table(rate_rows, colWidths=[18 * mm, 32 * mm, 28 * mm])
    rt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f6"))]))
    tt = Table(tot_rows, colWidths=[52 * mm, 34 * mm])
    tt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999"))]))
    both = Table([[rt, tt]], colWidths=[W - 88 * mm, 88 * mm])
    both.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [both, Spacer(1, 2 * mm)]
    if tong.get("chu"):
        story.append(Paragraph("Số tiền viết bằng chữ: <b>%s</b>" % esc(tong.get("chu")), st))
    story.append(Spacer(1, 5 * mm))
    sign = "<b>NGƯỜI BÁN HÀNG</b>"
    if v.get("signed") or v.get("mccqt"):
        sign += "<br/><font color='#1e8449'>Đã ký số</font><br/>Ký bởi: %s" % esc(nb.get("ten"))
        k = _to_dt(v.get("nky"))
        if k:
            sign += "<br/>Ký ngày: %s" % k.strftime("%d/%m/%Y")
    sig = Table([[Paragraph("<b>NGƯỜI MUA HÀNG</b>", center), Paragraph(sign, center)]], colWidths=[W / 2, W / 2])
    sig.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                             ("BOX", (1, 0), (1, 0), 0.6, colors.HexColor("#1e8449") if (v.get("signed") or v.get("mccqt"))
                              else colors.white)]))
    story += [sig, Spacer(1, 6 * mm)]
    meta = "Trạng thái: %s%s" % (esc(tthai or ""), (" · Kết quả kiểm tra: %s" % esc(kq)) if kq else "")
    story.append(Paragraph("%s<br/>Bản thể hiện dựng từ %s trên cổng hoadondientu.gdt.gov.vn (Tổng cục Thuế); "
                           "tra cứu hóa đơn tại https://hoadondientu.gdt.gov.vn" % (meta, esc(v.get("nguon"))), small))
    cancelled = bool(re.search(r"hủy|huỷ", tthai or "", re.I))

    def on_page(canvas, _doc):
        if cancelled:
            canvas.saveState()
            canvas.setFont(fb, 46)
            canvas.setFillColor(colors.Color(0.8, 0.1, 0.1, alpha=0.18))
            canvas.translate(A4[0] / 2, A4[1] / 2)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, "HÓA ĐƠN ĐÃ BỊ HỦY")
            canvas.restoreState()
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def tax_pdf(out_root, mst, kind, key, client=None):
    """Tạo (hoặc dùng lại) PDF bản thể hiện theo dữ liệu cổng thuế của một hoá đơn trong kho → đường dẫn tương đối."""
    base = company_dir(out_root, mst)
    idx = _load_json(index_file(out_root, mst), {})
    e = idx.get(kind, {}).get(key)
    if e is None:
        raise ValueError("Không tìm thấy hoá đơn")
    inv = dict(e.get("inv") or {})
    inv.setdefault("tdlap", e.get("tdlap"))
    src = ""
    if e.get("xml") and os.path.exists(os.path.join(base, e["xml"])):
        src = e["xml"]
    elif e.get("detail") and os.path.exists(os.path.join(base, e["detail"])):
        src = e["detail"]
    elif client is not None and client.token_valid():
        try:
            detail = client.get_detail(inv)
            src = save_detail_json(base, e, inv, detail)
            update_invoice_fields(out_root, mst, kind, key, detail=src)
        except PortalError:
            src = ""
    rel = (os.path.splitext(src)[0] if src else os.path.join("_cqt", invoice_basename(inv))) + "_CQT.pdf"
    path = os.path.join(base, rel)
    if os.path.exists(path) and (not src or os.path.getmtime(path) >= os.path.getmtime(os.path.join(base, src))):
        return rel
    if src.endswith(".xml"):
        with open(os.path.join(base, src), "rb") as f:
            v = invoice_view_from_xml(f.read())
    elif src:
        v = invoice_view_from_detail(_load_json(os.path.join(base, src), {}))
    else:
        v = invoice_view_from_summary(inv)
    _fill_view(v, invoice_view_from_summary(inv))
    tthai = TTHAI_NIBOT.get(_num(e.get("tthai")), "")
    kq = TTXLY_NIBOT.get(_num(inv.get("ttxly")), "") if inv else ""
    data = build_tax_pdf(v, tthai, kq)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f:
        f.write(data)
    return rel


def save_detail_json(base, e, inv, detail):
    """Lưu dữ liệu chi tiết của cổng cạnh thư mục XML của kỳ (hoặc _chi-tiet/) → đường dẫn tương đối với thư mục DN."""
    folder = os.path.dirname(e["xml"]) if e.get("xml") else "_chi-tiet"
    rel = os.path.join(folder, invoice_basename(inv) + ".json")
    os.makedirs(os.path.join(base, folder), exist_ok=True)
    _save_json(os.path.join(base, rel), detail)
    return rel


def update_invoice_fields(out_root, mst, kind, key, **fields):
    path = index_file(out_root, mst)
    with _INDEX_LOCK:
        cur = _load_json(path, {})
        e = cur.get(kind, {}).get(key)
        if e is not None:
            e.update(fields)
            _save_json(path, cur)


def render_invoice_html(out_root, mst, kind, key):
    """Bản thể hiện hoá đơn dựng từ XML (in hoặc lưu PDF từ trình duyệt)."""
    import html as H
    e = _load_json(index_file(out_root, mst), {}).get(kind, {}).get(key)
    if e is None:
        raise ValueError("Không tìm thấy hoá đơn")
    inv = e.get("inv") or {}
    xml_path = os.path.join(company_dir(out_root, mst), e["xml"]) if e.get("xml") else ""
    if xml_path and os.path.exists(xml_path):
        try:
            with open(xml_path, "rb") as fh:
                return render_xml_invoice(fh.read())
        except ValueError:
            pass
    info = parse_invoice_xml(xml_path) if xml_path else {}
    nb, nm, tc, tong = info.get("nb", {}), info.get("nm", {}), info.get("ttchung", {}), info.get("tong", {})
    d = invoice_date(inv)
    m = lambda v: "{:,.0f}".format(_float(v)).replace(",", ".")
    esc = lambda v: H.escape(str(v or ""))
    rows = "".join("<tr><td>%d</td><td>%s</td><td>%s</td><td class=n>%s</td><td class=n>%s</td><td class=n>%s</td>"
                   "<td class=n>%s</td></tr>" % (i, esc(it["name"]), esc(it["unit"]), ("%g" % it["qty"]) if it["qty"] else "",
                                                 m(it["price"]) if it["price"] else "", esc(it["rate"]), m(it["amount"]))
                   for i, it in enumerate(invoice_lines(dict(inv, _xmlinfo=info)), 1))
    tcu = e.get("tra_cuu") or {}
    party = lambda title, p, mst_, ten, dchi: (
        "<div class=box><b>%s</b><br>Tên: %s<br>MST: %s<br>Địa chỉ: %s</div>" % (title, esc(p.get("Ten") or ten),
                                                                          esc(p.get("MST") or mst_), esc(p.get("DChi") or dchi)))
    return """<!doctype html><html lang=vi><head><meta charset=utf-8><title>HĐ %s-%s</title><style>
body{font:14px/1.5 Arial,sans-serif;color:#111;max-width:900px;margin:20px auto;padding:0 16px}h1{text-align:center;font-size:20px;margin:4px}
.c{text-align:center}.grid{display:grid;grid-template-columns:1fr 1fr;gap:12px;margin:14px 0}.box{border:1px solid #999;padding:8px}
table{border-collapse:collapse;width:100%%}td,th{border:1px solid #999;padding:5px}.n{text-align:right}.mute{color:#666;font-size:12px}
@media print{.np{display:none}}</style></head><body>
<p class=np><button onclick="print()">In / Lưu PDF</button> <span class=mute>Bản thể hiện dựng từ XML – không thay thế bản PDF gốc của người bán.</span></p>
<h1>%s</h1><p class=c>Ký hiệu: <b>%s%s</b> &nbsp; Số: <b>%s</b> &nbsp; Ngày: <b>%s</b><br>Mã CQT: %s</p>
<div class=grid>%s%s</div><p>Hình thức thanh toán: %s</p>
<table><tr><th>STT</th><th>Tên hàng hoá, dịch vụ</th><th>ĐVT</th><th>SL</th><th>Đơn giá</th><th>TS</th><th>Thành tiền</th></tr>%s
<tr><td colspan=6 class=n>Cộng tiền hàng</td><td class=n>%s</td></tr><tr><td colspan=6 class=n>Tiền thuế GTGT</td><td class=n>%s</td></tr>
<tr><td colspan=6 class=n><b>Tổng tiền thanh toán</b></td><td class=n><b>%s</b></td></tr></table>
<p>Số tiền viết bằng chữ: %s</p><p class=mute>Nhà cung cấp HĐĐT: %s %s</p></body></html>""" % (
        esc(inv.get("khhdon")), esc(inv.get("shdon")), esc(tc.get("THDon") or "HÓA ĐƠN"),
        esc(inv.get("khmshdon")), esc(inv.get("khhdon")), esc(inv.get("shdon")), d.strftime("%d/%m/%Y") if d else "",
        esc(inv.get("mhdon")), party("Người bán", nb, inv.get("nbmst"), inv.get("nbten"), inv.get("nbdchi")),
        party("Người mua", nm, inv.get("nmmst"), inv.get("nmten"), inv.get("nmdchi")),
        esc(tc.get("HTTToan") or inv.get("thtttoan")), rows, m(tong.get("TgTCThue") or inv.get("tgtcthue")),
        m(tong.get("TgTThue") or inv.get("tgtthue")), m(tong.get("TgTTTBSo") or inv.get("tgtttbso")),
        esc(tong.get("TgTTTBChu")), esc(tcu.get("ncc")),
        ("– %s: %s" % (esc(tcu.get("field")), esc(tcu.get("code")))) if tcu.get("code") else "")


def update_invoices(out_root, mst, kind, keys, **fields):
    """Sửa duyệt / HĐ dịch vụ cho nhiều hoá đơn một lần."""
    path = index_file(out_root, mst)
    n = 0
    with _INDEX_LOCK:
        cur = _load_json(path, {})
        entries = cur.get(kind, {})
        for key in keys:
            e = entries.get(key)
            if e is None:
                continue
            for f in ("duyet", "dv"):
                if fields.get(f) is not None:
                    e[f] = bool(fields[f]) if f == "dv" else fields[f]
            n += 1
        _save_json(path, cur)
    return n


def _sibling(base, rel, ext):
    if not rel:
        return ""
    cand = os.path.splitext(rel)[0] + ext
    return cand if os.path.exists(os.path.join(base, cand)) else ""


def _cqt_rel(base, e):
    src = e.get("xml") or e.get("detail") or ""
    rel = (os.path.splitext(src)[0] + "_CQT.pdf") if src else ""
    return rel if rel and os.path.exists(os.path.join(base, rel)) else ""


def query_invoices(out_root, mst, f):
    """Lọc hoá đơn trong kho theo bộ lọc của màn hình Hoá đơn. Trả về list dòng (dict)."""
    kind = f.get("kind") or "purchase"
    only_dv = kind.endswith("_dv")
    kind = kind.replace("_dv", "")
    base = company_dir(out_root, mst)
    entries = _load_json(index_file(out_root, mst), {}).get(kind, {})
    start = parse_date(f["from"]) if f.get("from") else None
    end = parse_date(f["to"]) if f.get("to") else None
    q = (f.get("q") or "").strip().lower()
    only = set(f.get("keys") or [])  # các dòng người dùng đã tích chọn
    rows = []
    for key, e in entries.items():
        if only and key not in only:
            continue
        inv = e.get("inv") or {}
        d = invoice_date(inv) or invoice_date({"tdlap": e.get("tdlap")})
        if (start and (not d or d < start)) or (end and (not d or d > end)):
            continue
        if only_dv and not e.get("dv"):
            continue
        xml = e.get("xml", "") if e.get("xml") and os.path.exists(os.path.join(base, e["xml"])) else ""
        html, pdf = _sibling(base, xml, ".html"), _sibling(base, xml, ".pdf")
        file_f = f.get("file") or ""
        if not pdf and e.get("pdf") and os.path.exists(os.path.join(base, e["pdf"])):
            pdf = e["pdf"]
        if (file_f == "no_xml" and xml) or (file_f == "xml" and not xml) or (file_f == "pdf" and not pdf) \
                or (file_f == "no_pdf" and pdf) or (file_f == "xml_pdf" and not (xml and pdf)):
            continue
        duyet = e.get("duyet") or DUYET_OPTIONS[0]
        if f.get("duyet") and f["duyet"] != duyet:
            continue
        tthai, ttxly = _num(e.get("tthai")), _num(inv.get("ttxly"))
        if f.get("tthai") not in (None, "") and _num(f["tthai"]) != tthai:
            continue
        if f.get("kq") not in (None, "") and _num(f["kq"]) != ttxly:
            continue
        if f.get("khhdon") and f["khhdon"].upper() not in ("%s%s" % (inv.get("khmshdon") or "", inv.get("khhdon") or "")).upper():
            continue
        if f.get("shdon") and str(f["shdon"]).strip() != str(inv.get("shdon") or ""):
            continue
        mst_dt, ten_dt = ((inv.get("nbmst"), inv.get("nbten")) if kind == "purchase"
                          else (inv.get("nmmst"), inv.get("nmten")))
        if q and q not in " ".join(str(x or "") for x in (mst_dt, ten_dt, e.get("note"), e.get("mat_hang"))).lower():
            continue
        rows.append({
            "key": key, "mst": str(mst_dt or ""), "ten": ten_dt or "", "ngay": d.strftime("%d/%m/%Y") if d else "",
            "_sort": (d.isoformat() if d else "", _num(inv.get("shdon"))),
            "khhdon": "%s%s" % (inv.get("khmshdon") or "", inv.get("khhdon") or ""), "shdon": str(inv.get("shdon") or ""),
            "cthue": _float(inv.get("tgtcthue")), "thue": _float(inv.get("tgtthue")), "ck": _float(inv.get("ttcktmai")),
            "phi": _float(inv.get("tgtphi")), "tt": _float(inv.get("tgtttbso")),
            "tthai": TTHAI_NIBOT.get(tthai, ""), "kq": TTXLY_NIBOT.get(ttxly, "") if inv else "",
            "duyet": duyet, "dv": bool(e.get("dv")), "mat_hang": e.get("mat_hang", ""), "note": e.get("note", ""),
            "xml": xml, "html": html, "pdf": pdf or (e.get("pdf") if e.get("pdf") and
                                                     os.path.exists(os.path.join(base, e["pdf"])) else ""),
            "tra_cuu": e.get("tra_cuu") or {}, "pdf_url": browser_pdf_url(e.get("tra_cuu")),
            "can_fetch": can_fetch_pdf(e.get("tra_cuu"))[0], "cqt": _cqt_rel(base, e)})
    rows.sort(key=lambda r: r.pop("_sort"))
    return rows


def export_invoices(out_root, mst, company_name, f, fmt):
    """Kết xuất các HĐ đang lọc: 'xlsx' (mẫu Nibot) hoặc 'xml' / 'html' / 'pdf' (gói .zip). Trả về đường dẫn."""
    kind = (f.get("kind") or "purchase").replace("_dv", "")
    rows = query_invoices(out_root, mst, f)
    if not rows:
        raise ValueError("Không có hoá đơn nào để kết xuất")
    base = company_dir(out_root, mst)
    entries = _load_json(index_file(out_root, mst), {}).get(kind, {})
    out_dir = os.path.join(base, "_ket-xuat")
    os.makedirs(out_dir, exist_ok=True)
    label = "MUA_VAO" if kind == "purchase" else "BAN_RA"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if fmt == "xlsx":
        invoices = []
        for r in rows:
            e = entries[r["key"]]
            inv = dict(e.get("inv") or {})
            inv.setdefault("tthai", e.get("tthai"))
            inv.setdefault("tdlap", e.get("tdlap"))
            inv["_duyet"], inv["_note"] = r["duyet"], r["note"]
            if r["xml"]:
                inv["_xml"] = r["xml"]
                inv["_xmlinfo"] = parse_invoice_xml(os.path.join(base, r["xml"]))
            invoices.append(inv)
        dates = [invoice_date(i) for i in invoices if invoice_date(i)]
        start = parse_date(f["from"]) if f.get("from") else min(dates)
        end = parse_date(f["to"]) if f.get("to") else max(dates)
        path = os.path.join(out_dir, "%s_%s_%s.xlsx" % (label, mst, stamp))
        return write_nibot_workbook(path, kind, invoices, mst, company_name, start, end)
    if fmt == "cqt":
        files = [tax_pdf(out_root, mst, kind, r["key"]) for r in rows]
        path = os.path.join(out_dir, "%s_PDF_THUE_%s_%s.zip" % (label, mst, stamp))
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for rel in files:
                z.write(os.path.join(base, rel), os.path.basename(rel))
        return path
    if fmt not in ("xml", "html", "pdf"):
        raise ValueError("Định dạng kết xuất không hợp lệ")
    files = [r[fmt] for r in rows if r[fmt]]
    if not files:
        raise ValueError("Các hoá đơn đang lọc chưa có file %s" % fmt.upper())
    path = os.path.join(out_dir, "%s_%s_%s_%s.zip" % (label, fmt.upper(), mst, stamp))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for rel in files:
            z.write(os.path.join(base, rel), os.path.basename(rel))
    return path


# ---------------------------------------------------------------------------
# Tiện ích: đọc hoá đơn XML, kiểm tra MST hàng loạt, nối / tách file PDF
# ---------------------------------------------------------------------------

def read_invoice_xml(xml_bytes):
    """Đọc toàn bộ nội dung hiển thị của một hoá đơn XML (chuẩn TT78/ND123)."""
    import xml.etree.ElementTree as ET
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError as e:
        raise ValueError("File không phải XML hoá đơn hợp lệ (%s)" % e)

    def find(el, name):
        for c in el.iter():
            if _local(c.tag) == name:
                return c
        return None

    def flat(el):
        return {_local(c.tag): (c.text or "").strip() for c in el if len(c) == 0} if el is not None else {}

    dl = find(root, "DLHDon")
    if dl is None:
        raise ValueError("Không thấy phần dữ liệu hoá đơn (DLHDon) trong file XML")
    tc, nd = find(dl, "TTChung"), find(dl, "NDHDon")
    nb, nm, tt = (find(nd, n) if nd is not None else None for n in ("NBan", "NMua", "TToan"))
    items = []
    for el in dl.iter():
        if _local(el.tag) == "HHDVu":
            f = flat(el)
            items.append({k: f.get(k, "") for k in ("STT", "TChat", "THHDVu", "DVTinh", "SLuong", "DGia", "TLCKhau",
                                                    "STCKhau", "TSuat", "ThTien")})
    rates = [flat(el) for el in dl.iter() if _local(el.tag) == "LTSuat"]
    mccqt = find(root, "MCCQT")
    sigs = [_local(el.tag) for el in root.iter() if _local(el.tag) == "Signature"]
    info = {"ttchung": flat(tc), "nb": flat(nb), "nm": flat(nm), "tong": flat(tt), "items": items, "rates": rates,
            "mccqt": (mccqt.text or "").strip() if mccqt is not None else "", "so_chu_ky": len(sigs),
            "ttkhac": {}, "msttcgp": "", "dlhdon_id": dl.get("Id", "")}
    info["msttcgp"] = info["ttchung"].get("MSTTCGP", "")
    for el in dl.iter():
        if _local(el.tag) == "TTin":
            f = flat(el)
            if f.get("TTruong") and f.get("DLieu"):
                info["ttkhac"].setdefault(f["TTruong"], f["DLieu"])
    return info


def render_xml_invoice(xml_bytes, title_note=""):
    """Bản thể hiện hoá đơn (HTML) dựng từ file XML – xem, in hoặc lưu PDF."""
    import html as H
    x = read_invoice_xml(xml_bytes)
    tc, nb, nm, tong = x["ttchung"], x["nb"], x["nm"], x["tong"]
    esc = lambda v: H.escape(str(v or ""))

    def m(v):
        if v in (None, ""):
            return ""
        f = _float(v)
        return ("{:,.0f}" if f == int(f) else "{:,.2f}").format(f)
    d = _to_dt(tc.get("NLap"))
    ngay = "Ngày %02d tháng %02d năm %d" % (d.day, d.month, d.year) if d else esc(tc.get("NLap"))
    rows = "".join(
        "<tr><td class=c>%s</td><td>%s</td><td class=c>%s</td><td class=n>%s</td><td class=n>%s</td>"
        "<td class=n>%s</td><td class=c>%s</td><td class=n>%s</td></tr>" % (
            esc(it["STT"]), esc(it["THHDVu"]), esc(it["DVTinh"]), m(it["SLuong"]), m(it["DGia"]),
            m(it["STCKhau"]) or "0", esc(it["TSuat"]), m(it["ThTien"]))
        for it in x["items"])
    rate_rows = "".join("<tr><td class=c>%s</td><td class=n>%s</td><td class=n>%s</td></tr>" % (
        esc(r.get("TSuat")), m(r.get("ThTien")), m(r.get("TThue"))) for r in x["rates"])
    lk = lookup_info(x)

    def party(title, p):
        rows_ = [("Tên " + title, p.get("Ten")), ("Mã số thuế", p.get("MST") or p.get("CCCDan")), ("Địa chỉ", p.get("DChi")),
                 ("Điện thoại", p.get("SDThoai")), ("Số tài khoản", p.get("STKNHang"))]
        return "".join("<div>%s: <b>%s</b></div>" % (k, esc(v)) for k, v in rows_ if v or k.startswith("Tên"))
    tot = [("Tổng tiền chưa thuế", tong.get("TgTCThue")), ("Tổng tiền thuế", tong.get("TgTThue")),
           ("Tổng tiền phí", tong.get("TgTPhi") or "0"), ("Tổng tiền CKTM", tong.get("TTCKTMai") or "0"),
           ("Tổng tiền thanh toán", tong.get("TgTTTBSo"))]
    sig = ("Có %d chữ ký số trong file (phần mềm chưa kiểm tra tính hợp lệ của chữ ký)" % x["so_chu_ky"]
           if x["so_chu_ky"] else "Không có chữ ký số trong file")
    return """<!doctype html><html lang=vi><head><meta charset=utf-8><title>%s %s-%s</title><style>
body{font:14px/1.55 Arial,sans-serif;color:#111;max-width:960px;margin:16px auto;padding:0 16px;background:#fff}
.top{display:grid;grid-template-columns:1fr auto;gap:12px;border-bottom:1px solid #ccc;padding-bottom:8px}
h1{font-size:21px;margin:0;text-align:center}.c{text-align:center}.n{text-align:right}.sec{border-bottom:1px solid #ccc;padding:6px 0}
table{border-collapse:collapse;width:100%%;margin-top:8px}td,th{border:1px solid #888;padding:4px 6px}th{background:#f3f3f3}
.two{display:grid;grid-template-columns:1fr 1fr;gap:12px;align-items:start}.mute{color:#666;font-size:12px}.np{margin-bottom:8px}
@media print{.np{display:none}}</style></head><body>
<div class=np><button onclick="print()">In / Lưu PDF</button> <span class=mute>%s</span></div>
<div class=top><div><h1>%s</h1><p class=c>%s<br>MCCQT: %s</p></div>
<div>Mẫu số: <b>%s</b><br>Ký hiệu: <b>%s</b><br>Số: <b>%s</b></div></div>
<div class=sec>%s</div><div class=sec>%s<div class=two><div>Hình thức thanh toán: <b>%s</b></div><div>Đơn vị tiền tệ: <b>%s</b></div></div></div>
<table><tr><th>STT</th><th>Tên hàng hoá, dịch vụ</th><th>ĐVT</th><th>SL</th><th>Đơn giá</th><th>Tiền CK</th><th>Thuế suất</th><th>Thành tiền</th></tr>%s</table>
<div class=two><table><tr><th>Thuế suất</th><th>Tổng tiền chưa thuế</th><th>Tiền thuế</th></tr>%s</table>
<table>%s<tr><td colspan=2><b>Bằng chữ:</b> %s</td></tr></table></div>
<p class=mute>%s<br>Nhà cung cấp HĐĐT: %s%s</p></body></html>""" % (
        esc(tc.get("KHHDon")), esc(tc.get("SHDon")), esc(tc.get("THDon")), esc(title_note) or "Bản thể hiện dựng từ XML",
        esc(tc.get("THDon") or "HÓA ĐƠN"), ngay, esc(x["mccqt"] or "(không mã)"), esc(tc.get("KHMSHDon")),
        esc(tc.get("KHHDon")), esc(tc.get("SHDon")), party("người bán", nb), party("người mua", nm),
        esc(tc.get("HTTToan")), esc(tc.get("DVTTe") or "VND"), rows, rate_rows,
        "".join("<tr><td>%s</td><td class=n>%s</td></tr>" % (k, m(v)) for k, v in tot), esc(tong.get("TgTTTBChu")), sig,
        esc(lk["ncc"]), (" – %s: %s" % (esc(lk["field"]), esc(lk["code"]))) if lk["code"] else "")


def _pypdf():
    try:
        import pypdf
        return pypdf
    except ImportError:
        raise ValueError("Chưa cài thư viện pypdf. Đóng phần mềm và chạy lại chay.bat (hoặc: pip install pypdf)")


def merge_pdfs(files):
    """files: list (tên, bytes) theo thứ tự → bytes PDF đã nối."""
    pypdf = _pypdf()
    if len(files) < 2:
        raise ValueError("Chọn ít nhất 2 file PDF để nối")
    w = pypdf.PdfWriter()
    for name, data in files:
        try:
            r = pypdf.PdfReader(io.BytesIO(data))
            for page in r.pages:
                w.add_page(page)
        except Exception as e:  # file hỏng / có mật khẩu
            raise ValueError("Không đọc được file %s: %s" % (name, e))
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def parse_page_ranges(text, total):
    """'1-3, 5, 7-8' → [[1,2,3],[5],[7,8]]; để trống → mỗi trang một file."""
    text = (text or "").strip()
    if not text:
        return [[i] for i in range(1, total + 1)]
    groups = []
    for part in re.split(r"[;,]", text):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)|(\d+)", part)
        if not m:
            raise ValueError("Khoảng trang không hợp lệ: %r (ví dụ đúng: 1-3, 5, 7-8)" % part)
        a, b = (int(m.group(1)), int(m.group(2))) if m.group(1) else (int(m.group(3)), int(m.group(3)))
        if not (1 <= a <= b <= total):
            raise ValueError("Trang %s nằm ngoài file (file có %d trang)" % (part, total))
        groups.append(list(range(a, b + 1)))
    return groups


def split_pdf(name, data, ranges_text):
    """Tách PDF theo khoảng trang → bytes file .zip chứa các PDF con."""
    pypdf = _pypdf()
    try:
        r = pypdf.PdfReader(io.BytesIO(data))
        total = len(r.pages)
    except Exception as e:
        raise ValueError("Không đọc được file %s: %s" % (name, e))
    base = os.path.splitext(os.path.basename(name or "file"))[0]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for pages in parse_page_ranges(ranges_text, total):
            w = pypdf.PdfWriter()
            for p in pages:
                w.add_page(r.pages[p - 1])
            part = io.BytesIO()
            w.write(part)
            label = str(pages[0]) if len(pages) == 1 else "%d-%d" % (pages[0], pages[-1])
            z.writestr("%s_trang_%s.pdf" % (base, label), part.getvalue())
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Sao kê ngân hàng → file KTSC nhập phần mềm kế toán (Smart Pro), theo mẫu Nibot
# ---------------------------------------------------------------------------

BANK_NAMES = (("VIETINBANK", "VIETIN"), ("CONG THUONG", "VIETIN"), ("VIETCOMBANK", "VCB"), ("NGOAI THUONG", "VCB"),
              ("TECHCOMBANK", "TCB"), ("KY THUONG", "TCB"), ("BIDV", "BIDV"), ("DAU TU VA PHAT TRIEN", "BIDV"),
              ("AGRIBANK", "AGRIBANK"), ("ACB", "ACB"), ("A CHAU", "ACB"), ("MBBANK", "MB"), ("QUAN DOI", "MB"),
              ("SACOMBANK", "SACOMBANK"), ("SAI GON THUONG TIN", "SACOMBANK"), ("VPBANK", "VPBANK"),
              ("TPBANK", "TPBANK"), ("TIEN PHONG", "TPBANK"), ("SHB", "SHB"), ("HDBANK", "HDBANK"), ("OCB", "OCB"),
              ("VIB", "VIB"), ("EXIMBANK", "EXIMBANK"), ("SEABANK", "SEABANK"), ("MSB", "MSB"), ("HANG HAI", "MSB"),
              ("SHINHAN", "SHINHAN"), ("WOORI", "WOORI"))
KTSC_BANK_COLS = ("LCTG", "NGAYCT", "SOCT", "DIENGIAI", "TKNO", "MADTPNNO", "TKCO", "MADTPNCO", "TTVND", "TTVND_TT",
                  "TENKH", "ID_NGHIEPVU", "GHICHU", "GUID")


def _bank_plain(v):
    return re.sub(r"[^a-z0-9/]+", "", _plain(v))


def read_table_file(name, data):
    """Đọc file bảng tính sao kê → [(tên sheet, [hàng [ô...]])]. Hỗ trợ .xlsx, .xls (Excel 97-2003, cần xlrd),
    .xls/.html dạng bảng HTML (nhiều ngân hàng xuất kiểu này) và .csv."""
    head = data[:8]
    if head[:2] == b"PK":
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Chưa cài openpyxl (pip install openpyxl)")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise ValueError("Không đọc được file Excel %s: %s" % (name, e))
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    if head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        try:
            import xlrd
        except ImportError:
            raise ValueError("Đọc file .xls (Excel 97-2003) cần thư viện xlrd: pip install xlrd "
                             "(bản web: khởi động lại dịch vụ để tự cài) – hoặc mở file bằng Excel, lưu lại dạng .xlsx")
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as e:
            raise ValueError("Không đọc được file .xls %s: %s" % (name, e))
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                            continue
                        except Exception:
                            pass
                    row.append(cell.value)
                rows.append(row)
            sheets.append((sh.name, rows))
        return sheets
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
        try:
            text = data.decode(enc)
            if enc != "utf-16" or "\x00" not in text:
                break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Không nhận ra định dạng file %s" % name)
    if re.search(r"<\s*(table|tr)\b", text[:200000], re.I):
        from html.parser import HTMLParser

        class P(HTMLParser):
            def __init__(self):
                super().__init__()
                self.rows, self.row, self.cell = [], None, None

            def handle_starttag(self, tag, attrs):
                if tag == "tr":
                    self.row = []
                elif tag in ("td", "th") and self.row is not None:
                    self.cell = []
                elif tag == "br" and self.cell is not None:
                    self.cell.append(" ")

            def handle_endtag(self, tag):
                if tag in ("td", "th") and self.cell is not None and self.row is not None:
                    self.row.append(re.sub(r"\s+", " ", "".join(self.cell)).strip())
                    self.cell = None
                elif tag == "tr" and self.row is not None:
                    self.rows.append(self.row)
                    self.row = None

            def handle_data(self, d):
                if self.cell is not None:
                    self.cell.append(d)
        p = P()
        p.feed(text)
        return [(os.path.splitext(name)[0], p.rows)]
    sample = text[:5000]
    delim = max((";", ",", "\t", "|"), key=sample.count)
    return [(os.path.splitext(name)[0], [r for r in csv.reader(io.StringIO(text), delimiter=delim)])]


def _bank_amount(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-()]", "", str(v))
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("-()")
    if not s:
        return 0.0
    if "," in s and "." in s:  # 1,234,567.89 hoặc 1.234.567,89
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        s = s.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif s.count(",") > 1 or s.count(".") > 1 or re.fullmatch(r"\d{1,3}([.,]\d{3})+", s):
        s = s.replace(",", "").replace(".", "")  # dấu ngăn cách hàng nghìn
    else:
        s = s.replace(",", ".")
    try:
        f = float(s)
    except ValueError:
        return 0.0
    return -f if neg else f


def _bank_date(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v or "").strip()
    m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})(?:\D+(\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        hh, mm, ss = (int(m.group(i) or 0) for i in (4, 5, 6))
    else:
        m = re.match(r"(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})(?:\D+(\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
        if not m:
            return None
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        hh, mm, ss = (int(m.group(i) or 0) for i in (4, 5, 6))
    try:
        return datetime(y, mo, d, hh, mm, ss)
    except ValueError:
        return None


def _bank_col(text):
    """Tên cột sao kê → loại cột (date, desc, out, in, amount, ref, name, None)."""
    p = _bank_plain(text)
    toks = set(re.split(r"[^a-z0-9]+", _plain(text)))
    if not p or "stt" in toks or p.startswith("stt") or "sodu" in p or "balance" in p or "luyke" in p:
        return None
    if "ten" in p and ("doiung" in p or "corresponsive" in p or "thuhuong" in p or "nguoichuyen" in p
                       or "nguoinhan" in p or "khachhang" in p or "counterpart" in p or "beneficiary" in p):
        return "name"
    if "sotaikhoan" in p or "accountno" in p or "taikhoandoiung" in p or "virtualaccount" in p:
        return None
    if any(k in p for k in ("sogiaodich", "transactionnumber", "sothamchieu", "reference", "refno", "magiaodich",
                            "sobut", "soct", "transactionid", "transno")):
        return "ref"
    if any(k in p for k in ("mota", "noidung", "diengiai", "description", "details", "remark", "narrative")):
        return "desc"
    if "ngay" in p or "date" in p:
        return "date"
    if any(k in p for k in ("ghino", "sotienno", "phatsinhno", "debit", "withdraw", "sotienchi", "tienra", "rut")) \
            or p.startswith("no/") or p == "no" or p.startswith("chi"):
        return "out"
    if any(k in p for k in ("ghico", "sotienco", "phatsinhco", "credit", "deposit", "sotienthu", "tienvao")) \
            or p.startswith("co/") or p == "co" or p.startswith("thu"):
        return "in"
    if "sotien" in p or "amount" in p:
        return "amount"
    return None


def parse_bank_statement(name, data):
    """Đọc sao kê ngân hàng → thông tin tài khoản + danh sách giao dịch (theo thứ tự trong file)."""
    if data[:5] == b"%PDF-":
        return _finish_statement(parse_bank_pdf(name, data))
    best = None
    for sheet, rows in read_table_file(name, data):
        for i, row in enumerate(rows[:80]):
            cols = {}
            for j, v in enumerate(row):
                kind = _bank_col(v) if isinstance(v, str) else None
                if kind and kind not in cols:
                    cols[kind] = j
            if "date" in cols and ("desc" in cols or "ref" in cols) and (
                    ("in" in cols and "out" in cols) or "amount" in cols):
                best = (sheet, rows, i, cols)
                break
        if best:
            break
    if not best:
        raise ValueError("%s: không tìm thấy bảng giao dịch (cần các cột Ngày, Mô tả/Nội dung, Nợ/Có hoặc Số tiền)" % name)
    sheet, rows, hdr, cols = best
    top = " ".join(_plain(c) if isinstance(c, str) else str(c or "") for r in rows[:hdr] for c in r)
    top += " " + _plain(name)
    if data[:2] != b"PK" and data[:4] != b"\xd0\xcf\x11\xe0":  # file dạng chữ (html/csv): tên ngân hàng có thể nằm ngoài bảng
        top += " " + _plain(re.sub(r"<[^>]+>", " ", data[:20000].decode("utf-8", "ignore")))
    info = {"file": name, "bank": "", "account": "", "company": "", "from": "", "to": "",
            "opening": None, "closing": None, "total_in": None, "total_out": None, "rows": [], "warnings": []}
    for key, code in BANK_NAMES:
        if _plain(key).replace(" ", "") in top.replace(" ", ""):
            info["bank"] = code
            break

    def label_value(*labels):
        for r in rows[:hdr]:
            for j, c in enumerate(r):
                if isinstance(c, str) and any(lb in _bank_plain(c) for lb in labels):
                    rest = [x for x in r[j + 1:] if x not in (None, "")]
                    if rest:
                        return rest[0]
                    m = re.search(r":\s*(.+)$", c)
                    if m:
                        return m.group(1).strip()
        return None
    acc = label_value("sotaikhoan", "accountno", "taikhoanso")
    info["account"] = re.sub(r"[^\dA-Za-z]", "", str(acc or ""))[:30]
    info["company"] = str(label_value("congty/companyname", "tenkhachhang", "tentaikhoan", "chutaikhoan",
                                      "companyname", "accountname") or "").strip()
    for key, labels in (("opening", ("sodudauky", "openingbalance")), ("closing", ("soducuoiky", "closingbalance")),
                        ("total_in", ("tonggiatrighico", "tongghico", "totalcredit", "tongsotienghico")),
                        ("total_out", ("tonggiatrighino", "tongghino", "totaldebit", "tongsotienghino"))):
        v = label_value(*labels)
        if v not in (None, ""):
            info[key] = _bank_amount(v)
    # Hàng "Tổng giá trị ghi có ... Tổng giá trị ghi nợ" nằm cùng một hàng: tìm riêng từng nhãn
    for r in rows[:hdr]:
        for j, c in enumerate(r):
            if not isinstance(c, str):
                continue
            p = _bank_plain(c)
            nxt = next((x for x in r[j + 1:] if x not in (None, "")), None)
            if nxt is None:
                continue
            if "ghico" in p and "tong" in p and "so" not in p.replace("tongso", "") and "giaodich" not in p:
                info["total_in"] = _bank_amount(nxt)
            elif "ghino" in p and "tong" in p and "giaodich" not in p:
                info["total_out"] = _bank_amount(nxt)
    ncol = max(cols.values()) + 1
    for r in rows[hdr + 1:]:
        r = list(r) + [None] * (ncol - len(r))
        d = _bank_date(r[cols["date"]])
        if not d:
            continue
        if "amount" in cols and not ("in" in cols and "out" in cols):
            a = _bank_amount(r[cols["amount"]])
            tin, tout = (a, 0.0) if a > 0 else (0.0, -a)
        else:
            tin, tout = abs(_bank_amount(r[cols["in"]])), abs(_bank_amount(r[cols["out"]]))
        if not tin and not tout:
            continue
        val = lambda k: re.sub(r"[\r\n\t]+", " ", str(r[cols[k]] if k in cols and r[cols[k]] is not None else "")).strip()
        ref = val("ref")
        if ref.endswith(".0") and ref[:-2].isdigit():
            ref = ref[:-2]
        info["rows"].append({"date": d.strftime("%Y-%m-%d"), "time": d.strftime("%H:%M:%S"), "ref": ref,
                             "desc": val("desc"), "name": val("name"), "in": tin, "out": tout})
    return _finish_statement(info)


def _finish_statement(info):
    """Kỳ sao kê, tổng thu/chi và đối chiếu với tổng / số dư in trên sao kê."""
    name = info["file"]
    if not info["rows"]:
        raise ValueError("%s: không có giao dịch nào" % name)
    ds = sorted(x["date"] for x in info["rows"])
    info["from"], info["to"] = ds[0], ds[-1]
    s_in, s_out = sum(x["in"] for x in info["rows"]), sum(x["out"] for x in info["rows"])
    info["sum_in"], info["sum_out"] = s_in, s_out
    if info["total_in"] is not None and abs(info["total_in"] - s_in) > 0.5:
        info["warnings"].append("Tổng ghi có trên sao kê %s ≠ tổng các dòng đọc được %s" % (
            "{:,.0f}".format(info["total_in"]), "{:,.0f}".format(s_in)))
    if info["total_out"] is not None and abs(info["total_out"] - s_out) > 0.5:
        info["warnings"].append("Tổng ghi nợ trên sao kê %s ≠ tổng các dòng đọc được %s" % (
            "{:,.0f}".format(info["total_out"]), "{:,.0f}".format(s_out)))
    if info["opening"] is not None and info["closing"] is not None and \
            abs(info["opening"] + s_in - s_out - info["closing"]) > 0.5:
        info["warnings"].append("Số dư đầu kỳ + thu − chi không bằng số dư cuối kỳ – kiểm tra lại file")
    return info


_AMOUNT_RE = re.compile(r"^-?(?:0|\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?|\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?|\d{1,3})$")
_PDF_SKIP_RE = re.compile(r"chungtunay|automaticallyexported|tongphatsinh|tongcong|soducuoi|closingbalance|bangchu|"
                          r"inwords|^trang\d|^page\d")


class _PdfSkip:
    """Dòng chân trang / tổng cộng / số dư cuối kỳ – không phải giao dịch (so khớp không dấu)."""
    @staticmethod
    def search(text):
        return _PDF_SKIP_RE.search(_bank_plain(text))


_PDF_SKIP = _PdfSkip()


def _unchunk35(s):
    """Nội dung giao dịch MB (core T24) bị cắt thành đoạn 35 ký tự nối bằng 2 dấu cách ("chuyen kh  oan") → nối lại."""
    i = 0
    while len(s) > i + 37 and s[i + 35:i + 37] == "  ":
        s = s[:i + 35] + s[i + 37:]
        i += 35
    return s


def _pdf_fragments(page):
    """Các đoạn chữ trên trang PDF kèm toạ độ đọc (x tăng sang phải, y tăng lên trên)."""
    out = []

    def visit(text, cm, tm, fd, fs):
        t = (text or "").replace("\n", " ").strip()
        if not t:
            return
        if abs(cm[1]) > 0.5 and abs(cm[2]) > 0.5:  # trang xoay ngang: toạ độ trong hệ của chữ
            x, y = tm[4], tm[5]
        else:
            x, y = tm[4] * cm[0] + cm[4], tm[5] * cm[3] + cm[5]
        out.append({"x": round(x, 1), "y": round(y, 1), "t": t})
    page.extract_text(visitor_text=visit)
    return out


def _bank_col_pdf(text):
    p = _bank_plain(text)
    if any(k in p for k in ("donvithuhuong", "donvichuyen", "beneficiary", "nguoithuhuong", "nguoichuyen")):
        return "name"
    if "hachtoan" in p or "accountingdate" in p:
        return "date_acc"
    return _bank_col(text)


def parse_bank_pdf(name, data):
    """Sổ phụ ngân hàng dạng PDF có chữ (xuất từ internet banking, không phải bản scan)."""
    pypdf = _pypdf()
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = [_pdf_fragments(p) for p in reader.pages]
        full = "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception as e:
        raise ValueError("Không đọc được PDF %s: %s" % (name, e))
    if len(re.sub(r"\s", "", full)) < 50:
        raise ValueError("%s: PDF không có chữ (bản scan/ảnh chụp) – cần file sổ phụ tải từ internet banking hoặc file Excel" % name)
    info = {"file": name, "bank": "", "account": "", "company": "", "from": "", "to": "", "opening": None, "closing": None,
            "total_in": None, "total_out": None, "rows": [], "warnings": []}
    head = full.split("\n")
    flat = _plain(" ".join(head[:6]))  # tên ngân hàng ở đầu sổ phụ (nội dung giao dịch có tên ngân hàng đối tác)
    for key, code in BANK_NAMES:
        if _plain(key).replace(" ", "") in flat.replace(" ", ""):
            info["bank"] = code
            break

    def grab(pattern):
        m = re.search(pattern, full, re.I)
        return m.group(1).strip() if m else ""
    info["account"] = re.sub(r"\D", "", grab(r"(?:Số tài khoản|Tài khoản|Account No\.?)\s*/?[^:\n]{0,25}:\s*([\d .\-]{6,25})"))[:20]
    info["company"] = grab(r"(?:Tên khách hàng|Tên tài khoản|Customer name|Account name)\s*/?[^:\n]{0,25}:\s*([^\n]+)")
    for key, pat in (("opening", r"(?:Số dư đầu kỳ|Opening Balance)[^:\n]*:\s*(-?[\d.,]+)"),
                     ("closing", r"(?:Số dư cuối kỳ|Closing Balance)[^:\n]*:\s*(-?[\d.,]+)")):
        v = grab(pat)
        if v:
            info[key] = _bank_amount(v)
    header = None

    def header_band(frags):
        """Dòng tiêu đề bảng: các nhãn cột ngắn nằm cùng dải với nhãn ngày."""
        for d in frags:
            if _bank_col_pdf(d["t"]) not in ("date", "date_acc") or len(d["t"]) > 30:
                continue
            cols = {}
            for f in sorted(frags, key=lambda f: (f["x"], -f["y"])):
                if abs(f["y"] - d["y"]) > 22 or len(f["t"]) > 40 or re.search(r":\s*\S", f["t"]):
                    continue  # nhãn cột không kèm giá trị ("Tên khách hàng: …" là thông tin tài khoản)
                k = _bank_col_pdf(f["t"])
                if k and (k not in cols or len(f["t"]) > len(cols[k]["t"])):  # nhãn đầy đủ nhất ("Phát sinh nợ", không phải "No")
                    cols[k] = f
            if ("date" in cols or "date_acc" in cols) and "in" in cols and "out" in cols:
                return cols
        return None
    for frags in pages:
        header = header_band(frags)
        if header:
            break
    if not header:
        raise ValueError("%s: không tìm thấy bảng giao dịch trong sổ phụ PDF" % name)
    amount_hdr = sorted((header[k]["x"], k) for k in ("out", "in"))
    amt_lo, amt_hi = amount_hdr[0][0] - 45, amount_hdr[-1][0] + 90
    date_x = header.get("date_acc", header.get("date"))["x"]
    ref_x = header["ref"]["x"] if "ref" in header else None
    text_hdrs = sorted((header[k]["x"], k) for k in ("desc", "name") if k in header)

    def is_amount(f):
        return amt_lo <= f["x"] <= amt_hi and _AMOUNT_RE.match(f["t"].replace(" ", ""))
    # Cột chữ (nội dung, tên đối ứng…): nhóm theo vị trí bắt đầu của chữ trên toàn bộ sổ phụ
    starts = {}
    body = []
    for pi, frags in enumerate(pages):
        band = header_band(frags)
        if not band:
            continue
        top = min(f["y"] for f in band.values()) - 6
        stop = max([f["y"] for f in frags if _PDF_SKIP.search(f["t"]) and f["y"] < top] or [-1e9])
        page_body = [f for f in frags if stop + 1 < f["y"] < top and not _PDF_SKIP.search(f["t"])]
        body.append(page_body)
        for f in page_body:
            if f["x"] > amt_hi - 60 and not is_amount(f):
                starts[round(f["x"] / 4)] = starts.get(round(f["x"] / 4), 0) + 1
    col_x = sorted({k * 4 for k, n in starts.items() if n >= 2})
    clusters = []
    for x in col_x:
        if clusters and x - clusters[-1][-1] <= 8:
            clusters[-1].append(x)
        else:
            clusters.append([x])
    text_cols = [min(c) for c in clusters]
    # gán cột chữ cho tiêu đề gần nhất (theo thứ tự trái → phải)
    col_kind = {}
    for hx, k in text_hdrs:
        free = [c for c in text_cols if c not in col_kind and c <= hx + 20]
        if free:
            col_kind[max(free, key=lambda c: (c <= hx + 20, -abs(c - hx)))] = k
    if text_hdrs and "desc" not in col_kind.values() and text_cols:
        col_kind[text_cols[0]] = "desc"

    def text_col(x):
        own = [c for c in text_cols if c <= x + 3]
        return col_kind.get(max(own)) if own else None
    for page_body in body:
        anchors = []
        for f in sorted((f for f in page_body if is_amount(f)), key=lambda f: -f["y"]):
            if anchors and abs(anchors[-1]["y"] - f["y"]) <= 2:
                anchors[-1]["amts"].append(f)
            else:
                anchors.append({"y": f["y"], "amts": [f], "frags": []})
        if not anchors:
            continue
        for f in page_body:
            if is_amount(f):
                continue
            a = min(anchors, key=lambda a: abs(a["y"] - f["y"]))
            if abs(a["y"] - f["y"]) <= 45:
                a["frags"].append(f)
        for a in anchors:
            amts = sorted(a["amts"], key=lambda f: f["x"])
            vals = {}
            if len(amts) == len(amount_hdr):
                for (hx, k), f in zip(amount_hdr, amts):
                    vals[k] = _bank_amount(f["t"])
            else:
                for f in amts:
                    k = min(amount_hdr, key=lambda h: abs(h[0] - f["x"]))[1]
                    vals[k] = _bank_amount(f["t"])
            dates = [f for f in a["frags"] if re.search(r"\d{1,2}/\d{1,2}/\d{4}", f["t"])]
            if not dates:
                continue
            dfrag = min(dates, key=lambda f: abs(f["x"] - date_x))
            d = _bank_date(re.search(r"\d{1,2}/\d{1,2}/\d{4}", dfrag["t"]).group(0))
            tin, tout = abs(vals.get("in", 0.0)), abs(vals.get("out", 0.0))
            if not d or (not tin and not tout):
                continue
            lines = {"desc": [], "name": []}
            ref = []
            for f in sorted(a["frags"], key=lambda f: (-f["y"], f["x"])):
                if f in dates or re.fullmatch(r"\d{1,2}:\d{2}(:\d{2})?", f["t"]):
                    continue
                if ref_x is not None and f["x"] < amt_lo and abs(f["x"] - ref_x) <= 30:
                    ref.append(f["t"])
                    continue
                k = text_col(f["x"]) if f["x"] >= amt_lo else None
                if k in lines:
                    lines[k].append(f["t"])
            clean = lambda parts: re.sub(r"\s+", " ", " ".join(parts)).strip()
            info["rows"].append({"date": d.strftime("%Y-%m-%d"), "time": "", "ref": "".join(ref).replace(" ", ""),
                                 "desc": clean([_unchunk35(" ".join(lines["desc"]))]), "name": clean(lines["name"]),
                                 "in": tin, "out": tout})
    # Tổng phát sinh (số có thể bị ngắt 2 dòng: "693,350,3" + "85")
    for frags in pages:
        lab = next((f for f in frags if re.search(r"^(tongphatsinh|tongcong|total)", _bank_plain(f["t"]))), None)
        if not lab:
            continue
        near = sorted((f for f in frags if abs(f["y"] - lab["y"]) <= 12 and f["x"] > lab["x"]
                       and re.fullmatch(r"[\d.,]+", f["t"])), key=lambda f: (-f["y"], f["x"]))
        groups = []  # số in 2 dòng: phần sau nằm dòng dưới, lệch phải
        for f in near:
            prev = [g for g in groups if g["x"] <= f["x"] + 2 and g["y"] > f["y"] + 2]
            if prev:
                max(prev, key=lambda g: g["x"])["t"] += f["t"]
            else:
                groups.append(dict(f))
        nums = [(g["x"], _bank_amount(g["t"])) for g in sorted(groups, key=lambda g: g["x"])]
        if len(nums) >= 2:
            for (hx, k), (x, v) in zip(amount_hdr, nums[:2]):
                info["total_" + ("in" if k == "in" else "out")] = v
    return info


def bank_ktsc_rows(rows, opts=None):
    """Giao dịch → các dòng KTSC. Mặc định giống Nibot: chỉ ghi TK ngân hàng, để trống TK đối ứng.
    opts: tk (1121), ma_dt (mã đối tượng của TK ngân hàng), tk_thu / tk_chi (TK đối ứng tiền vào / tiền ra),
    nghiep_vu (TIENHANG), ghi_chu."""
    o = opts or {}
    tk = str(o.get("tk") or "1121").strip()
    ma = str(o.get("ma_dt") or "").strip()
    tk_thu, tk_chi = str(o.get("tk_thu") or "").strip(), str(o.get("tk_chi") or "").strip()
    out = []
    for x in rows:
        incoming = x["in"] > 0
        amt = x["in"] if incoming else x["out"]
        amt = int(amt) if float(amt).is_integer() else amt
        out.append({"LCTG": "CTNH", "NGAYCT": datetime.strptime(x["date"], "%Y-%m-%d"), "SOCT": x.get("ref") or "",
                    "DIENGIAI": x.get("desc") or "",
                    "TKNO": tk if incoming else (str(x.get("tkdu") or "").strip() or tk_chi),
                    "MADTPNNO": ma if incoming else "",
                    "TKCO": (str(x.get("tkdu") or "").strip() or tk_thu) if incoming else tk,
                    "MADTPNCO": "" if incoming else ma,
                    "TTVND": amt, "TTVND_TT": amt, "TENKH": x.get("name") or "",
                    "ID_NGHIEPVU": str(o.get("nghiep_vu") or "TIENHANG"), "GHICHU": str(o.get("ghi_chu") or "TAIHOADON IMPORT"),
                    "GUID": ""})
    return out


def merge_bank_statements(infos):
    """Gộp nhiều sao kê, bỏ giao dịch trùng (sao kê tháng sau hay lặp giao dịch cuối tháng trước)."""
    seen, rows, dups = set(), [], 0
    for info in infos:
        for x in info["rows"]:
            k = (info.get("account"), x["date"], x["time"], x["ref"], x["in"], x["out"], x["desc"])
            if k in seen:
                dups += 1
                continue
            seen.add(k)
            rows.append(dict(x, account=info.get("account", ""), bank=info.get("bank", "")))
    return rows, dups


def write_bank_ktsc(rows, opts=None):
    """→ bytes file .xlsx sheet KTSC (14 cột như file Nibot)."""
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        raise ValueError("Chưa cài openpyxl (pip install openpyxl)")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KTSC"
    ws.append(list(KTSC_BANK_COLS))
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in bank_ktsc_rows(rows, opts):
        ws.append([r[k] for k in KTSC_BANK_COLS])
        ws.cell(ws.max_row, 2).number_format = "dd/mm/yyyy"
        for col in (9, 10):
            ws.cell(ws.max_row, col).number_format = "#,##0"
    for col, w in zip("ABCDEFGHIJKLMN", (7, 11, 20, 70, 7, 12, 7, 12, 14, 14, 40, 12, 16, 6)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Giao diện web chạy trên máy (127.0.0.1)
# ---------------------------------------------------------------------------

LOGIN_PAGE = r"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Đăng nhập – Tải hoá đơn</title>
<style>
:root{--bg:#f4f6f9;--card:#fff;--fg:#1d2733;--mute:#66758a;--line:#dde3ea;--pri:#2b3a8c;--acc:#5b5fe0;--err:#c0392b}
@media (prefers-color-scheme:dark){:root{--bg:#12161c;--card:#1b2129;--fg:#e6ebf1;--mute:#93a1b3;--line:#2c3540;--pri:#3a4bb0;--acc:#8b8ff5;--err:#ff6b5b}}
*{box-sizing:border-box}body{margin:0;min-height:100vh;display:flex;align-items:center;justify-content:center;padding:16px;
font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;background:var(--bg);color:var(--fg)}
form{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:28px;width:min(380px,100%)}
h1{font-size:22px;margin:0 0 4px;color:var(--pri)}p{margin:0 0 18px;color:var(--mute);font-size:13px}
label{display:block;font-size:13px;color:var(--mute);margin:12px 0 4px}
input{width:100%;padding:10px 12px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--fg);font:inherit}
button{width:100%;margin-top:18px;padding:11px;border:0;border-radius:8px;background:var(--acc);color:#fff;font:inherit;font-weight:600;cursor:pointer}
button:disabled{opacity:.6}.err{color:var(--err);font-size:13px;margin-top:10px;min-height:1em}
</style></head><body>
<form id="f"><h1>Tải hoá đơn điện tử</h1><p>Đăng nhập để làm việc với các doanh nghiệp được giao.</p>
<label for="u">Tên đăng nhập</label><input id="u" autocomplete="username" autofocus required>
<label for="p">Mật khẩu</label><input id="p" type="password" autocomplete="current-password" required>
<button id="b">Đăng nhập</button><div class="err" id="e"></div></form>
<script>
document.getElementById('f').onsubmit = async ev => {
  ev.preventDefault(); const b = document.getElementById('b'), e = document.getElementById('e'); b.disabled = true; e.textContent = '';
  try {
    const r = await fetch('/api/login', {method: 'POST', headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({username: document.getElementById('u').value, password: document.getElementById('p').value})});
    const d = await r.json().catch(() => ({}));
    if (r.ok && d.ok) { location.reload(); return; }
    e.textContent = d.error || ('Lỗi ' + r.status);
  } catch (x) { e.textContent = 'Không kết nối được máy chủ'; }
  b.disabled = false;
};
</script></body></html>
"""


PAGE = r"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Tải hoá đơn điện tử</title>
<style>
:root{--bg:#f4f6f9;--card:#fff;--fg:#1d2733;--mute:#66758a;--line:#dde3ea;--pri:#2b3a8c;--acc:#5b5fe0;--err:#c0392b;--ok:#1e8449;--v:#4361ee;--r:#d63447}
@media (prefers-color-scheme:dark){:root{--bg:#12161c;--card:#1b2129;--fg:#e6ebf1;--mute:#93a1b3;--line:#2c3540;--pri:#3a4bb0;--acc:#8b8ff5;--err:#ff6b5b;--ok:#4cc38a;--v:#8aa2ff;--r:#ff7b8a}}
*{box-sizing:border-box}body{margin:0;font:14px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;background:var(--bg);color:var(--fg)}
header{background:var(--pri);color:#fff;padding:10px 16px;display:flex;gap:16px;align-items:center;flex-wrap:wrap}
header b{font-size:20px;letter-spacing:.5px}header small{opacity:.8}
main{max-width:1400px;margin:0 auto;padding:16px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px;margin-bottom:14px}
.top{display:flex;gap:12px;flex-wrap:wrap;align-items:center;justify-content:space-between}
.stat{color:var(--fg);font-size:15px}.stat b{color:var(--acc)}
input[type=text],input[type=search],input[type=password],input[type=date],textarea{padding:8px 10px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg);font:inherit}
input.search{flex:1;min-width:220px;background:#fffbe6;color:#333}
button{padding:7px 14px;border:1px solid var(--acc);border-radius:6px;background:var(--acc);color:#fff;font:inherit;cursor:pointer}
button.sec{background:transparent;color:var(--acc)}button.red{background:transparent;color:var(--err);border-color:var(--err)}
button:disabled{opacity:.5;cursor:default}button.sm{padding:3px 9px;font-size:12px}
.bar{display:flex;gap:8px;flex-wrap:wrap;margin:10px 0}
table{border-collapse:collapse;width:100%}th,td{padding:8px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}
th{font-size:13px;color:var(--mute);font-weight:600}td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}
.dn{font-weight:600}.dn .m{color:var(--acc);font-weight:400}.loi{color:var(--err);font-size:12px}
.v{color:var(--v);font-weight:600}.r{color:var(--r);font-weight:600}
.note{border:0;background:transparent;width:100%;color:var(--mute);font-size:13px;padding:2px}
.note:focus{background:var(--bg);color:var(--fg)}.tbl{overflow:auto}
.modal{position:fixed;inset:0;background:rgba(0,0,0,.45);display:flex;align-items:center;justify-content:center;padding:16px;z-index:9}
.modal .card{width:min(560px,100%);max-height:90vh;overflow:auto;margin:0}
label{display:block;font-size:13px;color:var(--mute);margin:8px 0 3px}.full{width:100%}
.chk{display:flex;gap:16px;flex-wrap:wrap;margin:10px 0}.chk label{display:flex;gap:6px;align-items:center;color:var(--fg);margin:0}
.capimg{background:#fff;border-radius:6px;display:inline-block;padding:2px}.capimg svg{height:50px;width:auto;display:block}
.prog{height:6px;background:var(--line);border-radius:3px;overflow:hidden;margin:8px 0}.prog i{display:block;height:100%;background:var(--acc);width:0}
pre{background:var(--bg);border:1px solid var(--line);border-radius:6px;padding:8px;max-height:200px;overflow:auto;font-size:12px;margin:0;white-space:pre-wrap}
.hide{display:none!important}.mute{color:var(--mute);font-size:12px}.err{color:var(--err)}.ok{color:var(--ok)}
.hint{font-size:12px;color:var(--mute);margin-top:6px}
nav{display:flex;gap:4px;margin-left:12px}nav a{color:#fff;opacity:.7;text-decoration:none;padding:6px 12px;border-radius:6px;font-weight:600}
nav a.on{opacity:1;background:rgba(255,255,255,.15)}
.flt{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:8px;margin-bottom:8px}
.flt select,.flt input,select.sm{padding:6px 8px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg);font:inherit;width:100%}
#hTbl td,#hTbl th{font-size:13px;padding:6px}#hTbl select{min-width:112px}#hTbl td.mh{max-width:260px;white-space:normal}
#hTbl tfoot td{font-weight:700;color:var(--acc)}.pager{display:flex;gap:6px;align-items:center;flex-wrap:wrap;margin-top:8px}
.pager button{padding:3px 9px}
.tools{display:grid;grid-template-columns:200px 1fr;gap:20px}@media(max-width:800px){.tools{grid-template-columns:1fr}}
.tmenu a{display:block;padding:10px 12px;border-radius:8px;color:var(--fg);text-decoration:none;margin-bottom:4px;background:var(--bg)}
.tmenu a.on{background:var(--pri);color:#fff}.tpane h3{margin:0 0 10px;color:var(--acc)}#mergeList li{margin:4px 0}
.sync{display:grid;grid-template-columns:minmax(260px,360px) 1fr;gap:20px}@media(max-width:800px){.sync{grid-template-columns:1fr}}
.sync label{margin-top:10px}.sync input,.sync select{width:100%}.sync .chk label{margin:0}.sync .chk input{width:auto}.bigbtn{display:block;width:100%;margin-top:8px;padding:10px;font-weight:600;border:0}
.b-vao{background:#4361ee}.b-ra{background:#e5534b}.b-vr{background:#8c0f9d}.advice{background:#eef0ff;border:1px solid #c9ceff;border-radius:8px;padding:12px;margin-top:12px;color:#222;font-size:14px;line-height:1.7}
.advice h4{color:#5503ca;margin:0 0 6px}.badge{display:inline-block;padding:2px 8px;border-radius:5px;color:#fff;font-size:12px;background:#6c757d}
.badge.run{background:#f0a500}.badge.ok{background:#1e8449}.badge.bad{background:#c0392b}.steps{line-height:1.9;margin:10px 0}
.staff .adm,.local .srv{display:none!important}
.pw{-webkit-text-security:disc;text-security:disc}
/* Trang doanh nghiệp */
.wshead{padding-bottom:0;position:sticky;top:0;z-index:5}.wshead .crumb{margin-bottom:6px}
.wsrow{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
#wsPick{flex:1;min-width:260px;font-weight:600;font-size:15px;padding:8px 12px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--fg)}
.wsmeta{font-size:12px;color:var(--mute);display:flex;gap:12px;flex-wrap:wrap}.wsmeta b{color:var(--fg)}
.subtabs{display:flex;gap:2px;margin-top:10px;overflow-x:auto}
.subtabs a{padding:9px 16px;font-weight:700;font-size:13px;letter-spacing:.3px;text-transform:uppercase;color:var(--mute);
  text-decoration:none;border-bottom:3px solid transparent;white-space:nowrap}
.subtabs a.on{color:var(--acc);border-bottom-color:var(--acc)}.subtabs a:hover{color:var(--fg)}
.filters{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-bottom:8px}
.filters select,.filters input{padding:7px 9px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg);font:inherit}
.filters .grow{flex:1;min-width:200px}
.seg{display:inline-flex;border:1px solid var(--line);border-radius:8px;overflow:hidden}
.seg button{border:0;border-radius:0;background:transparent;color:var(--fg);padding:7px 12px}
.seg button+button{border-left:1px solid var(--line)}.seg button.on{background:var(--acc);color:#fff}
.bulk{display:flex;gap:8px;flex-wrap:wrap;align-items:center;background:#eef0ff;border:1px solid #c9ceff;color:#222;border-radius:8px;padding:8px 10px;margin:6px 0}
.bulk select{width:auto}
.grid{overflow:auto;max-height:calc(100vh - 300px);min-height:240px;border:1px solid var(--line);border-radius:8px}
.grid table{border-collapse:separate;border-spacing:0}
.grid th,.grid td{font-size:13px;padding:6px 7px;border-bottom:1px solid var(--line);background:var(--card)}
.grid thead th{position:sticky;top:0;z-index:2;background:var(--bg);color:var(--acc);white-space:nowrap;user-select:none}
.grid thead tr.f th{top:var(--hh,32px);padding:4px;font-weight:400}
.grid thead th.s{cursor:pointer}.grid thead th.s:hover{color:var(--fg)}
.grid thead tr.f input,.grid thead tr.f select{width:100%;min-width:52px;padding:4px 6px;font-size:12px;border:1px solid var(--line);border-radius:5px;background:var(--card);color:var(--fg)}
.grid tfoot td{position:sticky;bottom:0;z-index:2;background:var(--bg);font-weight:700;color:var(--acc);border-top:2px solid var(--line)}
.grid tbody tr{cursor:default}.grid tbody tr:hover td{background:rgba(91,95,224,.06)}
.grid tbody tr.sel td{background:rgba(91,95,224,.14)}.grid tbody tr.sel td:first-child{box-shadow:inset 3px 0 var(--acc)}
.grid tbody tr.bad td{color:var(--err)}.grid td.mh{max-width:280px;white-space:normal;min-width:180px}
.grid td.ten{max-width:260px;min-width:160px;white-space:normal}.grid td.nw{white-space:nowrap}
.grid select{min-width:108px;padding:3px 4px;font-size:12px}
a.dn-link{color:inherit;text-decoration:none}a.dn-link:hover{color:var(--acc);text-decoration:underline}
.me{display:flex;gap:8px;align-items:center;font-size:13px}.me a{color:#fff;opacity:.85;text-decoration:none;cursor:pointer}.me a:hover{opacity:1;text-decoration:underline}
.bell{position:relative;background:transparent;border:0;color:#fff;font-size:18px;padding:2px 6px}.bell i{position:absolute;top:-4px;right:-4px;background:#e5534b;color:#fff;
font-style:normal;font-size:11px;border-radius:9px;padding:0 5px;min-width:16px;text-align:center}
#notes{position:absolute;right:16px;top:52px;width:min(520px,calc(100vw - 32px));max-height:70vh;overflow:auto;z-index:8;box-shadow:0 8px 30px rgba(0,0,0,.25)}
#notes .it{padding:8px 0;border-bottom:1px solid var(--line);font-size:13px}#notes .it.new{font-weight:600}#notes .warn{color:#b26a00}#notes .err2{color:var(--err)}
.ulist{max-height:260px;overflow:auto;border:1px solid var(--line);border-radius:6px;padding:6px}.ulist label{display:flex;gap:6px;color:var(--fg);margin:2px 0;font-size:13px}
.crumb{font-size:13px;color:var(--mute);margin-bottom:10px}.crumb b{color:var(--fg)}#sTbl td,#sTbl th{font-size:13px;padding:5px}.pager .cur{background:var(--err);border-color:var(--err)}a.lk{color:var(--acc);margin-right:6px}
</style></head><body>
<header><b>Tải hoá đơn</b><nav><a href="#" id="navDN" class="on" onclick="tab('DN');return false">Doanh nghiệp</a>
<a href="#" id="navHD" onclick="openCompany(wsMst || (companies.find(c => !c.an) || {}).mst);return false">Hoá đơn</a>
<a href="#" id="navTI" onclick="tab('TI');return false">Tiện ích</a>
<a href="#" id="navND" class="adm srv" onclick="tab('ND');return false">Quản trị</a></nav><span style="flex:1"></span><small id="capStat"></small>
<span class="me srv"><button class="bell" id="bell" title="Thông báo" onclick="toggleNotes()">&#128276;<i id="bellN" class="hide"></i></button>
<span id="meName"></span> · <a onclick="openPw()">Đổi mật khẩu</a> · <a onclick="logout()">Đăng xuất</a></span></header>
<div class="card hide" id="notes"><div class="top"><b>Thông báo</b><button class="sm sec" onclick="toggleNotes()">Đóng</button></div><div id="notesList"></div></div>
<main>
<section class="card hide" id="tabND">
  <div class="top"><b>Người dùng</b><button onclick="openUser()">+ Thêm người dùng</button></div>
  <div class="hint">Quản trị xem và làm việc với tất cả doanh nghiệp, quản lý người dùng. Nhân viên chỉ thấy các doanh nghiệp được giao.</div>
  <div class="tbl"><table><thead><tr><th>Tên đăng nhập</th><th>Họ tên</th><th>Vai trò</th><th class="n">Số DN được giao</th><th>Trạng thái</th><th>Ngày tạo</th><th></th></tr></thead>
    <tbody id="uRows"></tbody></table></div>
  <div class="err" id="uErr"></div>
  <h3 style="margin:24px 0 6px">Đồng bộ tự động hằng ngày</h3>
  <div class="hint">Máy chủ tự đồng bộ mua vào + bán ra cho tất cả doanh nghiệp đang hiển thị (có mật khẩu) theo giờ đã đặt, tự giải captcha.
    HĐ mới, HĐ đổi trạng thái (huỷ, thay thế, điều chỉnh) và lỗi được báo ở chuông thông báo cho người phụ trách doanh nghiệp.</div>
  <div class="chk" style="align-items:center">
    <label><input type="checkbox" id="aOn"> Bật</label>
    <label>Giờ chạy <input type="time" id="aGio" value="06:00" style="width:auto"></label>
    <label>Kỳ <select id="aKy" class="sm" style="width:auto"></select></label>
    <button class="sm" onclick="saveAuto()">Lưu</button><button class="sm sec" onclick="runAuto()">Chạy ngay</button><span class="ok" id="aMsg"></span>
  </div>
  <div class="mute" id="aInfo"></div><pre id="aLog" class="hide" style="max-height:180px"></pre>
</section>
<section class="card hide" id="tabTI">
  <div class="tools">
    <div class="tmenu">
      <a href="#" data-t="xml" class="on">Đọc hoá đơn XML</a><a href="#" data-t="mst">Kiểm tra MST DN</a>
      <a href="#" data-t="merge">Nối file PDF</a><a href="#" data-t="split">Tách file PDF</a>
      <a href="#" data-t="bank">Sao kê ngân hàng</a>
    </div>
    <div>
      <div class="tpane" id="t-xml"><h3>Đọc và xem hoá đơn XML</h3>
        Chọn file hoá đơn XML: <input type="file" id="xmlFile" accept=".xml,text/xml">
        <button class="sec sm hide" id="xmlPrint" onclick="$('xmlFrame').contentWindow.print()">In / Lưu PDF</button>
        <div class="err" id="xmlErr"></div>
        <iframe id="xmlFrame" class="hide" style="width:100%;height:75vh;border:1px solid var(--line);border-radius:8px;margin-top:10px;background:#fff"></iframe></div>
      <div class="tpane hide" id="t-mst"><h3>Kiểm tra thông tin / tình trạng MST</h3>
        <textarea id="mstText" rows="5" class="full" placeholder="Dán danh sách MST, mỗi dòng một MST (tối đa 200)"></textarea>
        <div class="bar"><button onclick="toolMst()">Kiểm tra</button></div><div class="err" id="mstErr"></div>
        <div class="tbl"><table id="mstTbl"></table></div></div>
      <div class="tpane hide" id="t-merge"><h3>Nối nhiều file PDF thành một</h3>
        <input type="file" id="mergeFiles" accept="application/pdf" multiple>
        <div class="hint">Chọn các file theo đúng thứ tự muốn nối (có thể chọn thêm nhiều lần; kéo thứ tự bằng nút ↑ ↓).</div>
        <ol id="mergeList"></ol><div class="bar"><button onclick="toolMerge()">Nối file</button>
        <button class="sec" onclick="mergeQueue=[];renderMerge()">Xoá danh sách</button></div><div class="err" id="mergeErr"></div></div>
      <div class="tpane hide" id="t-split"><h3>Tách file PDF</h3>
        <input type="file" id="splitFile" accept="application/pdf">
        <label>Khoảng trang (để trống = mỗi trang một file)</label><input type="text" id="splitRanges" placeholder="vd: 1-3, 4, 5-8">
        <div class="bar"><button onclick="toolSplit()">Tách file</button></div><div class="err" id="splitErr"></div></div>
      <div class="tpane hide" id="t-bank"><h3>Sao kê ngân hàng → file nhập phần mềm kế toán (KTSC)</h3>
        <div class="hint">Chọn một hoặc nhiều file sao kê / sổ phụ tải từ internet banking (.xls, .xlsx, .csv, hoặc sổ phụ .pdf có chữ – không nhận bản scan). Phần mềm tự tìm bảng giao dịch,
          đối chiếu với tổng ghi có / ghi nợ và số dư trên sao kê, bỏ giao dịch trùng giữa các tháng, rồi xuất Excel sheet KTSC
          theo mẫu Nibot để nhập Smart Pro. Tiền vào: Nợ TK ngân hàng; tiền ra: Có TK ngân hàng.</div>
        <input type="file" id="bankFiles" multiple accept=".xls,.xlsx,.csv,.htm,.html,.txt,.pdf" style="margin-top:8px">
        <div class="flt" style="margin-top:10px">
          <label>TK ngân hàng<input type="text" id="bkTk" value="1121"></label>
          <label>Mã đối tượng TK ngân hàng<input type="text" id="bkMa" placeholder="vd 112_VIETTIN"></label>
          <label>TK đối ứng tiền vào<input type="text" id="bkThu" placeholder="để trống như Nibot, vd 131"></label>
          <label>TK đối ứng tiền ra<input type="text" id="bkChi" placeholder="để trống như Nibot, vd 331"></label>
          <label>Mã nghiệp vụ<input type="text" id="bkNv" value="TIENHANG"></label>
          <label>Ghi chú<input type="text" id="bkGc" value="NIBOT IMPORT"></label>
        </div>
        <div class="err" id="bankErr"></div><div id="bankInfo"></div>
        <div class="bar"><button id="bankExp" class="hide" onclick="bankExport()">Xuất Excel KTSC</button>
          <span class="mute" id="bankSum"></span></div>
        <div class="tbl"><table id="bankTbl"></table></div></div>
    </div>
  </div>
</section>
<section class="hide" id="tabWS">
  <div class="card wshead">
    <div class="crumb"><a href="#/" onclick="tab('DN');return false">Doanh nghiệp</a> › <b id="wsCrumb"></b></div>
    <div class="wsrow">
      <button class="sec sm" id="wsPrev" title="Doanh nghiệp trước (Alt+←)" onclick="stepCompany(-1)">‹</button>
      <input type="search" id="wsPick" list="wsList" autocomplete="off" placeholder="Gõ tên hoặc MST để chuyển doanh nghiệp…">
      <datalist id="wsList"></datalist>
      <button class="sec sm" id="wsNext" title="Doanh nghiệp sau (Alt+→)" onclick="stepCompany(1)">›</button>
      <button class="sec sm" title="Sửa thông tin doanh nghiệp" onclick="openEdit(companies.find(c => c.mst === wsMst))">✎ Sửa</button>
      <span class="wsmeta" id="wsMeta"></span>
    </div>
    <input type="hidden" id="hMst">
    <nav class="subtabs" id="wsTabs">
      <a href="#" data-s="hoa-don">Hoá đơn</a><a href="#" data-s="dong-bo">Đồng bộ</a>
      <a href="#" data-s="sao-ke">Sao kê ng.hàng</a><a href="#" data-s="tra-cuu-mst">Tra cứu MST</a>
    </nav>
  </div>

  <div class="card" id="wsHD">
    <div class="filters">
      <div class="seg" id="hKindSeg"><button data-k="purchase" class="on">Mua vào</button><button data-k="sold">Bán ra</button>
        <button data-k="purchase_dv" title="Hoá đơn mua vào đã đánh dấu dịch vụ">Mua vào – HĐDV</button><button data-k="sold_dv">Bán ra – HĐDV</button></div>
      <input type="hidden" id="hKind" value="purchase">
      <select id="hPer"></select><input type="date" id="hFrom" title="Từ ngày"><input type="date" id="hTo" title="Đến ngày">
      <input id="hQ" type="search" autocomplete="off" placeholder="Tìm MST, tên, mặt hàng, ghi chú…" class="grow">
      <button onclick="hPage=0;loadInv()">Tìm kiếm</button>
      <button class="sec" id="hMoreBtn" onclick="$('hMore').classList.toggle('hide')">Bộ lọc khác ▾</button>
    </div>
    <div class="filters hide" id="hMore">
      <select id="hFile"><option value="">Tất cả file</option><option value="no_xml">Không có XML</option><option value="xml">Có XML</option>
        <option value="pdf">Có PDF</option><option value="no_pdf">Không có PDF</option><option value="xml_pdf">Có XML và PDF</option></select>
      <select id="hDuyet"><option value="">Duyệt nội bộ: tất cả</option><option>Chờ duyệt</option><option>Đã duyệt</option><option>Không duyệt</option></select>
      <select id="hTthai"><option value="">Trạng thái HĐ: tất cả</option></select>
      <select id="hKq"><option value="">Kết quả kiểm tra: tất cả</option></select>
      <input id="hKh" autocomplete="off" placeholder="Ký hiệu HĐ"><input id="hSo" autocomplete="off" placeholder="Số HĐ">
    </div>
    <div class="bar actions">
      <button class="sec" onclick="openSub('dong-bo')">⟳ Đồng bộ</button>
      <button class="sec" id="hBulkPdf" onclick="bulkPdf()">Tải HĐ gốc hàng loạt</button>
      <button class="sec" id="hImpPdf" onclick="$('hImpFiles').click()" title="Chọn các file PDF hoá đơn gốc có sẵn trên máy – phần mềm tự gắn vào đúng hoá đơn">Gắn PDF gốc có sẵn</button>
      <input type="file" id="hImpFiles" accept="application/pdf" multiple class="hide">
      <span style="flex:1"></span>
      <button class="sec sm" onclick="resetGrid()" title="Bỏ sắp xếp và lọc theo cột">Bỏ lọc cột</button>
      <select class="sm" id="hExp" style="width:auto" onchange="exportInv(this.value);this.value=''">
        <option value="">Kết xuất…</option><option value="xlsx">Excel (mẫu Nibot)</option><option value="xml">XML.ZIP</option>
        <option value="html">HTML.ZIP</option><option value="pdf">PDF gốc.ZIP</option><option value="cqt">PDF thuế.ZIP</option></select>
    </div>
    <div class="bulk hide" id="hBulk">
      <b id="hSelInfo"></b>
      <select class="sm" id="hBulkDuyet" onchange="bulkUpdate({duyet: this.value});this.value=''">
        <option value="">Duyệt nội bộ…</option><option>Đã duyệt</option><option>Chờ duyệt</option><option>Không duyệt</option></select>
      <button class="sm sec" onclick="bulkUpdate({dv: true})">Đánh dấu HĐ dịch vụ</button>
      <button class="sm sec" onclick="bulkUpdate({dv: false})">Bỏ HĐ dịch vụ</button>
      <button class="sm sec" onclick="bulkPdf()">Tải PDF gốc</button>
      <button class="sm sec" onclick="exportInv('xlsx')">Kết xuất Excel</button>
      <button class="sm red" onclick="hSel.clear();renderInv()">Bỏ chọn</button>
    </div>
    <div class="err" id="hErr"></div>
    <div class="grid"><table id="hTbl"><thead id="hHead"></thead><tbody id="hRows"></tbody><tfoot><tr id="hFoot"></tr></tfoot></table></div>
    <div class="pager" id="hPager"></div>
  </div>

  <div class="card hide" id="wsSYNC">
  <div class="sync">
    <div>
      <label>Chọn khoảng thời gian</label><div style="display:flex;gap:8px"><select id="sPer"></select>
        <input type="number" id="sYear" style="width:90px"></div>
      <label>Từ ngày</label><input type="date" id="sFrom"><label>Đến ngày</label><input type="date" id="sTo">
      <div class="chk"><label><input type="checkbox" id="sMtt" checked> Gồm máy tính tiền</label>
        <label><input type="checkbox" id="sXml" checked> Tải XML</label></div>
      <button class="bigbtn b-vao" onclick="runSync(['purchase'])">Đồng bộ HĐĐT ĐẦU VÀO</button>
      <button class="bigbtn b-ra" onclick="runSync(['sold'])">Đồng bộ HĐĐT ĐẦU RA</button>
      <button class="bigbtn b-vr" onclick="runSync(['purchase','sold'])">Đồng bộ HĐĐT VÀO/RA</button>
      <div class="advice"><h4>KHUYẾN CÁO SỬ DỤNG HIỆU QUẢ</h4>
        <b>Nên bấm Đồng bộ trước khi kiểm tra số liệu vì:</b><br>
        1. Có hoá đơn người bán ký trễ, hoặc hoá đơn không mã (ký hiệu K) gửi lên cơ quan thuế trễ 15–45 ngày – đồng bộ lại để lấy <b>đủ số lượng hoá đơn</b>.<br>
        2. Hoá đơn tải về là HĐ mới nhưng sau đó người bán huỷ / điều chỉnh / thay thế – đồng bộ lại để có <b>trạng thái mới nhất</b>.<br>
        <b style="color:#1f3fb0">Hoá đơn đã tải rồi sẽ không tải lại XML, nên đồng bộ nhiều lần không tốn thời gian.</b></div>
    </div>
    <div>
      <div>Trạng thái: <span class="badge" id="sBadge">chờ đồng bộ</span> <button class="red sm hide" id="sStop" onclick="post('/api/stop',{})">Dừng</button></div>
      <div id="sCapBox" class="card hide" style="border-color:var(--acc);margin-top:10px">
        <div>Nhập captcha:</div>
        <div style="display:flex;gap:8px;align-items:center;flex-wrap:wrap;margin-top:6px">
          <span class="capimg" id="sCapImg"></span><input type="text" id="sCapIn" placeholder="ký tự trong ảnh" autocomplete="off">
          <button onclick="sendCap('sCapIn')">Gửi</button></div>
        <div class="hint">Chỉ cần nhập khi đăng nhập lần đầu hoặc khi phiên đã hết hạn.</div>
      </div>
      <div class="steps" id="sSteps"></div>
      <div class="prog"><i id="sProg"></i></div>
      <div class="bar"><span style="flex:1"></span><button class="sec" onclick="openInvoicesFromSync()">Xem hoá đơn</button></div>
      <div class="tbl" style="max-height:460px"><table id="sTbl"><thead><tr><th>Loại HĐ</th><th>Ngày lập</th><th>Mẫu số</th>
        <th>Ký hiệu</th><th class="n">Số HĐ</th><th>Loại đồng bộ</th><th>Kết quả đồng bộ</th></tr></thead><tbody id="sRows"></tbody></table></div>
    </div>
  </div>
</div>

  <div class="card hide" id="wsTOOL"></div>
</section>
<section class="card" id="tabDN">
  <div class="top">
    <div class="stat">Số doanh nghiệp: <b id="nDN">0</b> hiển thị; <b id="nAn">0</b> bị ẩn
      <a href="#" id="toggleAn" class="mute">xem doanh nghiệp bị ẩn</a></div>
    <input class="search" id="q" type="search" name="tim-doanh-nghiep" autocomplete="off" placeholder="nhập MST hoặc tên doanh nghiệp cần làm việc nhanh">
  </div>
  <div class="bar">
    <button class="adm" onclick="openEdit()">+ Thêm doanh nghiệp</button>
    <button class="sec adm" onclick="show('mImport')">Nhập danh sách từ Excel</button>
    <button class="red" onclick="openBatch()">Xử lý hàng loạt</button>
  </div>
  <div class="tbl"><table>
    <thead><tr><th><input type="checkbox" id="all"></th><th>Doanh nghiệp</th><th>Ghi chú công việc</th>
      <th>Lịch sử đồng bộ</th><th>Ngày tạo</th><th class="n">Số HĐ<br>Mua vào</th><th class="n">Số HĐ<br>Bán ra</th>
      <th class="n">Tổng HĐ</th><th></th></tr></thead>
    <tbody id="rows"></tbody></table></div>
</section>

<section class="card hide" id="jobCard">
  <div class="top"><b id="jobTitle">Đang xử lý</b><button class="red sm" id="btnStop" onclick="post('/api/stop',{})">Dừng</button></div>
  <div class="prog"><i id="progI"></i></div>
  <div id="capBox" class="card hide" style="border-color:var(--acc)">
    <div>Nhập captcha cho <b id="capFor"></b>:</div>
    <div style="display:flex;gap:8px;align-items:center;flex-wrap:wrap;margin-top:6px">
      <span class="capimg" id="capImg"></span><input type="text" id="capIn" placeholder="ký tự trong ảnh" autocomplete="off">
      <button onclick="sendCap()">Gửi</button></div>
    <div class="hint">Phần mềm ghi nhớ hình từng ký tự bạn gõ; sau vài lần sẽ tự giải captcha.</div>
  </div>
  <div id="results"></div>
  <pre id="log"></pre>
</section>
</main>

<div class="modal hide" id="mEdit"><div class="card">
  <b id="editTitle">Doanh nghiệp</b>
  <label>Mã số thuế – tên đăng nhập trang hoadondientu.gdt.gov.vn</label>
  <div style="display:flex;gap:8px"><input type="text" id="eMst" style="flex:1"><button type="button" class="sec" id="eLookup" onclick="lookupMst()">Lấy tên DN</button></div>
  <div class="hint" id="eInfo"></div>
  <label>Tên doanh nghiệp (hiển thị trong bảng kê – theo giấy phép hoặc tên gợi nhớ)</label><input type="text" class="full" id="eTen">
  <label>Mật khẩu trang hoadondientu.gdt.gov.vn</label><input type="text" class="full pw" id="ePw" autocomplete="off" autocapitalize="off" spellcheck="false" data-lpignore="true">
  <div class="hint" id="ePwHint"></div>
  <div class="chk">
    <label><input type="checkbox" id="eVao"> Đồng bộ hoá đơn đầu vào</label>
    <label><input type="checkbox" id="eRa"> Đồng bộ hoá đơn đầu ra</label>
    <label><input type="checkbox" id="eAn"> Ẩn doanh nghiệp</label>
  </div>
  <div class="err" id="eErr"></div>
  <div class="bar"><button onclick="saveEdit()">Lưu</button><button class="sec" onclick="saveEdit(true)">Lưu &amp; Test đăng nhập</button>
    <span style="flex:1"></span><button class="red adm" id="eDel" onclick="delEdit()">Xoá</button><button class="sec" onclick="hide('mEdit')">Đóng</button></div>
</div></div>

<div class="modal hide" id="mUser"><div class="card">
  <b id="uTitle">Người dùng</b>
  <label>Tên đăng nhập</label><input type="text" class="full" id="uName" autocomplete="off">
  <label>Họ tên</label><input type="text" class="full" id="uTen">
  <label>Vai trò</label><select id="uRole" class="sm"><option value="staff">Nhân viên – chỉ các DN được giao</option><option value="admin">Quản trị – tất cả DN, quản lý người dùng</option></select>
  <label>Mật khẩu đăng nhập phần mềm</label><input type="text" class="full pw" id="uPw" autocomplete="off" autocapitalize="off" spellcheck="false" data-lpignore="true">
  <div class="hint" id="uPwHint"></div>
  <div class="chk"><label><input type="checkbox" id="uActive" checked> Đang hoạt động (bỏ tích để khoá tài khoản)</label></div>
  <div id="uMstBox"><label>Doanh nghiệp được giao (<span id="uCnt">0</span>)</label>
    <div style="display:flex;gap:6px;margin-bottom:6px"><input type="search" id="uQ" autocomplete="off" placeholder="lọc MST / tên" style="flex:1">
      <button class="sm sec" type="button" onclick="uPickAll(true)">Chọn hết đang lọc</button><button class="sm sec" type="button" onclick="uPickAll(false)">Bỏ chọn</button></div>
    <div class="ulist" id="uList"></div></div>
  <div class="err" id="uEditErr"></div>
  <div class="bar"><button onclick="saveUser()">Lưu</button><span style="flex:1"></span>
    <button class="red" id="uDel" onclick="delUser()">Xoá</button><button class="sec" onclick="hide('mUser')">Đóng</button></div>
</div></div>

<div class="modal hide" id="mPw"><div class="card">
  <b>Đổi mật khẩu đăng nhập</b>
  <label>Mật khẩu hiện tại</label><input type="text" class="full pw" id="pOld" autocomplete="off" autocapitalize="off" spellcheck="false" data-lpignore="true">
  <label>Mật khẩu mới (tối thiểu 8 ký tự)</label><input type="text" class="full pw" id="pNew" autocomplete="off" autocapitalize="off" spellcheck="false" data-lpignore="true">
  <label>Nhập lại mật khẩu mới</label><input type="text" class="full pw" id="pNew2" autocomplete="off" autocapitalize="off" spellcheck="false" data-lpignore="true">
  <div class="err" id="pErr"></div>
  <div class="bar"><button onclick="savePw()">Đổi mật khẩu</button><button class="sec" onclick="hide('mPw')">Đóng</button></div>
</div></div>

<div class="modal hide" id="mImport"><div class="card">
  <b>Nhập danh sách doanh nghiệp</b>
  <div class="hint">Mỗi dòng: MST, Tên, Mật khẩu – copy 3 cột từ Excel dán vào (hoặc ngăn cách bằng dấu |). Để trống tên thì phần mềm tự lấy tên theo MST; mật khẩu có thể để trống.</div>
  <textarea id="impText" rows="10" class="full" style="margin-top:8px" placeholder="0313046002	BẢO LỘC KT305	matkhau"></textarea>
  <div class="err" id="impErr"></div>
  <div class="bar"><button onclick="doImport()">Nhập</button><button class="sec" onclick="hide('mImport')">Đóng</button></div>
</div></div>

<div class="modal hide" id="mBatch"><div class="card">
  <b>Xử lý hàng loạt</b>
  <div class="hint" id="bWho"></div>
  <div style="display:flex;gap:12px;flex-wrap:wrap">
    <div><label>Từ ngày</label><input type="date" id="bFrom"></div>
    <div><label>Đến ngày</label><input type="date" id="bTo"></div>
  </div>
  <div class="chk">
    <label><input type="checkbox" id="bVao" checked> Mua vào</label>
    <label><input type="checkbox" id="bRa" checked> Bán ra</label>
    <label><input type="checkbox" id="bMtt" checked> Gồm máy tính tiền</label>
    <label><input type="checkbox" id="bXml" checked> Tải XML + chi tiết hàng hoá</label>
  </div>
  <div class="hint">Chỉ đồng bộ chiều đầu vào/đầu ra đang bật trong cài đặt từng doanh nghiệp.</div>
  <div class="err" id="bErr"></div>
  <div class="bar"><button onclick="runBatch()">Bắt đầu</button><button class="sec" onclick="hide('mBatch')">Đóng</button></div>
</div></div>

<script>
const KEY = "__TOKEN__";
const $ = id => document.getElementById(id);
// Ô mật khẩu trong trang làm việc là ô chữ che bằng CSS: Chrome không coi trang là form đăng nhập nên không tự điền
// tên/mật khẩu đăng nhập web vào ô tìm kiếm và ô mật khẩu cổng thuế. Trình duyệt không hỗ trợ che → dùng ô mật khẩu.
if (!(window.CSS && (CSS.supports('-webkit-text-security', 'disc') || CSS.supports('text-security', 'disc'))))
  document.querySelectorAll('input.pw').forEach(i => { i.type = 'password'; i.autocomplete = 'new-password'; });
const show = id => $(id).classList.remove('hide'), hide = id => $(id).classList.add('hide');
const fmt = n => (n || 0).toLocaleString('en-US');
let me = {role: 'admin', server: false}, companies = [], showHidden = false, editing = null, batchMsts = [], timer = null;

async function api(path, body) {
  const r = await fetch(path, {method: body ? 'POST' : 'GET', headers: {'X-App-Key': KEY, 'Content-Type': 'application/json'},
                               body: body ? JSON.stringify(body) : undefined});
  const d = await r.json().catch(() => ({error: 'Phản hồi không hợp lệ'}));
  if (r.status === 401 && d.login) { location.reload(); throw new Error(d.error); }
  if (!r.ok || d.error) throw new Error(d.error || ('HTTP ' + r.status));
  return d;
}
const post = (p, b) => api(p, b);
function el(tag, cls, text) { const e = document.createElement(tag); if (cls) e.className = cls; if (text != null) e.textContent = text; return e; }

function render() {
  const q = $('q').value.trim().toLowerCase();
  const vis = companies.filter(c => showHidden ? c.an : !c.an)
    .filter(c => !q || c.mst.includes(q) || (c.ten || '').toLowerCase().includes(q));
  $('nDN').textContent = companies.filter(c => !c.an).length; $('nAn').textContent = companies.filter(c => c.an).length;
  $('toggleAn').textContent = showHidden ? 'xem doanh nghiệp đang hiển thị' : 'xem doanh nghiệp bị ẩn';
  const tb = $('rows'); tb.innerHTML = '';
  vis.forEach((c, i) => {
    const tr = tb.insertRow();
    const cb = el('input'); cb.type = 'checkbox'; cb.className = 'pick'; cb.value = c.mst; tr.insertCell().append(cb);
    const td = tr.insertCell(); const dn = el('a', 'dn dn-link', String(i + 1).padStart(3, '0') + '. ' + (c.ten || '(chưa đặt tên)') + ' ');
    dn.href = '#/dn/' + c.mst + '/hoa-don'; dn.title = 'Mở trang làm việc của doanh nghiệp';
    dn.onclick = ev => { ev.preventDefault(); openCompany(c.mst); };
    dn.append(el('span', 'm', '(' + c.mst + ')')); td.append(dn);
    if (!c.co_mk) td.append(el('div', 'loi', 'Chưa nhập mật khẩu'));
    if (c.loi) td.append(el('div', 'loi', c.loi));
    const note = el('input', 'note'); note.value = c.ghichu || ''; note.placeholder = 'nhập ghi chú công việc vào đây';
    note.onchange = () => post('/api/company/save', {mst: c.mst, ghichu: note.value}); tr.insertCell().append(note);
    const ls = tr.insertCell(), h = c.lich_su || {};
    if (h.purchase) ls.append(el('div', 'v', 'V: ' + h.purchase.moi + ' - ' + h.purchase.luc));
    if (h.sold) ls.append(el('div', 'r', 'R: ' + h.sold.moi + ' - ' + h.sold.luc));
    tr.insertCell().textContent = c.tao || '';
    const s = c.so_hd || {}, v = s.purchase || 0, r = s.sold || 0;
    [v, r, v + r].forEach(n => { const x = tr.insertCell(); x.className = 'n'; x.textContent = fmt(n); });
    const act = tr.insertCell(); act.style.whiteSpace = 'nowrap';
    const b1 = el('button', 'sm', 'Đồng bộ'); b1.onclick = () => openCompany(c.mst, 'dong-bo');
    const b2 = el('button', 'sm sec', 'Sửa'); b2.onclick = () => openEdit(c); b2.style.marginLeft = '4px';
    const b3 = el('button', 'sm sec', 'Xem HĐ'); b3.onclick = () => openCompany(c.mst); b3.style.marginLeft = '4px';
    act.append(b1, b2, b3);
  });
}
async function refresh() {
  const s = await api('/api/state'); companies = s.companies; me = s.me;
  document.body.classList.toggle('staff', me.role !== 'admin'); document.body.classList.toggle('local', !me.server);
  $('meName').textContent = me.ten ? (me.ten + ' (' + me.username + ')') : me.username;
  $('bellN').textContent = s.unread; $('bellN').classList.toggle('hide', !s.unread);
  if (s.auto) renderAuto(s.auto);
  render(); if (curTab === 'WS') { fillMst(); renderWsHead(); }
  $('capStat').textContent = s.captcha.count ? ('Captcha đã học ' + s.captcha.chars.length + ' ký tự') : '';
  if (s.job.running || s.job.log.length) renderJob(s.job);
  return s;
}
function openEdit(c) {
  editing = c || null; $('eErr').textContent = '';
  $('editTitle').textContent = c ? 'Cập nhật thông tin' : 'Thêm doanh nghiệp';
  $('eMst').value = c ? c.mst : ''; $('eMst').disabled = !!c; $('eTen').value = c ? c.ten : ''; $('ePw').value = '';
  $('eInfo').textContent = ''; $('eLookup').classList.toggle('hide', !!c);
  $('ePwHint').textContent = c && c.co_mk ? 'Đã lưu mật khẩu – để trống nếu không đổi. Mật khẩu này khác mật khẩu đăng nhập phần mềm.' : 'Mật khẩu được mã hoá khi lưu.';
  $('eVao').checked = c ? c.vao : true; $('eRa').checked = c ? c.ra : true; $('eAn').checked = c ? c.an : false;
  $('eDel').classList.toggle('hide', !c); show('mEdit'); (c ? $('ePw') : $('eMst')).focus();
}
async function lookupMst() {
  const mst = $('eMst').value.trim(); if (!mst) return;
  $('eInfo').textContent = 'Đang tra cứu…'; $('eErr').textContent = '';
  try { const d = await post('/api/mst-lookup', {mst});
        $('eTen').value = d.ten; $('eInfo').textContent = [d.dia_chi, d.cqt, d.tthai_text && ('Tình trạng: ' + d.tthai_text)].filter(Boolean).join(' · '); }
  catch (e) { $('eInfo').textContent = ''; $('eErr').textContent = e.message; }
}
async function saveEdit(test) {
  try {
    const mst = $('eMst').value.trim();
    await post('/api/company/save', {mst, ten: $('eTen').value, password: $('ePw').value,
      vao: $('eVao').checked, ra: $('eRa').checked, an: $('eAn').checked});
    hide('mEdit'); await refresh();
    if (test) await start({msts: [mst], test: true});
  } catch (e) { $('eErr').textContent = e.message; }
}
async function delEdit() {
  if (!editing || !confirm('Xoá doanh nghiệp ' + editing.mst + ' khỏi danh sách? (Hoá đơn đã tải vẫn giữ nguyên)')) return;
  await post('/api/company/delete', {mst: editing.mst}); hide('mEdit'); refresh();
}
async function doImport() {
  try { const d = await post('/api/company/import', {text: $('impText').value});
    $('impErr').textContent = d.errors.join('\n'); if (!d.errors.length) { hide('mImport'); $('impText').value = ''; } refresh();
  } catch (e) { $('impErr').textContent = e.message; }
}
function openBatch(msts) {
  batchMsts = msts || [...document.querySelectorAll('.pick:checked')].map(x => x.value);
  $('bErr').textContent = '';
  $('bWho').textContent = batchMsts.length ? ('Doanh nghiệp đã chọn: ' + batchMsts.length) : 'Chưa tích chọn doanh nghiệp nào → xử lý TẤT CẢ doanh nghiệp đang hiển thị.';
  show('mBatch');
}
async function runBatch() {
  const kinds = []; if ($('bVao').checked) kinds.push('purchase'); if ($('bRa').checked) kinds.push('sold');
  const msts = batchMsts.length ? batchMsts : companies.filter(c => !c.an).map(c => c.mst);
  try { await start({msts, kinds, from: $('bFrom').value, to: $('bTo').value, mtt: $('bMtt').checked, xml: $('bXml').checked});
        hide('mBatch'); } catch (e) { $('bErr').textContent = e.message; }
}
async function start(body) { await post('/api/run', body); show('jobCard'); poll(); }
async function sendCap(id) { id = id || 'capIn'; const v = $(id).value.trim(); if (!v) return; $(id).value = ''; hide('capBox'); hide('sCapBox');
  await post('/api/captcha', {answer: v}); poll(); }
let lastCap = null;
function renderJob(j) {
  if (curTab === 'WS' && wsSub === 'dong-bo') { hide('jobCard'); renderSync(j); return; }
  if (curTab === 'WS' && !j.running) { hide('jobCard'); return; }
  show('jobCard');
  $('jobTitle').textContent = j.running ? ('Đang xử lý' + (j.current ? ' ' + j.current : '') + '…') : 'Kết quả lần chạy gần nhất';
  $('btnStop').classList.toggle('hide', !j.running);
  $('progI').style.width = (j.total ? Math.round(100 * j.done / j.total) : (j.running ? 5 : 100)) + '%';
  if (j.need_captcha) {
    const key = j.need_captcha.mst + j.need_captcha.svg.length + j.need_captcha.svg.slice(-40);
    if (key !== lastCap) { lastCap = key; $('capFor').textContent = (j.need_captcha.ten || '') + ' (' + j.need_captcha.mst + ')';
      $('capImg').innerHTML = j.need_captcha.svg; $('capIn').value = ''; show('capBox'); $('capIn').focus(); }
  } else { hide('capBox'); lastCap = null; }
  const res = $('results'); res.innerHTML = '';
  j.results.forEach(r => res.append(el('div', 'ok', r.mst + ' ' + (r.ten || '') + ' – ' + r.loai + ': ' + r.so_hd + ' HĐ (' + r.moi + ' mới, ' + r.doi + ' đổi trạng thái) → ' + r.thu_muc)));
  $('log').textContent = j.log.join('\n'); $('log').scrollTop = 1e9;
}
let wasRunning = false;
async function poll() {
  clearTimeout(timer);
  let s; try { s = await refresh(); } catch (e) { timer = setTimeout(poll, 2000); return; }
  if (wasRunning && !s.job.running && curTab === 'WS') { loadInv(); }
  wasRunning = s.job.running;
  if (s.job.running || (s.auto && s.auto.running && curTab === 'ND')) timer = setTimeout(poll, 1500);
  else if (me.server) timer = setTimeout(poll, 60000);  // cập nhật chuông thông báo
}

// ---- Bản web: thông báo, tài khoản, người dùng, đồng bộ tự động ----
async function toggleNotes() {
  if (!$('notes').classList.contains('hide')) { hide('notes'); return; }
  const d = await post('/api/notices', {}); const box = $('notesList'); box.innerHTML = '';
  if (!d.items.length) box.append(el('div', 'mute', 'Chưa có thông báo.'));
  const unread = d.unread;
  d.items.forEach((n, i) => { const it = el('div', 'it' + (i < unread ? ' new' : ''));
    it.append(el('div', 'mute', n.t)); it.append(el('div', n.level === 'warn' ? 'warn' : n.level === 'err' ? 'err2' : '', n.text));
    if (n.mst && companies.some(c => c.mst === n.mst)) { const a = el('a', 'lk', 'Xem hoá đơn'); a.href = '#';
      a.onclick = ev => { ev.preventDefault(); hide('notes'); openCompany(n.mst); }; it.append(a); }
    box.append(it); });
  show('notes'); if (unread) { await post('/api/notices', {read: true}); $('bellN').classList.add('hide'); }
}
function openPw() { ['pOld', 'pNew', 'pNew2'].forEach(i => $(i).value = ''); $('pErr').textContent = ''; show('mPw'); $('pOld').focus(); }
async function savePw() {
  if ($('pNew').value !== $('pNew2').value) { $('pErr').textContent = 'Hai lần nhập mật khẩu mới không khớp'; return; }
  try { await post('/api/me/password', {old: $('pOld').value, new: $('pNew').value}); hide('mPw'); alert('Đã đổi mật khẩu.'); }
  catch (e) { $('pErr').textContent = e.message; }
}
async function logout() { try { await post('/api/logout', {}); } catch (e) {} location.reload(); }
let users = [], allCos = [], editingUser = null, uPicked = new Set();
async function loadUsers() {
  $('uErr').textContent = '';
  try { const d = await post('/api/users', {}); users = d.users; allCos = d.companies; } catch (e) { $('uErr').textContent = e.message; return; }
  const tb = $('uRows'); tb.innerHTML = '';
  users.forEach(u => { const tr = tb.insertRow();
    tr.insertCell().textContent = u.username; tr.insertCell().textContent = u.ten || '';
    tr.insertCell().textContent = u.role === 'admin' ? 'Quản trị' : 'Nhân viên';
    const n = tr.insertCell(); n.className = 'n'; n.textContent = u.role === 'admin' ? 'tất cả' : (u.msts || []).length;
    const st = tr.insertCell(); st.textContent = u.active === false ? 'Đã khoá' : 'Hoạt động'; st.className = u.active === false ? 'err' : 'ok';
    tr.insertCell().textContent = u.tao || '';
    const b = el('button', 'sm sec', 'Sửa'); b.onclick = () => openUser(u); tr.insertCell().append(b); });
}
function renderUList() {
  const q = $('uQ').value.trim().toLowerCase(), box = $('uList'); box.innerHTML = '';
  allCos.filter(c => !q || c.mst.includes(q) || (c.ten || '').toLowerCase().includes(q)).forEach(c => {
    const l = el('label'), cb = el('input'); cb.type = 'checkbox'; cb.checked = uPicked.has(c.mst);
    cb.onchange = () => { cb.checked ? uPicked.add(c.mst) : uPicked.delete(c.mst); $('uCnt').textContent = uPicked.size; };
    l.append(cb, document.createTextNode((c.ten || '(chưa đặt tên)') + ' (' + c.mst + ')')); box.append(l); });
  $('uCnt').textContent = uPicked.size;
}
function uPickAll(on) {
  const q = $('uQ').value.trim().toLowerCase();
  allCos.filter(c => !q || c.mst.includes(q) || (c.ten || '').toLowerCase().includes(q)).forEach(c => on ? uPicked.add(c.mst) : uPicked.delete(c.mst));
  renderUList();
}
function openUser(u) {
  editingUser = u || null; $('uEditErr').textContent = '';
  $('uTitle').textContent = u ? 'Sửa người dùng' : 'Thêm người dùng';
  $('uName').value = u ? u.username : ''; $('uName').disabled = !!u; $('uTen').value = u ? (u.ten || '') : '';
  $('uRole').value = u ? u.role : 'staff'; $('uActive').checked = !u || u.active !== false; $('uPw').value = '';
  $('uPwHint').textContent = u ? 'Để trống nếu không đổi. Đặt mật khẩu mới sẽ đăng xuất người này khỏi các máy khác.' : 'Tối thiểu 8 ký tự – gửi riêng cho nhân viên.';
  uPicked = new Set(u ? (u.msts || []) : []); $('uQ').value = ''; renderUList();
  $('uDel').classList.toggle('hide', !u || u.username === me.username); uRoleChanged(); show('mUser');
}
function uRoleChanged() { $('uMstBox').classList.toggle('hide', $('uRole').value === 'admin'); }
$('uRole').onchange = uRoleChanged; $('uQ').oninput = renderUList;
async function saveUser() {
  try { await post('/api/user/save', {username: $('uName').value.trim(), ten: $('uTen').value, role: $('uRole').value,
          password: $('uPw').value, active: $('uActive').checked, msts: [...uPicked], create: !editingUser});
        hide('mUser'); loadUsers(); }
  catch (e) { $('uEditErr').textContent = e.message; }
}
async function delUser() {
  if (!editingUser || !confirm('Xoá người dùng ' + editingUser.username + '?')) return;
  try { await post('/api/user/delete', {username: editingUser.username}); hide('mUser'); loadUsers(); }
  catch (e) { $('uEditErr').textContent = e.message; }
}
const AUTO_KY = [['auto', 'Tháng này (tới ngày 20 gồm cả tháng trước)'], ['7', '7 ngày gần nhất'], ['month', 'Tháng này'], ['2month', 'Tháng này và tháng trước']];
AUTO_KY.forEach(([v, t]) => $('aKy').append(new Option(t, v)));
let autoLoaded = false;
function renderAuto(a) {
  if (!autoLoaded || document.activeElement === document.body) { $('aOn').checked = a.on; $('aGio').value = a.gio; $('aKy').value = a.ky; autoLoaded = true; }
  $('aInfo').textContent = (a.running ? 'Đang chạy' + (a.dn ? ' – doanh nghiệp ' + a.dn[0] + '/' + a.dn[1] : '') + (a.current ? ' (' + a.current + ')' : '') + '. '
    : '') + (a.lan_cuoi ? 'Lần chạy gần nhất: ' + a.lan_cuoi : 'Chưa chạy lần nào');
  $('aLog').classList.toggle('hide', !(a.log && a.log.length)); $('aLog').textContent = (a.log || []).join('\n');
}
async function saveAuto() {
  try { renderAuto(await post('/api/auto/save', {on: $('aOn').checked, gio: $('aGio').value, ky: $('aKy').value})); $('aMsg').textContent = 'Đã lưu.'; setTimeout(() => $('aMsg').textContent = '', 3000); }
  catch (e) { $('uErr').textContent = e.message; }
}
async function runAuto() {
  if (!confirm('Chạy đồng bộ tự động ngay cho tất cả doanh nghiệp?')) return;
  try { await post('/api/auto/run', {}); refresh(); } catch (e) { $('uErr').textContent = e.message; }
}

// ---- Tiện ích ----
document.querySelectorAll('.tmenu a').forEach(a => a.onclick = ev => { ev.preventDefault();
  document.querySelectorAll('.tmenu a').forEach(x => x.classList.toggle('on', x === a));
  document.querySelectorAll('.tpane').forEach(p => p.classList.toggle('hide', p.id !== 't-' + a.dataset.t)); });
const readB64 = f => new Promise((res, rej) => { const r = new FileReader(); r.onload = () => res(String(r.result).split(',')[1]); r.onerror = rej; r.readAsDataURL(f); });
function download(name, b64, type) {
  const bin = atob(b64), u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i);
  const a = el('a'); a.href = URL.createObjectURL(new Blob([u], {type})); a.download = name; document.body.append(a); a.click(); a.remove();
}
$('xmlFile').onchange = async () => {
  const f = $('xmlFile').files[0]; if (!f) return; $('xmlErr').textContent = '';
  try { const d = await post('/api/tool/xml-view', {data: await readB64(f)});
        $('xmlFrame').srcdoc = d.html; show('xmlFrame'); show('xmlPrint'); }
  catch (e) { $('xmlErr').textContent = e.message; hide('xmlFrame'); hide('xmlPrint'); }
};
async function toolMst() {
  $('mstErr').textContent = 'Đang kiểm tra…'; const t = $('mstTbl'); t.innerHTML = '';
  try { const d = await post('/api/tool/mst', {text: $('mstText').value}); $('mstErr').textContent = '';
    const h = t.insertRow(); ['MST', 'Tên người nộp thuế', 'Tình trạng', 'Địa chỉ', 'Cơ quan thuế'].forEach(x => { const th = el('th', '', x); h.append(th); });
    d.rows.forEach(r => { const tr = t.insertRow(); tr.insertCell().textContent = r.mst; tr.insertCell().textContent = r.ten || r.loi;
      const s = tr.insertCell(); s.textContent = r.tthai_text; s.className = r.tthai_text === 'Đang hoạt động' ? 'ok' : 'err';
      tr.insertCell().textContent = r.dia_chi; tr.insertCell().textContent = r.cqt; });
  } catch (e) { $('mstErr').textContent = e.message; }
}
let mergeQueue = [];
function renderMerge() {
  const ol = $('mergeList'); ol.innerHTML = '';
  mergeQueue.forEach((f, i) => { const li = el('li', '', f.name + ' ');
    const up = el('button', 'sm sec', '↑'); up.onclick = () => { if (i) { [mergeQueue[i - 1], mergeQueue[i]] = [mergeQueue[i], mergeQueue[i - 1]]; renderMerge(); } };
    const dn = el('button', 'sm sec', '↓'); dn.onclick = () => { if (i < mergeQueue.length - 1) { [mergeQueue[i + 1], mergeQueue[i]] = [mergeQueue[i], mergeQueue[i + 1]]; renderMerge(); } };
    const rm = el('button', 'sm red', '×'); rm.onclick = () => { mergeQueue.splice(i, 1); renderMerge(); };
    li.append(up, dn, rm); ol.append(li); });
}
$('mergeFiles').onchange = () => { mergeQueue.push(...$('mergeFiles').files); $('mergeFiles').value = ''; renderMerge(); };
async function toolMerge() {
  $('mergeErr').textContent = 'Đang nối…';
  try { const files = []; for (const f of mergeQueue) files.push({name: f.name, data: await readB64(f)});
        const d = await post('/api/tool/pdf-merge', {files}); download(d.name, d.data, 'application/pdf'); $('mergeErr').textContent = ''; }
  catch (e) { $('mergeErr').textContent = e.message; }
}
async function toolSplit() {
  const f = $('splitFile').files[0]; if (!f) { $('splitErr').textContent = 'Chọn file PDF'; return; }
  $('splitErr').textContent = 'Đang tách…';
  try { const d = await post('/api/tool/pdf-split', {name: f.name, data: await readB64(f), ranges: $('splitRanges').value});
        download(d.name, d.data, 'application/zip'); $('splitErr').textContent = ''; }
  catch (e) { $('splitErr').textContent = e.message; }
}

// ---- Sao kê ngân hàng ----
let bankRows = [], bankMeta = {};
const BK_OPTS = ['bkTk', 'bkMa', 'bkThu', 'bkChi', 'bkNv', 'bkGc'];
BK_OPTS.forEach(id => { try { const v = localStorage.getItem('opt_' + id); if (v !== null) $(id).value = v; } catch (e) {}
  $(id).addEventListener('change', () => { try { localStorage.setItem('opt_' + id, $(id).value); } catch (e) {} renderBank(); }); });
$('bankFiles').onchange = async () => {
  const fs = [...$('bankFiles').files]; $('bankFiles').value = ''; if (!fs.length) return;
  $('bankErr').textContent = 'Đang đọc ' + fs.length + ' file…'; $('bankInfo').innerHTML = ''; bankRows = []; renderBank();
  try {
    const files = []; for (const f of fs) files.push({name: f.name, data: await readB64(f)});
    const d = await post('/api/tool/bank-parse', {files});
    $('bankErr').innerHTML = ''; d.errors.forEach(x => $('bankErr').append(el('div', 'err', x)));
    const info = $('bankInfo');
    d.statements.forEach(st => {
      const ok = !st.warnings.length, box = el('div', ok ? 'ok' : 'err');
      box.textContent = st.file + ': ' + [st.bank, st.account, st.company].filter(Boolean).join(' – ') + ' · ' + vnd(st.from) + ' → ' + vnd(st.to) +
        ' · thu ' + fmt(st.sum_in) + ' · chi ' + fmt(st.sum_out) + (ok ? (st.total_in != null ? ' · khớp tổng trên sao kê ✓' : '') : '');
      info.append(box); st.warnings.forEach(w => info.append(el('div', 'err', '⚠ ' + w)));
    });
    if (d.dups) info.append(el('div', 'mute', 'Đã bỏ ' + d.dups + ' giao dịch trùng giữa các file.'));
    const st0 = d.statements[0] || {};
    bankMeta = {bank: st0.bank || '', account: st0.account || ''};
    if (!$('bkMa').value && st0.bank) $('bkMa').placeholder = 'vd 112_' + st0.bank;
    bankRows = d.rows; renderBank();
  } catch (e) { $('bankErr').textContent = e.message; }
};
function renderBank() {
  const t = $('bankTbl'); t.innerHTML = '';
  $('bankExp').classList.toggle('hide', !bankRows.length);
  if (!bankRows.length) { $('bankSum').textContent = ''; return; }
  const h = t.insertRow(); ['Ngày', 'Số GD', 'Diễn giải', 'Đối tượng', 'Tiền vào', 'Tiền ra', 'Nợ', 'Có'].forEach((x, i) => {
    const th = el('th', i === 4 || i === 5 ? 'n' : '', x); h.append(th); });
  const tk = $('bkTk').value || '1121';
  bankRows.forEach(r => {
    const tr = t.insertRow(), inc = r.in > 0;
    [vnd(r.date), r.ref, r.desc, r.name].forEach(v => tr.insertCell().textContent = v);
    [r.in, r.out].forEach(v => { const c = tr.insertCell(); c.className = 'n'; c.textContent = v ? fmt(v) : ''; });
    const du = el('input', 'note'); du.style.width = '70px'; du.value = r.tkdu || ''; du.placeholder = inc ? ($('bkThu').value || '') : ($('bkChi').value || '');
    du.title = 'TK đối ứng riêng cho dòng này (để trống = theo mặc định)'; du.onchange = () => { r.tkdu = du.value.trim(); };
    const no = tr.insertCell(), co = tr.insertCell();
    if (inc) { no.textContent = tk; co.append(du); } else { no.append(du); co.textContent = tk; }
  });
  const si = bankRows.reduce((a, r) => a + r.in, 0), so = bankRows.reduce((a, r) => a + r.out, 0);
  $('bankSum').textContent = bankRows.length + ' giao dịch · tiền vào ' + fmt(si) + ' · tiền ra ' + fmt(so);
}
async function bankExport() {
  try {
    const opts = {tk: $('bkTk').value, ma_dt: $('bkMa').value, tk_thu: $('bkThu').value, tk_chi: $('bkChi').value,
                  nghiep_vu: $('bkNv').value, ghi_chu: $('bkGc').value};
    const d = await post('/api/tool/bank-export', {rows: bankRows, opts, bank: bankMeta.bank, account: bankMeta.account});
    download(d.name, d.data, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');
  } catch (e) { $('bankErr').textContent = e.message; }
}

// ---- Trang Đồng bộ ----
let syncMst = '', syncRange = null, syncRows = [];
const SPER = [['today', 'Hôm nay'], ['week', '1 tuần'], ['month', 'Tháng này']]
  .concat([...Array(12).keys()].map(i => ['m' + (i + 1), 'Tháng ' + (i + 1)])).concat([1, 2, 3, 4].map(i => ['q' + i, 'Quý ' + i]));
SPER.forEach(([v, t]) => $('sPer').append(new Option(t, v)));
function sPeriod() {
  const v = $('sPer').value, now = new Date(), y = +$('sYear').value || now.getFullYear(); let a, b;
  if (v === 'today') a = b = now;
  else if (v === 'week') { a = new Date(now); a.setDate(now.getDate() - 7); b = now; }
  else if (v === 'month') { a = new Date(now.getFullYear(), now.getMonth(), 1); b = now; }
  else if (v[0] === 'm') { const m = +v.slice(1) - 1; a = new Date(y, m, 1); b = new Date(y, m + 1, 0); }
  else { const q = +v.slice(1) - 1; a = new Date(y, q * 3, 1); b = new Date(y, q * 3 + 3, 0); }
  if (b > now) b = now;
  $('sFrom').value = isoD(a); $('sTo').value = isoD(b);
  $('sYear').classList.toggle('hide', ['today', 'week', 'month'].includes(v));
}
$('sPer').onchange = sPeriod; $('sYear').onchange = sPeriod;
function openSync(mst) { openCompany(mst, 'dong-bo'); }
async function runSync(kinds) {
  syncRange = [$('sFrom').value, $('sTo').value];
  try { await start({msts: [syncMst], kinds, from: syncRange[0], to: syncRange[1], mtt: $('sMtt').checked, xml: $('sXml').checked,
                     pdf: false}); }
  catch (e) { alert(e.message); }
}
function openInvoicesFromSync() {
  if (syncRange) { $('hPer').value = ''; $('hFrom').value = syncRange[0]; $('hTo').value = syncRange[1]; }
  hPage = 0; openSub('hoa-don');
}
const vnd = s => s ? s.split('-').reverse().join('/') : '';
function renderSync(j) {
  const b = $('sBadge');
  const phase = j.running ? (j.need_captcha ? 'chờ nhập captcha' : j.phase) : (j.log.length ? j.phase : 'chờ đồng bộ');
  b.textContent = phase; b.className = 'badge ' + (j.running ? 'run' : phase === 'đồng bộ xong' ? 'ok' : phase === 'lỗi' ? 'bad' : '');
  $('sStop').classList.toggle('hide', !j.running);
  if (j.need_captcha) {
    const key = j.need_captcha.svg.length + j.need_captcha.svg.slice(-40);
    if (key !== lastCap) { lastCap = key; $('sCapImg').innerHTML = j.need_captcha.svg; $('sCapIn').value = ''; show('sCapBox'); $('sCapIn').focus(); }
  } else { hide('sCapBox'); lastCap = null; }
  const st = $('sSteps'); st.innerHTML = '';
  const line = (html) => { const d = el('div'); d.innerHTML = html; st.append(d); };
  const esc = t => String(t).replace(/[&<>"]/g, ch => ({'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;'}[ch]));
  if (j.auth || j.running) {
    line('1. Chứng thực tài khoản: <b>' + esc(j.auth || 'đang chờ') + '</b>');
    const f = j.found || {};
    if ('purchase' in f || 'sold' in f) {
      line('2. Từ ngày ' + vnd((syncRange || [])[0]) + ' đến ngày ' + vnd((syncRange || [])[1]) + ' có:');
      if ('purchase' in f) line('&nbsp;&nbsp;- <b>' + f.purchase + '</b> hoá đơn mua vào');
      if ('sold' in f) line('&nbsp;&nbsp;- <b>' + f.sold + '</b> hoá đơn bán ra');
      line('3. Đồng bộ hoá đơn: <b>' + (j.running ? 'đang chạy' + (j.total ? ' – XML ' + j.done + '/' + j.total : '') : esc(j.phase)) + '</b>');
      line('&nbsp;&nbsp;- Tổng số HĐ: <b>' + ((f.purchase || 0) + (f.sold || 0)) + '</b> hoá đơn');
    }
    line('&nbsp;&nbsp;- Thời gian xử lý: ' + j.elapsed + ' giây');
    const errs = j.log.filter(x => x.includes('LỖI'));
    if (errs.length) line('<span class="err">' + esc(errs[errs.length - 1]) + '</span>');
  }
  $('sProg').style.width = (j.total ? Math.round(100 * j.done / j.total) : (j.running ? 5 : (j.log.length ? 100 : 0))) + '%';
  syncRows = j.rows || []; renderSyncRows();
}
function renderSyncRows() {
  const tb = $('sRows'); tb.innerHTML = '';
  syncRows.slice(0, 1500).forEach(r => {
    const tr = tb.insertRow();
    [r.loai, r.ngay, r.mau, r.kh].forEach(v => tr.insertCell().textContent = v);
    const so = tr.insertCell(); so.className = 'n'; so.textContent = r.so;
    tr.insertCell().textContent = r.dongbo;
    const k = tr.insertCell(); k.textContent = r.ketqua; k.className = r.ketqua.startsWith('OK') ? 'ok' : 'err';
  });
}
$('sCapIn').addEventListener('keydown', e => { if (e.key === 'Enter') sendCap('sCapIn'); });

// ---- Tab Hoá đơn ----
const TTHAI = ['', 'HĐ Mới', 'HĐ Thay thế', 'HĐ Điều chỉnh', 'HĐ Đã bị thay thế', 'HĐ Đã bị điều chỉnh', 'HĐ Đã bị hủy'];
const KQ = ['TCT đã nhận', 'Đang k.tra', 'CQT t.chối HĐ', 'HĐ đủ đ.kiện', 'HĐ k đủ đ.kiện', 'Đã cấp MST', 'TCT k nhận mã', 'Đã k.tra định kỳ', 'HĐ có mã từ máy tính tiền'];
TTHAI.forEach((t, i) => { if (t) $('hTthai').append(new Option(t, i)); });
KQ.forEach((t, i) => $('hKq').append(new Option(t, i)));
let curTab = 'DN', hRows = [], hPage = 0, hSize = 20, hSel = new Set();
const isoD = d => d.getFullYear() + '-' + String(d.getMonth() + 1).padStart(2, '0') + '-' + String(d.getDate()).padStart(2, '0');
const PERIODS = [['', '--Tuỳ chọn ngày--'], ['today', 'Hôm nay'], ['month', 'Tháng này']]
  .concat([...Array(12).keys()].map(i => ['m' + (i + 1), 'Tháng ' + (i + 1)]))
  .concat([1, 2, 3, 4].map(i => ['q' + i, 'Quý ' + i])).concat([['year', 'Năm nay'], ['lastyear', 'Năm trước']]);
PERIODS.forEach(([v, t]) => $('hPer').append(new Option(t, v)));
$('hPer').onchange = () => {
  const v = $('hPer').value, now = new Date(), y = now.getFullYear(); let a, b;
  if (v === 'today') a = b = now;
  else if (v === 'month') { a = new Date(y, now.getMonth(), 1); b = now; }
  else if (v[0] === 'm') { const m = +v.slice(1) - 1; a = new Date(y, m, 1); b = new Date(y, m + 1, 0); }
  else if (v[0] === 'q') { const q = +v.slice(1) - 1; a = new Date(y, q * 3, 1); b = new Date(y, q * 3 + 3, 0); }
  else if (v === 'year') { a = new Date(y, 0, 1); b = now; }
  else if (v === 'lastyear') { a = new Date(y - 1, 0, 1); b = new Date(y - 1, 11, 31); }
  if (a) { $('hFrom').value = isoD(a); $('hTo').value = isoD(b); hPage = 0; loadInv(); }
};
function tab(t, keepHash) {
  curTab = t; $('navDN').classList.toggle('on', t === 'DN'); $('navHD').classList.toggle('on', t === 'WS');
  $('tabDN').classList.toggle('hide', t !== 'DN'); $('tabWS').classList.toggle('hide', t !== 'WS');
  $('tabTI').classList.toggle('hide', t !== 'TI'); $('navTI').classList.toggle('on', t === 'TI');
  $('tabND').classList.toggle('hide', t !== 'ND'); $('navND').classList.toggle('on', t === 'ND');
  if (t !== 'WS') toolsHome();
  if (t === 'ND') loadUsers();
  if (!keepHash) setHash(t === 'WS' ? '#/dn/' + wsMst + '/' + wsSub : t === 'DN' ? '#/' : '#/' + t.toLowerCase());
}

// ---- Trang doanh nghiệp ----
let wsMst = '', wsSub = 'hoa-don';
const coLabel = (c, i) => String(i + 1).padStart(3, '0') + '. ' + (c.ten || '(chưa đặt tên)') + ' (' + c.mst + ')';
const visibleCos = () => companies.filter(c => !c.an || c.mst === wsMst);
function setHash(h) { if (location.hash !== h) history.pushState(null, '', h); }
function fillMst() {
  const dl = $('wsList'); dl.innerHTML = '';
  visibleCos().forEach((c, i) => dl.append(new Option(coLabel(c, i), coLabel(c, i))));
}
function renderWsHead() {
  const list = visibleCos(), i = list.findIndex(c => c.mst === wsMst), c = list[i] || {mst: wsMst};
  $('wsPick').value = i >= 0 ? coLabel(c, i) : wsMst; $('wsCrumb').textContent = c.ten || wsMst;
  $('wsPrev').disabled = i <= 0; $('wsNext').disabled = i < 0 || i >= list.length - 1;
  const h = c.lich_su || {}, so = c.so_hd || {}, m = $('wsMeta'); m.innerHTML = '';
  const item = (label, val, cls) => { const sp = el('span', cls || '', label + ': '); sp.append(el('b', '', val)); m.append(sp); };
  item('MST', c.mst || '');
  item('HĐ mua vào', fmt(so.purchase || 0)); item('HĐ bán ra', fmt(so.sold || 0));
  if (h.purchase) item('Đồng bộ vào', h.purchase.luc, 'v'); if (h.sold) item('Đồng bộ ra', h.sold.luc, 'r');
  if (c.loi) m.append(el('span', 'err', c.loi));
  document.querySelectorAll('#wsTabs a').forEach(a => a.classList.toggle('on', a.dataset.s === wsSub));
}
function openCompany(mst, sub) {
  if (!mst) { tab('DN'); return; }
  const changed = mst !== wsMst;
  wsMst = mst; syncMst = mst; $('hMst').value = mst;
  if (changed) { hSel.clear(); hColF = {}; hRows = []; }
  tab('WS', true); fillMst(); openSub(sub || (changed ? 'hoa-don' : wsSub), changed);
}
function openSub(sub, force) {
  wsSub = sub; renderWsHead();
  ['wsHD', 'wsSYNC', 'wsTOOL'].forEach(id => hide(id)); toolsHome();
  if (sub === 'dong-bo') {
    show('wsSYNC'); if (!$('sFrom').value) { $('sYear').value = new Date().getFullYear(); $('sPer').value = 'week'; sPeriod(); }
    poll();
  } else if (sub === 'sao-ke' || sub === 'tra-cuu-mst') {
    const pane = $(sub === 'sao-ke' ? 't-bank' : 't-mst'); $('wsTOOL').append(pane); pane.classList.remove('hide'); show('wsTOOL');
    if (sub === 'tra-cuu-mst' && !$('mstText').value.trim()) $('mstText').value = wsMst;
  } else { wsSub = 'hoa-don'; show('wsHD'); if (force || !hRows.length) loadInv(); else renderInv(); }
  setHash('#/dn/' + wsMst + '/' + wsSub);
}
const toolHome = $('t-bank').parentNode;
function toolsHome() {
  ['t-bank', 't-mst'].forEach(id => { const p = $(id); if (p.parentNode !== toolHome) { toolHome.append(p);
    p.classList.toggle('hide', !document.querySelector('.tmenu a.on[data-t="' + id.slice(2) + '"]')); } });
}
function stepCompany(d) {
  const list = visibleCos(), i = list.findIndex(c => c.mst === wsMst);
  if (list[i + d]) openCompany(list[i + d].mst, wsSub);
}
document.querySelectorAll('#wsTabs a').forEach(a => a.onclick = ev => { ev.preventDefault(); openSub(a.dataset.s); });
$('wsPick').addEventListener('change', () => {
  const v = $('wsPick').value.trim(), m = v.match(/\((\d{10}(?:-\d{3})?)\)\s*$/) || v.match(/^(\d{10}(?:-\d{3})?)$/);
  let c = m && companies.find(x => x.mst === m[1]);
  if (!c && v) { const q = v.toLowerCase(), hits = companies.filter(x => (x.ten || '').toLowerCase().includes(q) || x.mst.includes(q)); if (hits.length === 1) c = hits[0]; }
  if (c) openCompany(c.mst, wsSub); else if (v) $('wsPick').select();
});
$('wsPick').addEventListener('focus', () => $('wsPick').select());
document.addEventListener('keydown', e => { if (curTab === 'WS' && e.altKey && (e.key === 'ArrowLeft' || e.key === 'ArrowRight')) {
  e.preventDefault(); stepCompany(e.key === 'ArrowLeft' ? -1 : 1); } });
function route() {
  const h = location.hash, m = h.match(/^#\/dn\/([\d-]+)(?:\/([\w-]+))?/);
  if (m) { if (m[1] !== wsMst || curTab !== 'WS') openCompany(m[1], m[2]); else if (m[2] && m[2] !== wsSub) openSub(m[2]); return; }
  const t = (h.match(/^#\/(\w+)/) || [])[1];
  tab(t === 'ti' ? 'TI' : t === 'nd' ? 'ND' : 'DN', true);
}
window.addEventListener('popstate', route);
document.querySelectorAll('#hKindSeg button').forEach(b => b.onclick = () => {
  document.querySelectorAll('#hKindSeg button').forEach(x => x.classList.toggle('on', x === b));
  $('hKind').value = b.dataset.k; hSel.clear(); hColF = {}; hPage = 0; loadInv();
});
function hFilters(useSel) {
  const keys = useSel ? hRows.filter(r => hSel.has(r.key)).map(r => r.key) : [];
  return {keys, kind: $('hKind').value, file: $('hFile').value, duyet: $('hDuyet').value, tthai: $('hTthai').value, kq: $('hKq').value,
          khhdon: $('hKh').value.trim(), shdon: $('hSo').value.trim(), q: $('hQ').value.trim(), from: $('hFrom').value, to: $('hTo').value};
}
async function loadInv() {
  $('hErr').textContent = '';
  if (!$('hMst').value) { hRows = []; buildHead(); renderInv(); return; }
  try { hRows = (await post('/api/invoices', {mst: $('hMst').value, filters: hFilters()})).rows; }
  catch (e) { $('hErr').textContent = e.message; hRows = []; }
  const live = new Set(hRows.map(r => r.key)); hSel.forEach(k => { if (!live.has(k)) hSel.delete(k); });
  buildHead(); renderInv();
}
function fileLink(rel, text) {
  const a = el('a', 'lk', text); a.href = '/file?k=' + encodeURIComponent(KEY) + '&mst=' + encodeURIComponent($('hMst').value) + '&p=' + encodeURIComponent(rel);
  a.target = '_blank'; return a;
}

// ---- Bảng hoá đơn: sắp xếp, lọc theo cột, chọn dòng ----
const HCOLS = [
  {k: 'mst', t: 'MST', f: 'text'}, {k: 'ten', t: 'Người bán', f: 'text', cls: 'ten'}, {k: 'ngay', t: 'Ngày', f: 'text', sort: 'date', cls: 'nw'},
  {k: 'khhdon', t: 'Ký hiệu', f: 'text', cls: 'nw'}, {k: 'shdon', t: 'Số HĐ', f: 'text', n: 1, sort: 'num'},
  {k: 'cthue', t: 'Tiền C.Thuế', f: 'num', n: 1, sum: 1}, {k: 'thue', t: 'Tiền thuế', f: 'num', n: 1, sum: 1},
  {k: 'ck', t: 'CK.TM', f: 'num', n: 1, sum: 1}, {k: 'phi', t: 'Phí', f: 'num', n: 1, sum: 1},
  {k: 'tt', t: 'Tổng T.Toán', f: 'num', n: 1, sum: 1}, {k: 'tthai', t: 'T.thái HĐ', f: 'pick', cls: 'nw'},
  {k: 'kq', t: 'Kết quả k.tra', f: 'pick'}, {k: 'duyet', t: 'Duyệt nội bộ', f: 'pick'}, {k: 'dv', t: 'HĐ DV', f: 'bool'},
  {k: 'mat_hang', t: 'Mặt hàng', f: 'text', cls: 'mh'}, {k: 'note', t: 'Ghi chú', f: 'text'}, {k: '_ct', t: 'Chi tiết'}];
let hSort = {k: 'ngay', dir: 1}, hColF = {}, hLastClick = null;
try { const g = JSON.parse(localStorage.getItem('grid') || '{}'); if (g.sort) hSort = g.sort; if (g.size) hSize = g.size; } catch (e) {}
const saveGrid = () => { try { localStorage.setItem('grid', JSON.stringify({sort: hSort, size: hSize})); } catch (e) {} };
const dkey = s => s ? s.split('/').reverse().join('') : '';
function resetGrid() { hColF = {}; hSort = {k: 'ngay', dir: 1}; saveGrid(); buildHead(); renderInv(); }
function viewRows() {
  let rows = hRows.filter(r => HCOLS.every(c => {
    const f = hColF[c.k]; if (f == null || f === '') return true;
    if (c.f === 'pick') return r[c.k] === f;
    if (c.f === 'bool') return String(!!r[c.k]) === f;
    if (c.f === 'num') { const m = f.replace(/\s/g, '').match(/^(>=|<=|>|<|=)?([\d.,]+)$/);
      if (m) { const v = +m[2].replace(/[.,]/g, ''), x = r[c.k] || 0;
        return {'>': x > v, '<': x < v, '>=': x >= v, '<=': x <= v}[m[1]] ?? x === v; }
      return fmt(r[c.k]).includes(f); }
    return String(r[c.k] ?? '').toLowerCase().includes(f.toLowerCase());
  }));
  const c = HCOLS.find(x => x.k === hSort.k);
  if (c) {
    const key = c.sort === 'date' ? (r => dkey(r.ngay) + String(r.shdon).padStart(12, '0'))
      : c.sort === 'num' ? (r => +r[c.k] || 0) : c.n ? (r => r[c.k] || 0) : (r => String(r[c.k] ?? '').toLowerCase());
    rows = rows.slice().sort((a, b) => { const x = key(a), y = key(b); return (x > y ? 1 : x < y ? -1 : 0) * hSort.dir; });
  }
  return rows;
}
function buildHead() {
  const sold = $('hKind').value.startsWith('sold');
  HCOLS[1].t = sold ? 'Người mua' : 'Người bán';
  const th = $('hHead'); th.innerHTML = '';
  const r1 = th.insertRow(), r2 = th.insertRow(); r2.className = 'f';
  const all = el('input'); all.type = 'checkbox'; all.id = 'hAll'; all.title = 'Chọn tất cả các dòng đang lọc';
  all.onchange = () => { const v = viewRows(); if (all.checked) v.forEach(r => hSel.add(r.key)); else v.forEach(r => hSel.delete(r.key)); renderInv(); };
  const c0 = el('th'); c0.append(all); r1.append(c0); r2.append(el('th'));
  HCOLS.forEach(c => {
    const h = el('th', (c.n ? 'n ' : '') + (c.k !== '_ct' ? 's' : ''), c.t);
    if (hSort.k === c.k) h.textContent += hSort.dir > 0 ? ' ▲' : ' ▼';
    if (c.k !== '_ct') { h.title = 'Bấm để sắp xếp'; h.onclick = () => { hSort = {k: c.k, dir: hSort.k === c.k ? -hSort.dir : 1}; saveGrid(); buildHead(); renderInv(); }; }
    r1.append(h);
    const f = el('th'); r2.append(f);
    if (c.f === 'pick' || c.f === 'bool') {
      const sel = el('select'); sel.append(new Option('(Tất cả)', ''));
      if (c.f === 'bool') { sel.append(new Option('Có', 'true')); sel.append(new Option('Không', 'false')); }
      else [...new Set(hRows.map(r => r[c.k]).filter(Boolean))].sort().forEach(v => sel.append(new Option(v, v)));
      sel.value = hColF[c.k] || ''; sel.onchange = () => { hColF[c.k] = sel.value; hPage = 0; renderInv(); }; f.append(sel);
    } else if (c.f) {
      const inp = el('input'); inp.type = 'search'; inp.value = hColF[c.k] || ''; inp.placeholder = '🔍';
      if (c.f === 'num') inp.title = 'Lọc số tiền: gõ 1000000, >1000000, <500000, >=…';
      inp.oninput = () => { hColF[c.k] = inp.value; hPage = 0; renderInv(); }; f.append(inp);
    }
  });
  requestAnimationFrame(() => $('hTbl').style.setProperty('--hh', r1.getBoundingClientRect().height + 'px'));
}
function selInfo(view) {
  const sel = hRows.filter(r => hSel.has(r.key));
  $('hBulk').classList.toggle('hide', !sel.length);
  $('hSelInfo').textContent = sel.length ? ('Đã chọn ' + sel.length + ' HĐ · tổng thanh toán ' + fmt(sel.reduce((a, r) => a + r.tt, 0))) : '';
  const all = $('hAll'); if (all) { const v = view || viewRows(); all.checked = v.length > 0 && v.every(r => hSel.has(r.key));
    all.indeterminate = !all.checked && v.some(r => hSel.has(r.key)); }
}
const BAD = /bị hủy|bị huỷ|bị thay thế|bị điều chỉnh|từ chối|k đủ/i;
function renderInv() {
  if (!$('hHead').rows.length) buildHead();
  const view = viewRows(), tb = $('hRows'); tb.innerHTML = '';
  const size = hSize || view.length || 1, pages = Math.max(1, Math.ceil(view.length / size)); hPage = Math.min(hPage, pages - 1);
  const kind = $('hKind').value.replace('_dv', '');
  view.slice(hPage * size, hPage * size + size).forEach(r => {
    const tr = tb.insertRow(); tr.dataset.key = r.key;
    if (hSel.has(r.key)) tr.classList.add('sel');
    if (BAD.test(r.tthai + ' ' + r.kq)) tr.classList.add('bad');
    const pick = el('input'); pick.type = 'checkbox'; pick.checked = hSel.has(r.key); tr.insertCell().append(pick);
    tr.onclick = ev => {
      if (ev.target.closest('a,select,button,input:not([type=checkbox]),label')) return;
      const keys = view.map(x => x.key), on = !hSel.has(r.key);
      if (ev.shiftKey && hLastClick && keys.includes(hLastClick)) {
        const [a, b] = [keys.indexOf(hLastClick), keys.indexOf(r.key)].sort((x, y) => x - y);
        keys.slice(a, b + 1).forEach(k => hSel.add(k));
      } else on ? hSel.add(r.key) : hSel.delete(r.key);
      hLastClick = r.key; renderInv();
    };
    HCOLS.forEach(c => {
      if (c.k === 'duyet') { const sel = el('select'); ['Chờ duyệt', 'Đã duyệt', 'Không duyệt'].forEach(o => sel.append(new Option(o, o)));
        sel.value = r.duyet; sel.onchange = () => updInv(r, {duyet: sel.value}); tr.insertCell().append(sel); return; }
      if (c.k === 'dv') { const cb = el('input'); cb.type = 'checkbox'; cb.checked = r.dv; cb.title = 'Hoá đơn dịch vụ';
        cb.onclick = ev => ev.stopPropagation(); cb.onchange = () => updInv(r, {dv: cb.checked}); tr.insertCell().append(cb); return; }
      if (c.k === 'note') { const note = el('input', 'note'); note.value = r.note; note.placeholder = 'ghi chú…';
        note.onchange = () => updInv(r, {note: note.value}); tr.insertCell().append(note); return; }
      if (c.k === '_ct') { renderDetails(tr, r, kind); return; }
      const td = tr.insertCell(); td.className = (c.n ? 'n ' : '') + (c.cls || '');
      td.textContent = c.sum ? fmt(r[c.k]) : (r[c.k] ?? '');
    });
  });
  const f = $('hFoot'); f.innerHTML = '';
  const c0 = f.insertCell(); c0.colSpan = 6; c0.textContent = view.length + ' HĐ' + (view.length !== hRows.length ? ' (lọc từ ' + hRows.length + ')' : '');
  HCOLS.filter(c => c.sum).forEach(c => { const x = f.insertCell(); x.className = 'n'; x.textContent = fmt(view.reduce((a, r) => a + (r[c.k] || 0), 0)); });
  f.insertCell().colSpan = HCOLS.length - 9;
  const pg = $('hPager'); pg.innerHTML = '';
  [20, 50, 100, 0].forEach(n => { const b = el('button', 'sm sec' + (n === hSize ? ' cur' : ''), n ? String(n) : 'Tất cả'); if (n === hSize) b.style.color = '#fff';
    b.onclick = () => { hSize = n; hPage = 0; saveGrid(); renderInv(); }; pg.append(b); });
  pg.append(el('span', 'mute', ' Trang ' + (hPage + 1) + '/' + pages + ' '));
  const prev = el('button', 'sm sec', '‹'); prev.disabled = hPage === 0; prev.onclick = () => { hPage--; renderInv(); };
  const next = el('button', 'sm sec', '›'); next.disabled = hPage >= pages - 1; next.onclick = () => { hPage++; renderInv(); };
  pg.append(prev, next, el('span', 'mute', '  Mẹo: bấm vào dòng để chọn, Shift+bấm để chọn liên tiếp, bấm tiêu đề cột để sắp xếp.'));
  selInfo(view);
}
function renderDetails(tr, r, kind) {
    const ct = tr.insertCell(); ct.style.whiteSpace = 'nowrap';
    if (r.xml) ct.append(fileLink(r.xml, 'XML')); if (r.html) ct.append(fileLink(r.html, 'HTML')); if (r.pdf) ct.append(fileLink(r.pdf, 'PDF gốc'));
    const cq = el('a', 'lk', 'PDF thuế'); cq.target = '_blank'; cq.title = 'Bản thể hiện hoá đơn theo dữ liệu cổng thuế (hoadondientu.gdt.gov.vn)';
    cq.href = '/cqt?' + new URLSearchParams({k: KEY, mst: $('hMst').value, kind, key: r.key}); ct.append(cq);
    const t = r.tra_cuu || {};
    if (t.url || t.code) {
      const a = el('a', 'lk', 'Tra cứu'); a.href = '#';
      a.title = (t.ncc || '') + (t.code ? '\n' + t.field + ': ' + t.code + ' (đã chép, dán vào trang tra cứu)' : '') +
                (t.ncc && t.ncc.startsWith('Viettel') ? '\nMST bên bán: ' + r.mst : '');
      a.onclick = ev => { ev.preventDefault(); if (t.code && navigator.clipboard) navigator.clipboard.writeText(t.code).catch(() => {});
        if (t.url) window.open(t.url, '_blank'); else prompt((t.ncc || '') + ' – mã tra cứu', t.code); };
      ct.append(a);
    }
    if (!r.pdf && (r.pdf_url || r.can_fetch)) {
      const g = el('a', 'lk', 'Tải PDF gốc'); g.href = '#'; g.title = 'Tải PDF gốc từ ' + t.ncc + ' rồi tự gắn vào hoá đơn';
      g.onclick = async ev => { ev.preventDefault();
        if (me.server || !r.pdf_url) {  // máy chủ / phần mềm tự tải (EasyInvoice tự giải captcha)
          g.textContent = 'Đang tải…'; $('hErr').textContent = '';
          try { await post('/api/invoice/fetch-pdf', {mst: $('hMst').value, kind, key: r.key}); loadInv(); return; }
          catch (e) { g.textContent = 'Tải PDF gốc'; if (!r.pdf_url) { $('hErr').textContent = r.shdon + ': ' + e.message; return; } } }
        browserDownload([r.pdf_url]); };
      ct.append(g);
    }
    const up = el('a', 'lk', r.pdf ? '↻PDF' : '+PDF'); up.href = '#'; up.title = 'Gắn file PDF gốc đã tải từ trang tra cứu';
    up.onclick = ev => { ev.preventDefault(); const fi = el('input'); fi.type = 'file'; fi.accept = 'application/pdf';
      fi.onchange = async () => { const f = fi.files[0]; if (!f) return;
        const b64 = await new Promise(res => { const rd = new FileReader(); rd.onload = () => res(String(rd.result).split(',')[1]); rd.readAsDataURL(f); });
        try { await post('/api/invoice/pdf', {mst: $('hMst').value, kind, key: r.key, data: b64}); loadInv(); }
        catch (e) { $('hErr').textContent = e.message; } };
      fi.click(); };
    ct.append(up);
}
async function bulkUpdate(fields) {
  const keys = hRows.filter(r => hSel.has(r.key)).map(r => r.key);
  if (!keys.length || Object.values(fields).every(v => v === '')) return;
  try { await post('/api/invoice/update-bulk', Object.assign({mst: wsMst, kind: $('hKind').value.replace('_dv', ''), keys}, fields));
        hRows.forEach(r => { if (hSel.has(r.key)) Object.assign(r, fields); }); buildHead(); renderInv(); }
  catch (e) { $('hErr').textContent = e.message; }
}
$('hImpFiles').onchange = async () => {
  const fs = [...$('hImpFiles').files]; $('hImpFiles').value = ''; if (!fs.length) return;
  const b = $('hImpPdf'); b.disabled = true; b.textContent = 'Đang gắn ' + fs.length + ' file…'; $('hErr').textContent = '';
  try { const files = []; for (const f of fs) files.push({name: f.name, data: await readB64(f)});
    const d = await post('/api/invoice/pdf-import', {mst: $('hMst').value, files});
    const ok = d.results.filter(r => r.ok), bad = d.results.filter(r => !r.ok);
    $('hErr').innerHTML = ''; $('hErr').append(el('div', 'ok', 'Đã gắn ' + ok.length + '/' + d.results.length + ' file PDF.'));
    bad.forEach(r => $('hErr').append(el('div', 'err', r.file + ': ' + r.loi)));
    loadInv(); }
  catch (e) { $('hErr').textContent = e.message; }
  b.disabled = false; b.textContent = 'Gắn PDF gốc có sẵn';
};
async function bulkPdf() {
  const sel = hRows.filter(r => hSel.has(r.key)), src = sel.length ? sel : hRows;
  let need = src.filter(r => !r.pdf && (r.pdf_url || r.can_fetch));
  if (!need.length) { $('hErr').textContent = 'Không có hoá đơn nào cần tải PDF gốc (hiện tự tải được: MISA, EasyInvoice). Nhà cung cấp khác: bấm "Tra cứu" rồi "Gắn PDF gốc có sẵn".'; return; }
  // Phần mềm tự tải trước (EasyInvoice; bản web: cả MISA), từng nhóm 10 HĐ; MISA bản máy: trình duyệt tải.
  const auto = need.filter(r => r.can_fetch && (me.server || !r.pdf_url)), errs = [];
  if (auto.length) {
    const b = $('hBulkPdf'); b.disabled = true; let done = 0;
    for (let i = 0; i < auto.length; i += 10) {
      b.textContent = 'Đang tải PDF gốc ' + Math.min(i + 10, auto.length) + '/' + auto.length + '…';
      try { const f = hFilters(false); f.keys = auto.slice(i, i + 10).map(r => r.key);
        const d = await post('/api/invoice/fetch-pdf-bulk', {mst: $('hMst').value, filters: f}); done += d.ok; errs.push(...d.errors); }
      catch (e) { errs.push(e.message); }
    }
    b.disabled = false; b.textContent = 'Tải HĐ gốc hàng loạt';
    await loadInv();
    const keys = new Set(need.map(r => r.key)); need = hRows.filter(r => keys.has(r.key) && !r.pdf && r.pdf_url);
    $('hErr').innerHTML = ''; $('hErr').append(el('div', 'ok', 'Đã tải và gắn ' + done + '/' + auto.length + ' PDF gốc.'));
    errs.slice(0, 10).forEach(x => $('hErr').append(el('div', 'err', x)));
    if (!need.length) return;
  }
  browserDownload(need.map(r => r.pdf_url));
}
// Trình duyệt tải PDF gốc về thư mục Downloads (MISA chặn chương trình tự động nhưng không chặn trình duyệt);
// phần mềm theo dõi Downloads và tự gắn file mới vào đúng hoá đơn.
let scanTimer = null, promptOk = false;
async function browserDownload(urls) {
  if (me.server) {  // máy chủ không thấy thư mục Downloads trên máy bạn → tải về rồi gắn bằng "Gắn PDF gốc có sẵn"
    urls.forEach((u, i) => setTimeout(() => {
      const f = el('iframe'); f.style.display = 'none'; f.src = u; document.body.append(f); setTimeout(() => f.remove(), 120000);
    }, i * 1200));
    $('hErr').innerHTML = ''; $('hErr').append(el('div', 'ok', 'Trình duyệt đang tải ' + urls.length + ' PDF gốc về máy bạn (thư mục Downloads; nếu Chrome hỏi "tải nhiều tệp", chọn Cho phép). ' +
      'Tải xong bấm "Gắn PDF gốc có sẵn" và chọn các file vừa tải – phần mềm tự gắn vào đúng hoá đơn.'));
    const box = el('div', 'mute'); box.append('Nếu không thấy tải, bấm từng link: '); $('hErr').append(box);
    urls.slice(0, 50).forEach(u => { const a = el('a', 'lk', 'link'); a.href = u; a.target = '_blank'; box.append(a); });
    return;
  }
  if (urls.length > 1 && !promptOk) {
    try { const d = await post('/api/invoice/scan-downloads', {mst: $('hMst').value, info_only: true});
      if (d.prompt && !confirm('Chrome/Edge đang bật "Hỏi vị trí lưu từng tệp trước khi tải xuống" nên sẽ hiện hộp thoại Lưu cho TỪNG file.\n\n' +
          'Tắt một lần: mở tab mới, gõ  chrome://settings/downloads  (Edge: edge://settings/downloads)  rồi TẮT mục ' +
          '"Hỏi vị trí lưu từng tệp trước khi tải xuống".\n\nBấm OK để vẫn tải (phải bấm Lưu từng file), Cancel để đi tắt trước.')) return;
      promptOk = true;
    } catch (e) {}
  }
  const since = Date.now() / 1000;
  urls.forEach((u, i) => setTimeout(() => {
    const f = el('iframe'); f.style.display = 'none'; f.src = u; document.body.append(f); setTimeout(() => f.remove(), 120000);
  }, i * 1200));
  $('hErr').innerHTML = ''; $('hErr').append(el('div', 'ok', 'Trình duyệt đang tải ' + urls.length + ' PDF gốc về thư mục Downloads… ' +
    '(nếu Chrome hỏi "tải nhiều tệp", chọn Cho phép). Phần mềm sẽ tự gắn khi file về.'));
  const box = el('div', 'mute'); $('hErr').append(box);
  urls.slice(0, 50).forEach(u => { const a = el('a', 'lk', 'link'); a.href = u; a.target = '_blank'; box.append(a); });
  if (urls.length) box.prepend('Nếu không thấy tải, bấm từng link: ');
  let tries = 0, attached = 0; clearInterval(scanTimer); $('hBulkPdf').disabled = true;
  scanTimer = setInterval(async () => {
    tries++;
    try { const d = await post('/api/invoice/scan-downloads', {mst: $('hMst').value, since});
      const ok = d.results.filter(r => r.ok); attached += ok.length;
      if (ok.length) { loadInv(); }
      $('hBulkPdf').textContent = attached ? ('Đã gắn ' + attached + '/' + urls.length + ' PDF') : 'Đang chờ file tải về…';
      if (attached >= urls.length || tries > 20 + urls.length * 3) { clearInterval(scanTimer); $('hBulkPdf').textContent = 'Tải HĐ gốc hàng loạt'; $('hBulkPdf').disabled = false;
        if (attached < urls.length) $('hErr').append(el('div', 'err', 'Mới gắn được ' + attached + '/' + urls.length +
          ' file. Kiểm tra thư mục tải về (' + d.folder + ') hoặc dùng "Gắn PDF gốc có sẵn".')); }
    } catch (e) { clearInterval(scanTimer); $('hBulkPdf').disabled = false; $('hErr').append(el('div', 'err', e.message)); }
  }, 3000);
}
async function updInv(r, fields) {
  try { await post('/api/invoice/update', Object.assign({mst: $('hMst').value, kind: $('hKind').value.replace('_dv', ''), key: r.key}, fields));
        Object.assign(r, fields); } catch (e) { $('hErr').textContent = e.message; }
}
async function exportInv(fmtx) {
  if (!fmtx) return; $('hErr').textContent = '';
  try { const d = await post('/api/export', {mst: $('hMst').value, filters: hFilters(true), fmt: fmtx}); window.open(d.url, '_blank'); }
  catch (e) { $('hErr').textContent = e.message; }
}
['hFile', 'hDuyet', 'hTthai', 'hKq'].forEach(id => $(id).onchange = () => { hPage = 0; loadInv(); });
['hKh', 'hSo', 'hQ'].forEach(id => $(id).addEventListener('keydown', e => { if (e.key === 'Enter') { hPage = 0; loadInv(); } }));
$('q').oninput = render;
$('toggleAn').onclick = e => { e.preventDefault(); showHidden = !showHidden; render(); };
$('all').onchange = e => document.querySelectorAll('.pick').forEach(x => x.checked = e.target.checked);
$('capIn').addEventListener('keydown', e => { if (e.key === 'Enter') sendCap(); });
$('eMst').addEventListener('change', () => { if (!editing && !$('eTen').value.trim()) lookupMst(); });
(() => {
  const now = new Date(), first = new Date(now.getFullYear(), now.getMonth() - 1, 1), last = new Date(now.getFullYear(), now.getMonth(), 0);
  const iso = d => d.getFullYear() + '-' + String(d.getMonth() + 1).padStart(2, '0') + '-' + String(d.getDate()).padStart(2, '0');
  $('bFrom').value = iso(first); $('bTo').value = iso(last);
  $('hPer').value = 'month'; $('hFrom').value = iso(new Date(now.getFullYear(), now.getMonth(), 1)); $('hTo').value = iso(now);
  poll().then(() => { if (location.hash.length > 2) route(); });
})();
</script></body></html>
"""


AUTO_USER = {"username": "__tu-dong__", "ten": "Đồng bộ tự động", "role": "admin", "msts": [], "active": True}
AUTO_KY = {"auto": "Tháng này (tới ngày 20 gồm cả tháng trước)", "7": "7 ngày gần nhất", "month": "Tháng này",
           "2month": "Tháng này và tháng trước"}


def auto_range(ky, today=None):
    today = today or date.today()
    prev_first = (today.replace(day=1) - timedelta(days=1)).replace(day=1)
    if ky == "7":
        return today - timedelta(days=7), today
    if ky == "month":
        return today.replace(day=1), today
    if ky == "2month" or today.day <= 20:
        return prev_first, today
    return today.replace(day=1), today


def auto_sync_loop(app, interval=30):
    """Mỗi ngày đến giờ đã đặt thì đồng bộ tự động (bản web)."""
    while True:
        time.sleep(interval)
        try:
            a = app.settings()["auto"]
            now = datetime.now()
            if not a.get("on") or now.strftime("%H:%M") < a.get("gio", "06:00"):
                continue
            if (a.get("lan_cuoi") or "")[:10] == now.strftime("%Y-%m-%d"):
                continue
            app.run_auto(wait=True)
        except Exception as e:  # không để vòng lặp chết
            print("Đồng bộ tự động lỗi: %s" % e)


LOCAL_USER = {"username": "admin", "ten": "Máy này", "role": "admin", "msts": [], "active": True}
SESSION_TTL = 12 * 3600      # phiên đăng nhập bản web hết hạn sau 12 giờ không dùng
LOGIN_WINDOW = 900           # đăng nhập sai quá LOGIN_MAX lần trong 15 phút → tạm khoá
LOGIN_MAX = 8


class App:
    def __init__(self, out_root, client_factory=None, server=False):
        self.out_root = out_root
        self.server = server  # True: bản web nhiều người dùng (máy chủ); False: chạy trên máy, không cần đăng nhập
        cfg = os.path.join(out_root, "_cau-hinh")
        self.cfg = cfg
        self.store = Store(os.path.join(cfg, "doanh-nghiep.json"))
        self.users = UserStore(os.path.join(cfg, "nguoi-dung.json"))
        self.solver = CaptchaSolver(os.path.join(cfg, "captcha-mau.json"))
        self.client_factory = client_factory or HoaDonClient
        self.clients = {}  # MST → client đã đăng nhập (giữ phiên)
        self.downloads = default_downloads_dir()  # nơi trình duyệt lưu PDF gốc tải về
        self.seen_downloads = set()
        self.key = secrets.token_urlsafe(24)
        self.lock = threading.Lock()
        self.jobs = {}          # tên đăng nhập → Job (mỗi người một lượt đồng bộ riêng)
        self.sessions = {}      # mã phiên (cookie) → {user, csrf, exp}
        self.fails = {}         # IP / tên đăng nhập → các lần đăng nhập sai gần đây
        self.allowed_hosts = set()  # bản web: tên miền được phép (rỗng = không kiểm tra)
        self.trust_proxy = False    # chạy sau Caddy/Nginx: tin X-Forwarded-For / X-Forwarded-Proto
        self.max_jobs = 4           # số lượt đồng bộ chạy cùng lúc trên máy chủ

    def job_for(self, user):
        with self.lock:
            return self.jobs.setdefault(user["username"], Job())

    @property
    def job(self):  # bản chạy trên máy: một người dùng
        return self.job_for(LOCAL_USER)

    def running_jobs(self):
        with self.lock:
            return sum(1 for j in self.jobs.values() if j.running)

    def audit(self, who, action):
        try:
            os.makedirs(self.cfg, exist_ok=True)
            with open(os.path.join(self.cfg, "nhat-ky.log"), "a", encoding="utf-8") as f:
                f.write("%s\t%s\t%s\n" % (datetime.now().strftime("%Y-%m-%d %H:%M:%S"), who, action))
        except OSError:
            pass

    # ---- đăng nhập bản web ----
    def too_many_fails(self, *keys):
        now = time.time()
        with self.lock:
            for k in keys:
                self.fails[k] = [t for t in self.fails.get(k, []) if now - t < LOGIN_WINDOW]
            return any(len(self.fails[k]) >= LOGIN_MAX for k in keys)

    def add_fail(self, *keys):
        with self.lock:
            for k in keys:
                self.fails.setdefault(k, []).append(time.time())

    def new_session(self, username):
        sid, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(24)
        with self.lock:
            now = time.time()
            for k in [k for k, v in self.sessions.items() if v["exp"] < now]:
                del self.sessions[k]
            self.sessions[sid] = {"user": username, "csrf": csrf, "exp": now + SESSION_TTL}
        return sid

    def session_user(self, sid):
        if not sid:
            return None, None
        with self.lock:
            s = self.sessions.get(sid)
            if not s or s["exp"] < time.time():
                self.sessions.pop(sid, None)
                return None, None
            s["exp"] = time.time() + SESSION_TTL
        u = self.users.get(s["user"])
        if not u or not u.get("active", True):
            self.drop_sessions(s["user"])
            return None, None
        return u, s["csrf"]

    def drop_sessions(self, username, keep=None):
        with self.lock:
            for k in [k for k, v in self.sessions.items() if v["user"] == username and k != keep]:
                del self.sessions[k]

    def visible(self, user):
        return [c for c in self.store.public() if self.users.can(user, c["mst"])]

    # ---- chạy đồng bộ ----
    def start_job(self, user, msts, kinds, start, end, mtt=True, xml=True, test=False, pdf=False, unattended=False):
        job = Job()
        job.running = True
        job.want_pdf = pdf
        job.unattended = unattended
        job.owner = user["username"]
        with self.lock:
            self.jobs[user["username"]] = job
        t = threading.Thread(target=self._run_job, daemon=True,
                             args=(job, msts, kinds, start, end, mtt, xml, test))
        t.start()
        return job, t

    def _run_job(self, job, msts, kinds, start, end, mtt, xml, test):
        run_batch(job, self.store, self.solver, msts, kinds, start, end, mtt, xml, self.out_root, test,
                  self.client_factory, self.clients)
        if not test:
            try:
                self.post_notices(job)
            except Exception as e:  # thông báo lỗi không được làm hỏng lượt đồng bộ
                job.say("Không ghi được thông báo: %s" % e)

    # ---- thông báo (chuông) ----
    def _notes_path(self):
        return os.path.join(self.cfg, "thong-bao.json")

    def notify(self, mst, text, level="info"):
        with self.lock:
            data = _load_json(self._notes_path(), {})
            seq = data.get("seq", 0) + 1
            items = data.get("items", [])
            items.append({"id": seq, "t": datetime.now().strftime("%d/%m/%Y %H:%M"), "mst": mst, "text": text,
                          "level": level})
            data.update(seq=seq, items=items[-2000:])
            _save_json(self._notes_path(), data)

    def notices_for(self, user, limit=100):
        with self.lock:
            data = _load_json(self._notes_path(), {})
        read = data.get("read", {}).get(user["username"], 0)
        items = [n for n in data.get("items", []) if not n["mst"] or self.users.can(user, n["mst"])]
        return {"items": items[::-1][:limit], "unread": sum(1 for n in items if n["id"] > read)}

    def mark_read(self, user):
        with self.lock:
            data = _load_json(self._notes_path(), {})
            data.setdefault("read", {})[user["username"]] = data.get("seq", 0)
            _save_json(self._notes_path(), data)

    def post_notices(self, job):
        """Sau mỗi lượt đồng bộ: báo HĐ đổi trạng thái (huỷ, thay thế, điều chỉnh); lượt tự động báo thêm HĐ mới và lỗi."""
        per = {}
        for r in job.rows:
            if r.get("mst"):
                per.setdefault(r["mst"], []).append(r)
        who = "Đồng bộ tự động" if job.unattended else "Đồng bộ (%s)" % job.owner
        for mst, rows in per.items():
            c = self.store.get(mst) or {}
            name = "%s (%s)" % (c.get("ten") or "", mst)
            changed = [r for r in rows if r["dongbo"] == "Đổi trạng thái"]
            new = [r for r in rows if r["dongbo"] == "Mới"]
            if changed:
                ex = ", ".join("%s %s%s/%s – %s" % (r["loai"], r["mau"], r["kh"], r["so"], r.get("tthai") or "?")
                               for r in changed[:5])
                self.notify(mst, "%s: %s có %d HĐ đổi trạng thái: %s%s" % (
                    who, name, len(changed), ex, " …" if len(changed) > 5 else ""), "warn")
            if new and job.unattended:
                self.notify(mst, "%s: %s có %d HĐ mới" % (who, name, len(new)))
        if job.unattended:
            for line in job.log:
                m = re.match(r"\S+\s+(\d{10}(?:-\d{3})?): LỖI (.*)", line)
                if m:
                    c = self.store.get(m.group(1)) or {}
                    self.notify(m.group(1), "%s: %s (%s) lỗi – %s" % (who, c.get("ten") or "", m.group(1), m.group(2)), "err")

    # ---- cài đặt đồng bộ tự động ----
    def _settings_path(self):
        return os.path.join(self.cfg, "cai-dat.json")

    def settings(self):
        d = _load_json(self._settings_path(), {})
        auto = {"on": False, "gio": "06:00", "ky": "auto", "lan_cuoi": ""}
        auto.update(d.get("auto") or {})
        return {"auto": auto}

    def save_auto(self, **fields):
        with self.lock:
            d = _load_json(self._settings_path(), {})
            auto = self.settings()["auto"]
            auto.update({k: v for k, v in fields.items() if v is not None})
            d["auto"] = auto
            _save_json(self._settings_path(), d)
            return auto

    def run_auto(self, wait=True):
        """Đồng bộ tự động tất cả DN đang hiển thị, có mật khẩu (mua vào + bán ra theo cài đặt từng DN)."""
        job = self.jobs.get(AUTO_USER["username"])
        if job and job.running:
            return job
        msts = [c["mst"] for c in self.store.public() if not c.get("an") and c.get("co_mk")]
        start, end = auto_range(self.settings()["auto"]["ky"])
        self.save_auto(lan_cuoi=datetime.now().strftime("%Y-%m-%d %H:%M"))
        job, t = self.start_job(AUTO_USER, msts, ["purchase", "sold"], start, end, unattended=True)
        job.say("Đồng bộ tự động %d doanh nghiệp, từ %s đến %s" % (
            len(msts), start.strftime("%d/%m/%Y"), end.strftime("%d/%m/%Y")))
        if wait:
            t.join()
        return job


def make_handler(app, port):
    allowed_hosts = {"127.0.0.1:%d" % port, "localhost:%d" % port}
    COOKIE = "thd_sid"

    class Handler(BaseHTTPRequestHandler):
        server_version = "TaiHoaDon/" + __version__

        def log_message(self, fmt, *args):  # yên lặng trên console
            pass

        def _send(self, code, body, ctype="application/json; charset=utf-8"):
            if not isinstance(body, bytes):
                body = json.dumps(body, ensure_ascii=False).encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self._sec_headers()
            for h in getattr(self, "_extra", []):
                self.send_header(*h)
            self.end_headers()
            self.wfile.write(body)

        def _sec_headers(self):
            self.send_header("X-Frame-Options", "DENY")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Referrer-Policy", "same-origin")

        def _https(self):
            return app.trust_proxy and (self.headers.get("X-Forwarded-Proto", "").lower() == "https"
                                        or self.headers.get("X-Forwarded-Ssl", "").lower() == "on")

        def _public_host(self):
            """Tên miền người dùng gõ; nginx (FASTPANEL) có thể đổi Host thành 127.0.0.1 và giữ tên thật ở X-Forwarded-Host."""
            if app.trust_proxy and self.headers.get("X-Forwarded-Host"):
                return self.headers["X-Forwarded-Host"].split(",")[0].strip()
            return self.headers.get("Host") or ""

        def _host_ok(self):
            if self._public_host().rsplit(":", 1)[0].lower() in app.allowed_hosts:
                return True
            # nginx của bảng điều khiển (FASTPANEL…) có thể gửi Host 127.0.0.1:cổng – chấp nhận khi app chỉ nghe trong máy
            host = (self.headers.get("Host") or "").rsplit(":", 1)[0]
            return (app.trust_proxy and not self.headers.get("X-Forwarded-Host")
                    and host in ("127.0.0.1", "localhost") and self.client_address[0] == "127.0.0.1")

        def _ip(self):
            if app.trust_proxy and self.headers.get("CF-Connecting-IP"):  # tên miền đi qua Cloudflare
                return self.headers["CF-Connecting-IP"].strip()
            if app.trust_proxy and self.headers.get("X-Forwarded-For"):
                return self.headers["X-Forwarded-For"].split(",")[-1].strip()
            return self.client_address[0]

        def _cookie(self, name):
            import http.cookies
            try:
                c = http.cookies.SimpleCookie(self.headers.get("Cookie", ""))
            except http.cookies.CookieError:
                return ""
            return c[name].value if name in c else ""

        def _set_cookie(self, value, max_age):
            self._extra = getattr(self, "_extra", []) + [("Set-Cookie", "%s=%s; Path=/; HttpOnly; SameSite=Lax; Max-Age=%d%s" % (
                COOKIE, value, max_age, "; Secure" if self._https() else ""))]

        def _k(self):
            """Khoá đặt trong link /file, /view (chỉ bản chạy trên máy; bản web dùng cookie đăng nhập)."""
            return "" if app.server else app.key

        def _guard(self):
            # Chặn DNS rebinding và yêu cầu giả mạo từ trang web khác.
            self._extra = []
            host = self.headers.get("Host") or ""
            if app.server:
                if app.allowed_hosts and not self._host_ok():
                    self._send(403, {"error": "Host không hợp lệ"})
                    return False
                self.user, self.csrf = app.session_user(self._cookie(COOKIE))
            else:
                if host not in allowed_hosts:
                    self._send(403, {"error": "Host không hợp lệ"})
                    return False
                self.user, self.csrf = LOCAL_USER, app.key
            path = self.path.split("?")[0]
            if path.startswith("/api/") and path != "/api/login":
                if not self.user:
                    self._send(401, {"error": "Phiên đăng nhập đã hết – tải lại trang để đăng nhập lại", "login": True})
                    return False
                if not secrets.compare_digest(self.headers.get("X-App-Key", ""), self.csrf):
                    self._send(403, {"error": "Phiên giao diện không hợp lệ, hãy tải lại trang"})
                    return False
            return True

        def _auto_state(self):
            a = app.settings()["auto"]
            job = app.jobs.get(AUTO_USER["username"])
            a["ky_text"] = AUTO_KY.get(a["ky"], "")
            a["running"] = bool(job and job.running)
            if job:
                a["current"], a["dn"] = job.current, getattr(job, "dn", None)
                a["log"] = job.log[-30:]
            return a

        def _can(self, mst):
            return bool(mst) and app.users.can(self.user, mst) and app.store.get(mst) is not None

        def _admin(self):
            if self.user["role"] != "admin":
                raise PermissionError("Chỉ quản trị mới được làm việc này")

        def _body(self):
            n = int(self.headers.get("Content-Length") or 0)
            if n > 80000000:
                raise ValueError("Dữ liệu gửi lên quá lớn (tối đa ~60MB một lần) – chọn ít file hơn")
            try:
                return json.loads(self.rfile.read(n).decode("utf-8") or "{}")
            except ValueError:
                return {}

        def do_GET(self):
            if not self._guard():
                return
            path = self.path.split("?")[0]
            if path == "/":
                if not self.user:
                    return self._send(200, LOGIN_PAGE.encode("utf-8"), "text/html; charset=utf-8")
                return self._send(200, PAGE.replace("__TOKEN__", self.csrf).encode("utf-8"), "text/html; charset=utf-8")
            if path == "/api/state":
                u = self.user
                return self._send(200, {"companies": app.visible(u), "job": app.job_for(u).snapshot(),
                                        "me": {"username": u["username"], "ten": u.get("ten", ""), "role": u["role"],
                                               "server": app.server},
                                        "unread": app.notices_for(u, 0)["unread"] if app.server else 0,
                                        "auto": self._auto_state() if app.server and u["role"] == "admin" else None,
                                        "captcha": {"count": app.solver.count(), "chars": app.solver.chars()}})
            if path == "/file":
                return self._file()
            if path == "/view":
                q = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(self.path).query))
                if not self._file_ok(q):
                    return self._send(403, {"error": "Không có quyền"})
                try:
                    page = render_invoice_html(app.out_root, q["mst"], q.get("kind"), q.get("key"))
                except (ValueError, OSError) as e:
                    return self._send(404, {"error": str(e)})
                return self._send(200, page.encode("utf-8"), "text/html; charset=utf-8")
            if path == "/cqt":  # PDF bản thể hiện theo dữ liệu cổng thuế (tạo khi cần, lưu lại)
                q = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(self.path).query))
                if not self._file_ok(q):
                    return self._send(403, {"error": "Không có quyền"})
                try:
                    rel = tax_pdf(app.out_root, q["mst"], q.get("kind"), q.get("key"), app.clients.get(q["mst"]))
                    with open(os.path.join(company_dir(app.out_root, q["mst"]), rel), "rb") as f:
                        body = f.read()
                except (ValueError, OSError) as e:
                    return self._send(404, {"error": str(e)})
                self._extra = [("Content-Disposition", 'inline; filename="%s"' % os.path.basename(rel))]
                return self._send(200, body, "application/pdf")
            self._send(404, {"error": "Không tìm thấy"})

        def _file_ok(self, q):
            if not self.user or not self._can(q.get("mst", "")):
                return False
            return app.server or secrets.compare_digest(q.get("k", ""), app.key)

        def _file(self):
            """Mở file đã tải (XML/HTML/PDF/Excel/ZIP) – chỉ trong thư mục của doanh nghiệp."""
            q = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(self.path).query))
            if not self._file_ok(q):
                return self._send(403, {"error": "Không có quyền"})
            base = os.path.realpath(company_dir(app.out_root, q["mst"]))
            full = os.path.realpath(os.path.join(base, q.get("p", "")))
            if not full.startswith(base + os.sep) or not os.path.isfile(full):
                return self._send(404, {"error": "Không tìm thấy file"})
            ext = os.path.splitext(full)[1].lower()
            ctype = {".xml": "application/xml", ".html": "text/html; charset=utf-8", ".htm": "text/html; charset=utf-8",
                     ".pdf": "application/pdf", ".zip": "application/zip",
                     ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(ext, "application/octet-stream")
            with open(full, "rb") as fh:
                body = fh.read()
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "private, no-store")
            self._sec_headers()
            if ext in (".zip", ".xlsx"):
                self.send_header("Content-Disposition", 'attachment; filename="%s"' % os.path.basename(full))
            if ext in (".html", ".htm"):
                self.send_header("Content-Security-Policy", "script-src 'none'")  # HTML hoá đơn chỉ để xem
            self.end_headers()
            self.wfile.write(body)

        def do_POST(self):
            if not self._guard():
                return
            path = self.path.split("?")[0]
            try:
                return self._post(path, self._body())
            except PermissionError as e:
                return self._send(403, {"error": str(e)})
            except ValueError as e:
                return self._send(400, {"error": str(e)})

        def _auth_post(self, path, data):
            """Đăng nhập / đăng xuất / đổi mật khẩu / quản lý người dùng (bản web)."""
            if path == "/api/login":
                if not app.server:
                    return self._send(200, {"ok": True})
                origin = self.headers.get("Origin")
                o = urllib.parse.urlparse(origin or "")
                if origin and o.netloc.lower() != self._public_host().lower() \
                        and (o.hostname or "").lower() not in app.allowed_hosts:
                    return self._send(403, {"error": "Yêu cầu không hợp lệ"})
                name = str(data.get("username") or "").strip().lower()[:60]
                ip = self._ip()
                if app.too_many_fails("ip:" + ip, "u:" + name):
                    return self._send(429, {"error": "Đăng nhập sai quá nhiều lần – thử lại sau 15 phút"})
                u = app.users.authenticate(name, str(data.get("password") or ""))
                if not u:
                    app.add_fail("ip:" + ip, "u:" + name)
                    app.audit(name or "?", "đăng nhập sai từ " + ip)
                    return self._send(401, {"error": "Sai tên đăng nhập hoặc mật khẩu"})
                with app.lock:
                    app.fails.pop("u:" + name, None)
                self._set_cookie(app.new_session(u["username"]), SESSION_TTL)
                app.audit(u["username"], "đăng nhập từ " + ip)
                return self._send(200, {"ok": True})
            if path == "/api/logout":
                sid = self._cookie(COOKIE)
                with app.lock:
                    app.sessions.pop(sid, None)
                self._set_cookie("", 0)
                return self._send(200, {"ok": True})
            if path == "/api/me/password":
                if not app.server:
                    raise ValueError("Bản chạy trên máy không có mật khẩu đăng nhập")
                if not UserStore.check(str(data.get("old") or ""), self.user["pw"]):
                    raise ValueError("Mật khẩu hiện tại không đúng")
                app.users.upsert(self.user["username"], password=str(data.get("new") or "") or None)
                app.drop_sessions(self.user["username"], keep=self._cookie(COOKIE))
                app.audit(self.user["username"], "đổi mật khẩu")
                return self._send(200, {"ok": True})
            self._admin()
            if path == "/api/users":
                return self._send(200, {"users": app.users.public(),
                                        "companies": [{"mst": c["mst"], "ten": c["ten"]} for c in app.store.public()]})
            if path == "/api/user/save":
                msts = data.get("msts")
                if msts is not None:
                    msts = [m for m in msts if app.store.get(m)]
                u = app.users.upsert(data.get("username"), ten=data.get("ten"), role=data.get("role"),
                                     password=data.get("password") or None, msts=msts, active=data.get("active"),
                                     create=bool(data.get("create")))
                if data.get("password") or not u.get("active", True):
                    app.drop_sessions(u["username"], keep=self._cookie(COOKIE))
                app.audit(self.user["username"], "lưu người dùng " + u["username"])
                return self._send(200, {"ok": True})
            if path == "/api/user/delete":
                if str(data.get("username") or "").lower() == self.user["username"]:
                    raise ValueError("Không tự xoá tài khoản đang đăng nhập")
                app.users.delete(data.get("username"))
                app.drop_sessions(str(data.get("username") or "").lower())
                app.audit(self.user["username"], "xoá người dùng %s" % data.get("username"))
                return self._send(200, {"ok": True})
            self._send(404, {"error": "Không tìm thấy"})

        def _tool(self, name, data):
            import base64

            def b64(v):
                try:
                    return base64.b64decode(v or "", validate=True)
                except ValueError:
                    raise ValueError("Dữ liệu file không hợp lệ")
            if name == "xml-view":
                return self._send(200, {"html": render_xml_invoice(b64(data.get("data")))})
            if name == "mst":
                msts = [m for m in re.split(r"[\s,;]+", data.get("text") or "") if m][:200]
                client, rows = app.client_factory(), []
                for m in msts:
                    try:
                        d = client.lookup_company(m)
                        rows.append(dict(d, loi=""))
                    except (PortalError, ValueError) as e:
                        rows.append({"mst": m, "ten": "", "dia_chi": "", "cqt": "", "tthai_text": "", "loi": str(e)})
                return self._send(200, {"rows": rows})
            if name == "bank-parse":
                infos, errors = [], []
                for f in (data.get("files") or [])[:24]:
                    try:
                        infos.append(parse_bank_statement(f.get("name") or "sao-ke", b64(f.get("data"))))
                    except ValueError as e:
                        errors.append(str(e))
                rows, dups = merge_bank_statements(infos)
                for i in infos:
                    i.pop("rows")
                return self._send(200, {"statements": infos, "rows": rows, "dups": dups, "errors": errors})
            if name == "bank-export":
                rows = [r for r in (data.get("rows") or []) if isinstance(r, dict) and r.get("date")]
                if not rows:
                    raise ValueError("Chưa có giao dịch nào")
                for r in rows:
                    r["in"], r["out"] = _float(r.get("in")), _float(r.get("out"))
                    datetime.strptime(str(r["date"]), "%Y-%m-%d")
                out = write_bank_ktsc(rows, data.get("opts") or {})
                ds = sorted(r["date"] for r in rows)
                label = safe_name("_".join(x for x in ("SAO_KE", str(data.get("bank") or ""), str(data.get("account") or ""),
                                                       ds[0].replace("-", ""), ds[-1].replace("-", "")) if x))
                return self._send(200, {"name": label + ".xlsx", "data": base64.b64encode(out).decode()})
            if name == "pdf-merge":
                files = [(f.get("name"), b64(f.get("data"))) for f in (data.get("files") or [])]
                out = merge_pdfs(files)
                return self._send(200, {"name": "noi_file_%s.pdf" % datetime.now().strftime("%Y%m%d_%H%M%S"),
                                        "data": base64.b64encode(out).decode()})
            if name == "pdf-split":
                out = split_pdf(data.get("name"), b64(data.get("data")), data.get("ranges"))
                base = os.path.splitext(os.path.basename(data.get("name") or "file"))[0]
                return self._send(200, {"name": "%s_tach.zip" % safe_name(base), "data": base64.b64encode(out).decode()})
            self._send(404, {"error": "Không tìm thấy"})

        def _post(self, path, data):
            if path in ("/api/login", "/api/logout", "/api/me/password", "/api/users", "/api/user/save",
                        "/api/user/delete"):
                return self._auth_post(path, data)
            if path == "/api/company/save":
                existing = app.store.get(str(data.get("mst", "")).strip())
                if existing and not self._can(existing["mst"]):
                    raise PermissionError("Không có quyền với doanh nghiệp này")
                if not existing:
                    self._admin()
                if not existing and not data.get("password"):
                    raise ValueError("Nhập mật khẩu trang hoadondientu.gdt.gov.vn")
                fields = {k: data.get(k) for k in ("ghichu", "an")}
                for k in ("vao", "ra"):
                    fields[k] = data.get(k)
                c = app.store.upsert(data.get("mst"), data.get("ten"), data.get("password") or None, **fields)
                return self._send(200, {"mst": c["mst"]})
            if path.startswith("/api/tool/"):
                return self._tool(path[len("/api/tool/"):], data)
            if path == "/api/mst-lookup":
                try:
                    return self._send(200, app.client_factory().lookup_company(data.get("mst")))
                except PortalError as e:
                    raise ValueError(str(e))
            if path == "/api/company/delete":
                self._admin()
                app.store.delete(data.get("mst"))
                app.users.remove_mst(data.get("mst"))
                app.audit(self.user["username"], "xoá DN %s" % data.get("mst"))
                return self._send(200, {"ok": True})
            if path == "/api/company/import":
                self._admin()
                added, errors = app.store.import_text(data.get("text"))
                client, filled = None, 0
                for c in app.store.public():
                    if c["ten"] or filled >= 100:
                        continue
                    try:
                        client = client or app.client_factory()
                        app.store.upsert(c["mst"], client.lookup_company(c["mst"])["ten"])
                        filled += 1
                    except (PortalError, ValueError) as e:
                        errors.append("%s: không lấy được tên (%s)" % (c["mst"], e))
                return self._send(200, {"added": added, "errors": errors})
            if path == "/api/run":
                job = app.job_for(self.user)
                if job.running:
                    return self._send(409, {"error": "Đang xử lý, vui lòng chờ hoặc bấm Dừng"})
                if app.server and app.running_jobs() >= app.max_jobs:
                    return self._send(429, {"error": "Máy chủ đang chạy %d lượt đồng bộ – chờ một chút rồi thử lại" % app.max_jobs})
                msts = [m for m in (data.get("msts") or []) if self._can(m)]
                if not msts:
                    raise ValueError("Chưa có doanh nghiệp nào để xử lý")
                busy = busy_msts()
                if all(m in busy for m in msts):
                    raise ValueError("Doanh nghiệp đang được đồng bộ (bởi bạn hoặc người khác) – chờ xong rồi thử lại")
                test = bool(data.get("test"))
                start = end = None
                kinds = []
                if not test:
                    start, end = parse_date(data.get("from")), parse_date(data.get("to"))
                    if end < start:
                        raise ValueError("Ngày kết thúc phải sau ngày bắt đầu")
                    kinds = [k for k in (data.get("kinds") or []) if k in ("purchase", "sold")]
                    if not kinds:
                        raise ValueError("Chọn ít nhất mua vào hoặc bán ra")
                app.start_job(self.user, msts, kinds, start, end, bool(data.get("mtt")), bool(data.get("xml")), test,
                              bool(data.get("pdf")))
                return self._send(200, {"ok": True})
            if path == "/api/notices":
                if data.get("read"):
                    app.mark_read(self.user)
                return self._send(200, app.notices_for(self.user))
            if path in ("/api/auto/save", "/api/auto/run"):
                self._admin()
                if path == "/api/auto/run":
                    job = app.jobs.get(AUTO_USER["username"])
                    if job and job.running:
                        raise ValueError("Đồng bộ tự động đang chạy")
                    app.run_auto(wait=False)
                    app.audit(self.user["username"], "chạy đồng bộ tự động")
                    return self._send(200, {"ok": True})
                gio = str(data.get("gio") or "06:00")
                if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", gio):
                    raise ValueError("Giờ không hợp lệ (vd 06:00)")
                if data.get("ky") not in AUTO_KY:
                    raise ValueError("Kỳ không hợp lệ")
                return self._send(200, app.save_auto(on=bool(data.get("on")), gio=gio, ky=data["ky"]))
            if path == "/api/captcha":
                app.job_for(self.user).answer(data.get("answer"))
                return self._send(200, {"ok": True})
            if path in ("/api/invoices", "/api/export", "/api/invoice/update", "/api/invoice/update-bulk", "/api/invoice/pdf",
                        "/api/invoice/fetch-pdf", "/api/invoice/fetch-pdf-bulk", "/api/invoice/pdf-import",
                        "/api/invoice/scan-downloads"):
                mst = data.get("mst", "")
                company = app.store.get(mst)
                if not company:
                    raise ValueError("Chọn doanh nghiệp")
                if not self._can(mst):
                    raise PermissionError("Không có quyền với doanh nghiệp này")
                if path == "/api/invoices":
                    return self._send(200, {"rows": query_invoices(app.out_root, mst, data.get("filters") or {})})
                if path == "/api/invoice/scan-downloads":
                    if app.server:  # máy chủ không nhìn thấy thư mục Downloads trên máy người dùng
                        return self._send(200, {"results": [], "folders": [], "folder": "", "prompt": False, "server": True})
                    folders, prompt = watch_dirs(app)
                    if data.get("info_only"):
                        return self._send(200, {"folders": folders, "prompt": prompt})
                    res = scan_downloads(app.out_root, mst, folders, float(data.get("since") or 0), app.seen_downloads)
                    return self._send(200, {"results": res, "folder": ", ".join(folders), "prompt": prompt})
                if path == "/api/invoice/pdf-import":
                    import base64
                    files = []
                    for f in (data.get("files") or [])[:500]:
                        try:
                            files.append((f.get("name"), base64.b64decode(f.get("data") or "", validate=True)))
                        except ValueError:
                            files.append((f.get("name"), b""))
                    return self._send(200, {"results": import_pdfs(app.out_root, mst, files)})
                if path == "/api/invoice/fetch-pdf":
                    try:
                        return self._send(200, {"pdf": fetch_original_pdf(app.out_root, mst, data.get("kind"), data.get("key"))})
                    except PortalError as e:
                        raise ValueError(str(e))
                if path == "/api/invoice/fetch-pdf-bulk":
                    f = data.get("filters") or {}
                    kind = (f.get("kind") or "purchase").replace("_dv", "")
                    rows = [r for r in query_invoices(app.out_root, mst, f)
                            if not r["pdf"] and can_fetch_pdf(r.get("tra_cuu"))[0]][:300]
                    ok, errs = 0, []
                    for r in rows:
                        try:
                            fetch_original_pdf(app.out_root, mst, kind, r["key"])
                            ok += 1
                        except (PortalError, ValueError) as e:
                            errs.append("%s/%s: %s" % (r["khhdon"], r["shdon"], e))
                        time.sleep(0.3)
                    return self._send(200, {"ok": ok, "total": len(rows), "errors": errs[:20]})
                if path == "/api/invoice/pdf":
                    import base64
                    try:
                        raw = base64.b64decode(data.get("data") or "", validate=True)
                    except ValueError:
                        raise ValueError("Dữ liệu file không hợp lệ")
                    return self._send(200, {"pdf": attach_pdf(app.out_root, mst, data.get("kind"), data.get("key"), raw)})
                if path == "/api/invoice/update-bulk":
                    if data.get("duyet") is not None and data["duyet"] not in DUYET_OPTIONS:
                        raise ValueError("Trạng thái duyệt không hợp lệ")
                    keys = [str(k) for k in (data.get("keys") or [])][:5000]
                    n = update_invoices(app.out_root, mst, data.get("kind"), keys, duyet=data.get("duyet"),
                                        dv=data.get("dv"))
                    return self._send(200, {"ok": True, "updated": n})
                if path == "/api/invoice/update":
                    if data.get("duyet") is not None and data["duyet"] not in DUYET_OPTIONS:
                        raise ValueError("Trạng thái duyệt không hợp lệ")
                    update_invoice(app.out_root, mst, data.get("kind"), data.get("key"), note=data.get("note"),
                                   duyet=data.get("duyet"), dv=data.get("dv"))
                    return self._send(200, {"ok": True})
                out = export_invoices(app.out_root, mst, company.get("ten"), data.get("filters") or {}, data.get("fmt"))
                rel = os.path.relpath(out, company_dir(app.out_root, mst))
                return self._send(200, {"url": "/file?" + urllib.parse.urlencode({"k": self._k(), "mst": mst, "p": rel})})
            if path == "/api/stop":
                app.job_for(self.user).stop()
                return self._send(200, {"ok": True})
            self._send(404, {"error": "Không tìm thấy"})

    return Handler


def main(argv=None):
    ap = argparse.ArgumentParser(description="Phần mềm tải hoá đơn điện tử (hoadondientu.gdt.gov.vn)")
    ap.add_argument("--port", type=int, default=int(os.environ.get("TAIHOADON_PORT") or 8765),
                    help="cổng giao diện (mặc định 8765)")
    ap.add_argument("--out", default=os.environ.get("TAIHOADON_DATA") or os.path.join(os.getcwd(), "HoaDon"),
                    help="thư mục lưu (mặc định ./HoaDon)")
    ap.add_argument("--no-browser", action="store_true", help="không tự mở trình duyệt")
    ap.add_argument("--insecure", action="store_true", help="bỏ kiểm tra chứng chỉ SSL (chỉ dùng khi máy báo lỗi SSL)")
    ap.add_argument("--downloads", help="thư mục trình duyệt lưu file tải về (mặc định: Downloads của Windows)")
    web = ap.add_argument_group("bản web (máy chủ, nhiều người dùng)")
    web.add_argument("--server", action="store_true", default=os.environ.get("TAIHOADON_SERVER") == "1",
                     help="chạy bản web: đăng nhập, phân quyền nhân viên theo doanh nghiệp")
    web.add_argument("--bind", default=os.environ.get("TAIHOADON_BIND") or "0.0.0.0", help="địa chỉ lắng nghe (mặc định 0.0.0.0)")
    web.add_argument("--domain", default=os.environ.get("TAIHOADON_DOMAIN") or "",
                     help="tên miền được phép, ngăn cách bằng dấu phẩy (vd hoadon.congty.vn)")
    web.add_argument("--trust-proxy", action="store_true", default=os.environ.get("TAIHOADON_TRUST_PROXY") == "1",
                     help="chạy sau Caddy/Nginx (HTTPS): tin X-Forwarded-For / X-Forwarded-Proto")
    web.add_argument("--max-jobs", type=int, default=int(os.environ.get("TAIHOADON_MAX_JOBS") or 4),
                     help="số lượt đồng bộ chạy cùng lúc (mặc định 4)")
    web.add_argument("--set-password", metavar="TÊN", help="tạo / đặt lại mật khẩu quản trị TÊN rồi thoát")
    args = ap.parse_args(argv)

    out = os.path.abspath(args.out)
    cfg = os.path.join(out, "_cau-hinh")
    if args.server or args.set_password:
        set_secret_key(os.path.join(cfg, "khoa-bi-mat.key"), os.environ.get("TAIHOADON_SECRET_KEY"))
    app = App(out, lambda: HoaDonClient(verify_ssl=not args.insecure), server=args.server)

    if args.set_password:
        import getpass
        pw = os.environ.get("TAIHOADON_NEW_PASSWORD") or getpass.getpass("Mật khẩu mới cho %s: " % args.set_password)
        exists = app.users.get(args.set_password)
        app.users.upsert(args.set_password, password=pw, role="admin", active=True,
                         ten=None if exists else "Quản trị")
        print("Đã đặt mật khẩu quản trị cho %s" % args.set_password.lower())
        return 0

    if args.downloads:
        app.downloads = os.path.abspath(args.downloads)
    host = "127.0.0.1"
    if app.server:
        host = args.bind
        app.allowed_hosts = {h.strip().lower() for h in args.domain.split(",") if h.strip()}
        app.trust_proxy = args.trust_proxy
        app.max_jobs = max(1, args.max_jobs)
        n = app.store.reencrypt()
        if n:
            print("Đã mã hoá lại %d mật khẩu doanh nghiệp bằng khoá máy chủ." % n)
        if not app.users.count():
            name = (os.environ.get("TAIHOADON_ADMIN_USER") or "admin").lower()
            pw = os.environ.get("TAIHOADON_ADMIN_PASSWORD") or secrets.token_urlsafe(9)
            app.users.upsert(name, ten="Quản trị", role="admin", password=pw)
            print("=" * 60)
            print("Tạo tài khoản quản trị đầu tiên: %s" % name)
            if not os.environ.get("TAIHOADON_ADMIN_PASSWORD"):
                print("Mật khẩu: %s   (đăng nhập rồi đổi ngay)" % pw)
                # để xem được qua Trình quản lý file của bảng điều khiển khi không có SSH; xoá file sau khi đăng nhập
                note = os.path.join(cfg, "MAT-KHAU-ADMIN-BAN-DAU.txt")
                fd = os.open(note, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
                with os.fdopen(fd, "w", encoding="utf-8") as f:
                    f.write("Tên đăng nhập: %s\nMật khẩu: %s\nĐăng nhập, đổi mật khẩu rồi XOÁ file này.\n" % (name, pw))
            print("=" * 60)
    server = ThreadingHTTPServer((host, args.port), make_handler(app, args.port))
    if app.server:
        threading.Thread(target=auto_sync_loop, args=(app,), daemon=True).start()
        print("Bản web %s đang chạy tại %s:%d%s" % (__version__, host, args.port,
                                                    (" – tên miền: " + ", ".join(sorted(app.allowed_hosts))) if app.allowed_hosts else ""))
    else:
        url = "http://127.0.0.1:%d/" % args.port
        print("Phần mềm tải hoá đơn %s đang chạy tại %s" % (__version__, url))
        if os.name != "nt":
            print("Lưu ý: ngoài Windows, mật khẩu lưu dạng mã hoá đơn giản – hãy bảo vệ thư mục _cau-hinh.")
        print("Đóng cửa sổ này (hoặc Ctrl+C) để thoát.")
        if not args.no_browser:
            threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    print("Hoá đơn lưu vào: %s" % app.out_root)
    sys.stdout.flush()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    sys.exit(main())
